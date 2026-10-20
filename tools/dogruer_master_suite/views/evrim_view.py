# -*- coding: utf-8 -*-
"""
evrim_view.py — Evrim Canlı Veri Gezgini ve Detaylı Dosya Sorgulama
MasterSuite için Evrim gümrük kayıtları tablosu, filtreleme, detay künyesi, vergi maili, QNB/GLP/Talep/GTB entegrasyonları.
"""
import os
import sys
import re
import time
import json
import threading
from datetime import datetime
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

try:
    import openpyxl
    import openpyxl.styles
    import openpyxl.utils
except ImportError:
    openpyxl = None

from .base_view import BaseView


class EvrimView(BaseView):
    """Evrim Canlı Veri Gezgini Sekmesi."""
    def __init__(self, parent, app, *args, **kwargs):
        super().__init__(parent, app, *args, **kwargs)
        self.current_evrim_filter = "ALL"
        self.filtered_evrim_records = []
        self.selected_evrim_rec = None
        self.ev_detail_labels = {}
        self._sort_reverse_map = {}
        self.setup_ui()

    @property
    def evrim_records(self):
        return getattr(self.app, "evrim_records", [])

    @evrim_records.setter
    def evrim_records(self, val):
        self.app.evrim_records = val

    @property
    def gtb_tracker(self):
        return getattr(self.app, "gtb_tracker", None)

    def setup_ui(self):
        # 1. Üst Metrik Bar
        metric_bar = tk.Frame(self, bg="#f8fafc")
        metric_bar.pack(fill=tk.X, pady=(0, 10))

        self.card_ev_total = self.create_metric_widget(
            metric_bar, "TOPLAM EVRİM KAYDI", "0", "Hafızadaki Dosyalar", "#0f172a",
            lambda: self.set_evrim_filter("ALL")
        )
        self.card_ev_beyan = self.create_metric_widget(
            metric_bar, "BEYANNAMELİ / TESCİLLİ", "0", "Tescil No Alınmış", "#16a34a",
            lambda: self.set_evrim_filter("TESCILLI")
        )
        self.card_ev_pending = self.create_metric_widget(
            metric_bar, "İŞLEMDE / BEKLEYEN", "0", "Tescil Bekleyen", "#f59e0b",
            lambda: self.set_evrim_filter("PENDING")
        )
        self.card_ev_closed = self.create_metric_widget(
            metric_bar, "TESLİM EDİLENLER", "0", "Teslimatı Yapılan", "#2563eb",
            lambda: self.set_evrim_filter("CLOSED")
        )
        self.card_ev_ports = self.create_metric_widget(
            metric_bar, "LİMAN / ANTREPOLAR", "0", "Farklı Nokta", "#7c3aed",
            lambda: self.focus_location_filter()
        )

        # 2. Arama ve Filtreleme Barı (2 Satırlı Modern Tasarım)
        filter_container = tk.Frame(self, bg="#ffffff", padx=14, pady=10, relief=tk.FLAT, bd=0)
        filter_container.pack(fill=tk.X, pady=(0, 10))

        # ── Satır 1: Anlık Arama, Hızlı Filtre Butonları ve Ana Aksiyonlar ──
        row1 = tk.Frame(filter_container, bg="#ffffff")
        row1.pack(fill=tk.X, pady=(0, 8))

        tk.Label(row1, text="🔍 Anlık Arama:", font=("Segoe UI", 9, "bold"), fg="#0f172a", bg="#ffffff").pack(side=tk.LEFT, padx=(0, 6))
        self.txt_ev_search = tk.Entry(row1, font=("Segoe UI", 10), width=28, bd=1, relief=tk.SOLID)
        self.txt_ev_search.pack(side=tk.LEFT, ipady=3, padx=(0, 12))
        self.txt_ev_search.bind("<KeyRelease>", lambda e: self.apply_evrim_filters())

        # Hızlı Butonlar
        self.btn_ev_f_all = tk.Button(row1, text="Tümü", font=("Segoe UI", 8, "bold"), bg="#0f172a", fg="#ffffff", relief=tk.FLAT, padx=8, pady=3, cursor="hand2", command=lambda: self.set_evrim_filter("ALL"))
        self.btn_ev_f_all.pack(side=tk.LEFT, padx=2)

        self.btn_ev_f_beyan = tk.Button(row1, text="🟢 Tescilliler", font=("Segoe UI", 8), bg="#f1f5f9", fg="#16a34a", relief=tk.FLAT, bd=0, padx=8, pady=3, cursor="hand2", command=lambda: self.set_evrim_filter("TESCILLI"))
        self.btn_ev_f_beyan.pack(side=tk.LEFT, padx=2)

        self.btn_ev_f_pend = tk.Button(row1, text="⏳ İşlemde", font=("Segoe UI", 8), bg="#f1f5f9", fg="#d97706", relief=tk.FLAT, bd=0, padx=8, pady=3, cursor="hand2", command=lambda: self.set_evrim_filter("PENDING"))
        self.btn_ev_f_pend.pack(side=tk.LEFT, padx=2)

        self.btn_ev_f_close = tk.Button(row1, text="✅ Teslim Edilen", font=("Segoe UI", 8), bg="#f1f5f9", fg="#2563eb", relief=tk.FLAT, bd=0, padx=8, pady=3, cursor="hand2", command=lambda: self.set_evrim_filter("CLOSED"))
        self.btn_ev_f_close.pack(side=tk.LEFT, padx=2)

        # Sağ Aksiyon Butonları
        right_f = tk.Frame(row1, bg="#ffffff")
        right_f.pack(side=tk.RIGHT)

        btn_ev_sync = tk.Button(
            right_f,
            text="🔄 Doğruer'dan Canlı Çek",
            font=("Segoe UI", 8, "bold"),
            bg="#2563eb",
            fg="#ffffff",
            activebackground="#1d4ed8",
            activeforeground="#ffffff",
            relief=tk.FLAT,
            padx=10,
            pady=4,
            cursor="hand2",
            command=self.trigger_dogruer_sync
        )
        btn_ev_sync.pack(side=tk.LEFT, padx=4)

        btn_ev_excel = tk.Button(
            right_f,
            text="📥 Excel'e Aktar",
            font=("Segoe UI", 8, "bold"),
            bg="#107c41",
            fg="#ffffff",
            activebackground="#0b5e31",
            activeforeground="#ffffff",
            relief=tk.FLAT,
            padx=10,
            pady=4,
            cursor="hand2",
            command=self.export_evrim_to_excel
        )
        btn_ev_excel.pack(side=tk.LEFT, padx=4)

        # ── Satır 2: Detaylı Combobox Filtreleri (Liman, Antrepo, Statü) ──
        row2 = tk.Frame(filter_container, bg="#ffffff")
        row2.pack(fill=tk.X)

        # 1. Liman Filtresi
        tk.Label(row2, text="⚓ Liman / Konum:", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#ffffff").pack(side=tk.LEFT, padx=(0, 4))
        self.cmb_ev_port = ttk.Combobox(row2, font=("Segoe UI", 8), state="readonly", width=16)
        self.cmb_ev_port.pack(side=tk.LEFT, padx=(0, 10))
        self.cmb_ev_port.bind("<<ComboboxSelected>>", lambda e: self.apply_evrim_filters())

        # 2. Antrepo Filtresi
        tk.Label(row2, text="🏢 Antrepo / Depo:", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#ffffff").pack(side=tk.LEFT, padx=(0, 4))
        self.cmb_ev_antrepo = ttk.Combobox(row2, font=("Segoe UI", 8), state="readonly", width=16)
        self.cmb_ev_antrepo.pack(side=tk.LEFT, padx=(0, 10))
        self.cmb_ev_antrepo.bind("<<ComboboxSelected>>", lambda e: self.apply_evrim_filters())

        # 3. Statü Filtresi
        tk.Label(row2, text="📌 Statü:", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#ffffff").pack(side=tk.LEFT, padx=(0, 4))
        self.cmb_ev_statu = ttk.Combobox(row2, font=("Segoe UI", 8), state="readonly", width=22)
        self.cmb_ev_statu.pack(side=tk.LEFT, padx=(0, 10))
        self.cmb_ev_statu.bind("<<ComboboxSelected>>", lambda e: self.apply_evrim_filters())

        # Filtreleri Temizle Butonu
        btn_reset_f = tk.Button(
            row2,
            text="🧹 Filtreleri Temizle",
            font=("Segoe UI", 8),
            bg="#f1f5f9",
            fg="#64748b",
            activebackground="#e2e8f0",
            relief=tk.FLAT,
            padx=8,
            pady=2,
            cursor="hand2",
            command=self.reset_filters
        )
        btn_reset_f.pack(side=tk.LEFT, padx=4)

        # Sonuç Sayacı Rozeti
        self.lbl_ev_result_count = tk.Label(
            row2,
            text="0 Kayıt Listeleniyor",
            font=("Segoe UI", 8, "bold"),
            fg="#1e40af",
            bg="#eff6ff",
            padx=8,
            pady=2
        )
        self.lbl_ev_result_count.pack(side=tk.RIGHT)

        # 3. İkili Bölüntülü Tablo ve Künye Paneli
        paned = tk.PanedWindow(self, orient=tk.HORIZONTAL, bg="#f8fafc", bd=0, sashwidth=6)
        paned.pack(fill=tk.BOTH, expand=True)

        # Sol Tablo Çerçevesi
        table_frame = tk.Frame(paned, bg="#ffffff", relief=tk.FLAT, bd=0)
        paned.add(table_frame, minsize=650)

        ev_cols = (
            "dosya_no", "hankook_ref", "beyanname_no", "beyanname_tarihi",
            "statu", "asama", "liman", "gumruk", "teslim_tarihi", "yurt_disi_firma"
        )
        self.tree_evrim = ttk.Treeview(
            table_frame,
            columns=ev_cols,
            show="headings",
            selectmode="extended"
        )

        headers = [
            ("dosya_no", "Evrim Dosya No", 105),
            ("hankook_ref", "Hankook Ref (F.Ref)", 140),
            ("beyanname_no", "Beyanname No", 145),
            ("beyanname_tarihi", "Tescil Tarihi", 95),
            ("statu", "Statü Açıklama", 205),
            ("asama", "Aşama / %", 105),
            ("liman", "Bulunduğu Yer / Liman", 140),
            ("gumruk", "Gümrük Müdürlüğü", 150),
            ("teslim_tarihi", "Teslim Tarihi", 95),
            ("yurt_disi_firma", "Yurt Dışı Firma", 190)
        ]

        for col_id, col_name, width in headers:
            self.tree_evrim.heading(col_id, text=col_name, anchor="w", command=lambda c=col_id: self.sort_evrim_tree(c))
            self.tree_evrim.column(col_id, width=width, anchor="w")

        ev_scroll_y = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.tree_evrim.yview)
        ev_scroll_x = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL, command=self.tree_evrim.xview)
        self.tree_evrim.configure(yscrollcommand=ev_scroll_y.set, xscrollcommand=ev_scroll_x.set)

        self.tree_evrim.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ev_scroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        ev_scroll_x.pack(side=tk.BOTTOM, fill=tk.X)

        self.tree_evrim.bind("<<TreeviewSelect>>", self.on_evrim_row_select)
        self.tree_evrim.bind("<Control-a>", self.select_all_evrim_rows)
        self.tree_evrim.bind("<Control-A>", self.select_all_evrim_rows)
        self.tree_evrim.bind("<Double-1>", lambda e: self.copy_evrim_details())

        # Sağ Tık Menüsü Kurulumu
        self.setup_context_menu()

        # Tablo Renk Etiketleri
        self.tree_evrim.tag_configure("tescilli", background="#f0fdf4", foreground="#166534")
        self.tree_evrim.tag_configure("pending", background="#fffbeb", foreground="#92400e")
        self.tree_evrim.tag_configure("closed", background="#eff6ff", foreground="#1e40af")

        # Sağ Detay Künyesi
        detail_frame = tk.Frame(paned, bg="#ffffff", padx=16, pady=16, relief=tk.FLAT, bd=0, width=380)
        paned.add(detail_frame, minsize=320)

        lbl_det_title = tk.Label(
            detail_frame,
            text="📋 SEÇİLEN DOSYA KÜNYESİ",
            font=("Segoe UI", 11, "bold"),
            fg="#0f172a",
            bg="#ffffff"
        )
        lbl_det_title.pack(anchor="w", pady=(0, 10))

        self.lbl_ev_status_badge = tk.Label(
            detail_frame,
            text="🔍 Lütfen tablodan bir dosya seçin",
            font=("Segoe UI", 9, "bold"),
            bg="#f1f5f9",
            fg="#64748b",
            padx=10,
            pady=6
        )
        self.lbl_ev_status_badge.pack(fill=tk.X, pady=(0, 14))

        # 11 Bilgi Satırı
        self.ev_detail_labels = {}
        fields = [
            ("Hankook Referansı:", "ref"),
            ("Evrim Dosya No:", "dosya"),
            ("Beyanname No:", "beyan_no"),
            ("Tescil Tarihi:", "beyan_tarih"),
            ("Gümrük Müdürlüğü:", "gumruk"),
            ("Bulunduğu Yer / Liman:", "liman"),
            ("Rejim / Tip:", "rejim"),
            ("Statü Durumu:", "statu"),
            ("Aşama & İlerleme:", "asama"),
            ("Teslimat Tarihi:", "teslim"),
            ("Yurt Dışı Firma:", "yurt_disi")
        ]

        for label_text, key in fields:
            row = tk.Frame(detail_frame, bg="#ffffff")
            row.pack(fill=tk.X, pady=2)

            lbl_k = tk.Label(row, text=label_text, font=("Segoe UI", 8, "bold"), fg="#64748b", bg="#ffffff", width=18, anchor="w")
            lbl_k.pack(side=tk.LEFT)

            lbl_v = tk.Label(row, text="—", font=("Segoe UI", 8), fg="#0f172a", bg="#ffffff", anchor="w", wraplength=180, justify=tk.LEFT)
            lbl_v.pack(side=tk.LEFT, fill=tk.X, expand=True)

            self.ev_detail_labels[key] = lbl_v

        sep = tk.Frame(detail_frame, bg="#e2e8f0", height=1)
        sep.pack(fill=tk.X, pady=14)

        # Hızlı Aksiyon Butonları
        lbl_act = tk.Label(detail_frame, text="⚡ HIZLI AKSİYONLAR", font=("Segoe UI", 9, "bold"), fg="#0f172a", bg="#ffffff")
        lbl_act.pack(anchor="w", pady=(0, 8))

        self.btn_ev_mail = tk.Button(
            detail_frame,
            text="📧 Tek Tıkla Vergi Maili Hazırla",
            font=("Segoe UI", 9, "bold"),
            bg="#2563eb",
            fg="#ffffff",
            activebackground="#1d4ed8",
            activeforeground="#ffffff",
            relief=tk.FLAT,
            pady=6,
            cursor="hand2",
            command=self.create_tax_mail_from_evrim
        )
        self.btn_ev_mail.pack(fill=tk.X, pady=3)

        btn_copy = tk.Button(
            detail_frame,
            text="📋 Dosya & Beyanname No Kopyala",
            font=("Segoe UI", 8),
            bg="#f1f5f9",
            fg="#0f172a",
            relief=tk.SOLID,
            bd=1,
            pady=5,
            cursor="hand2",
            command=self.copy_evrim_details
        )
        btn_copy.pack(fill=tk.X, pady=3)

        btn_to_gtb = tk.Button(
            detail_frame,
            text="🏛️ GTB Takip Havuzuna Ekle",
            font=("Segoe UI", 8, "bold"),
            bg="#0284c7",
            fg="#ffffff",
            activebackground="#0369a1",
            activeforeground="#ffffff",
            relief=tk.FLAT,
            pady=5,
            cursor="hand2",
            command=self.send_to_gtb_tracking
        )
        btn_to_gtb.pack(fill=tk.X, pady=3)

        btn_to_qnb = tk.Button(
            detail_frame,
            text="📦 QNB Fatura Eşleştirici'ye Aktar",
            font=("Segoe UI", 9, "bold"),
            bg="#6b21a8",
            fg="#ffffff",
            activebackground="#581c87",
            activeforeground="#ffffff",
            relief=tk.FLAT,
            pady=6,
            cursor="hand2",
            command=self.send_to_qnb_matcher
        )
        btn_to_qnb.pack(fill=tk.X, pady=3)

        btn_to_glp = tk.Button(
            detail_frame,
            text="🚛 GLP Konşimento / Ref Sorgula",
            font=("Segoe UI", 8),
            bg="#f1f5f9",
            fg="#0f172a",
            relief=tk.SOLID,
            bd=1,
            pady=5,
            cursor="hand2",
            command=self.send_to_glp_query
        )
        btn_to_glp.pack(fill=tk.X, pady=3)

        btn_to_talep = tk.Button(
            detail_frame,
            text="💸 Bu Dosya İçin Para Talebi Aç",
            font=("Segoe UI", 9, "bold"),
            bg="#107c41",
            fg="#ffffff",
            activebackground="#0b5e31",
            activeforeground="#ffffff",
            relief=tk.FLAT,
            pady=6,
            cursor="hand2",
            command=self.send_selected_evrim_to_talep
        )
        btn_to_talep.pack(fill=tk.X, pady=3)

        # İlk filtreleme ve render
        self.apply_evrim_filters()

    def setup_context_menu(self):
        """Sağ tık bağlam menüsü kurulumu."""
        self.context_menu = tk.Menu(self.tree_evrim, tearoff=0)
        self.context_menu.add_command(
            label="📧 Vergi Talebi Maili Aç",
            command=self.create_tax_mail_from_evrim
        )
        self.context_menu.add_command(
            label="🌐 QNB Fatura Eşleştiriciye Gönder",
            command=self.send_to_qnb_matcher
        )
        self.context_menu.add_command(
            label="🚛 GLP Referans Sorgusuna Gönder",
            command=self.send_to_glp_query
        )
        self.context_menu.add_command(
            label="💸 Para Talebine Gönder",
            command=self.send_selected_evrim_to_talep
        )
        self.context_menu.add_command(
            label="🏛️ GTB Takip Havuzuna Ekle",
            command=self.send_to_gtb_tracking
        )
        self.context_menu.add_separator()
        self.context_menu.add_command(
            label="📊 Excel'e Aktar",
            command=self.export_evrim_to_excel
        )
        self.context_menu.add_command(
            label="📋 Dosya / Beyanname Bilgilerini Kopyala",
            command=self.copy_evrim_details
        )
        self.tree_evrim.bind("<Button-3>", self.show_context_menu)

    def show_context_menu(self, event):
        """Sağ tıklandığında seçili satırları algılar ve dinamik menüyü açar."""
        item_id = self.tree_evrim.identify_row(event.y)
        selected_ids = list(self.tree_evrim.selection())

        if item_id and item_id not in selected_ids:
            self.tree_evrim.selection_set(item_id)
            selected_ids = [item_id]
            self.on_evrim_row_select(None)

        if not selected_ids:
            return

        sel_count = len(selected_ids)
        menu = tk.Menu(self.tree_evrim, tearoff=0)

        # 1. Operasyonel Aksiyonlar
        menu.add_command(label="📧 Vergi Talebi Maili Aç", command=self.create_tax_mail_from_evrim)
        menu.add_command(label="🤖 Otonom RPA İstasyonuna Gönder", command=self.send_selected_to_rpa)
        menu.add_command(label="🌐 QNB Fatura Eşleştiriciye Gönder", command=self.send_to_qnb_matcher)
        menu.add_command(label="🚛 GLP Referans Sorgusuna Gönder", command=self.send_to_glp_query)
        menu.add_command(label="💸 Para Talebine Gönder", command=self.send_selected_evrim_to_talep)
        menu.add_command(label="🏛️ GTB Takip Havuzuna Ekle", command=self.send_to_gtb_tracking)
        menu.add_separator()

        # 2. Seçili Sütun Kopyalama
        lbl_dosya = f"📋 Seçili ({sel_count}) Dosya No'ları Kopyala (Alt Alta)" if sel_count > 1 else "📋 Dosya Numarasını Kopyala"
        lbl_ref = f"📋 Seçili ({sel_count}) Referans No'ları Kopyala (Alt Alta)" if sel_count > 1 else "📋 Referans Numarasını Kopyala"
        lbl_beyan = f"📋 Seçili ({sel_count}) Beyanname No'ları Kopyala (Alt Alta)" if sel_count > 1 else "📋 Beyanname Numarasını Kopyala"

        menu.add_command(label=lbl_dosya, command=lambda: self.copy_evrim_column("dosya_no", "Dosya No", only_selected=True))
        menu.add_command(label=lbl_ref, command=lambda: self.copy_evrim_column("hankook_ref", "Referans No", only_selected=True))
        menu.add_command(label=lbl_beyan, command=lambda: self.copy_evrim_column("beyanname_no", "Beyanname No", only_selected=True))
        if sel_count > 1:
            menu.add_command(label=f"📋 Seçili ({sel_count}) Satırı Excel Tablosu Olarak Kopyala", command=self.copy_selected_evrim_rows_as_table)

        menu.add_separator()

        # 3. Tablodaki TÜM Sütunları Kopyalama
        all_count = len(self.filtered_evrim_records or [])
        all_menu = tk.Menu(menu, tearoff=0)
        all_menu.add_command(label=f"📑 Tablodaki TÜM Dosya No'ları Kopyala ({all_count} Satır)", command=lambda: self.copy_evrim_column("dosya_no", "Dosya No", only_selected=False))
        all_menu.add_command(label=f"📑 Tablodaki TÜM Referans No'ları Kopyala ({all_count} Satır)", command=lambda: self.copy_evrim_column("hankook_ref", "Referans No", only_selected=False))
        all_menu.add_command(label=f"📑 Tablodaki TÜM Beyanname No'ları Kopyala ({all_count} Satır)", command=lambda: self.copy_evrim_column("beyanname_no", "Beyanname No", only_selected=False))
        menu.add_cascade(label=f"📑 Tablodaki TÜM Sütunu Kopyala ({all_count} Satır)...", menu=all_menu)

        menu.add_separator()
        menu.add_command(label="📊 Excel'e Aktar (.xlsx)", command=self.export_evrim_to_excel)

        menu.tk_popup(event.x_root, event.y_root)

    def select_all_evrim_rows(self, event=None):
        """Ctrl+A ile tüm satırları seçer."""
        all_items = self.tree_evrim.get_children()
        if all_items:
            self.tree_evrim.selection_set(all_items)
            self.on_evrim_row_select(None)
        return "break"

    def focus_location_filter(self):
        """Metrik kardına tıklandığında Liman filtre combobox'ına odaklar."""
        if hasattr(self, "cmb_ev_port"):
            self.cmb_ev_port.focus_set()
            try:
                self.cmb_ev_port.event_generate("<Button-1>")
            except Exception:
                pass

    def reset_filters(self):
        """Tüm arama ve combobox filtrelerini varsayılana sıfırlar."""
        if hasattr(self, "txt_ev_search"):
            self.txt_ev_search.delete(0, tk.END)
        if hasattr(self, "cmb_ev_port"):
            self.cmb_ev_port.set("Tüm Limanlar")
        if hasattr(self, "cmb_ev_antrepo"):
            self.cmb_ev_antrepo.set("Tüm Antrepolar")
        if hasattr(self, "cmb_ev_statu"):
            self.cmb_ev_statu.set("Tüm Statüler")
        self.set_evrim_filter("ALL")

    def on_activate(self):
        """Sekme aktifleştiğinde filtreleri yeniden uygular."""
        self.apply_evrim_filters()

    def set_evrim_filter(self, filter_type):
        self.current_evrim_filter = filter_type
        btns = [
            ("ALL", self.btn_ev_f_all),
            ("TESCILLI", self.btn_ev_f_beyan),
            ("PENDING", self.btn_ev_f_pend),
            ("CLOSED", self.btn_ev_f_close)
        ]
        for ftype, btn in btns:
            if ftype == filter_type:
                btn.config(bg="#0f172a", fg="#ffffff")
            else:
                btn.config(bg="#f1f5f9", fg="#475569")
        self.apply_evrim_filters()

    def is_evrim_record_closed(self, r):
        st = (r.get("statu") or "").upper()
        asama = (r.get("asama") or "").upper()
        ilerleme = (r.get("ilerleme") or "").strip()
        teslim = (r.get("teslim_tarihi") or "").strip()

        if teslim and len(teslim) > 3:
            return True
        if "MUHASEBE EVRAĞI DAĞITIM" in st or "DAGITIMA GONDERILDI" in st or "DAĞITIMA GÖNDERİLDİ" in st or "MUHASEBE EVRAK" in st:
            return True
        if ilerleme in ["%100", "100%", "100"]:
            return True
        if "KAPAN" in st or "TESLİM" in st or "TAMAMLANDI" in st:
            return True
        if "İŞLEM BİTTİ" in asama or "KAPANDI" in asama:
            return True
        return False

    def is_antrepo_entry(self, val_str):
        """Değerin antrepo / özel depo olup olmadığını kontrol eder."""
        u = (val_str or "").upper()
        return bool(
            "ANTREPO" in u or "DEPO" in u or
            re.search(r"\bG[0-9]{5,}\b", u) or
            re.search(r"\bA[0-9]{4,}\b", u)
        )

    def apply_evrim_filters(self):
        records = self.evrim_records
        query = self.txt_ev_search.get().strip().upper() if hasattr(self, "txt_ev_search") else ""
        selected_port = self.cmb_ev_port.get() if hasattr(self, "cmb_ev_port") else "Tüm Limanlar"
        selected_antrepo = self.cmb_ev_antrepo.get() if hasattr(self, "cmb_ev_antrepo") else "Tüm Antrepolar"
        selected_statu = self.cmb_ev_statu.get() if hasattr(self, "cmb_ev_statu") else "Tüm Statüler"

        ports = set()
        antrepos = set()
        statuses = set()

        for r in records:
            p = (r.get("liman") or "").strip()
            g = (r.get("gumruk") or "").strip()
            st = (r.get("statu") or "").strip()

            if p:
                if self.is_antrepo_entry(p):
                    antrepos.add(p)
                else:
                    ports.add(p)
            if g and self.is_antrepo_entry(g):
                antrepos.add(g)
            if st:
                statuses.add(st)

        # Liman Combobox Güncelleme
        if hasattr(self, "cmb_ev_port"):
            current_vals = list(self.cmb_ev_port["values"])
            new_vals = ["Tüm Limanlar"] + sorted(list(ports))
            if current_vals != new_vals:
                self.cmb_ev_port["values"] = new_vals
                if not selected_port or selected_port not in new_vals:
                    self.cmb_ev_port.set("Tüm Limanlar")
                    selected_port = "Tüm Limanlar"

        # Antrepo Combobox Güncelleme
        if hasattr(self, "cmb_ev_antrepo"):
            current_antrepo_vals = list(self.cmb_ev_antrepo["values"])
            new_antrepo_vals = ["Tüm Antrepolar"] + sorted(list(antrepos))
            if current_antrepo_vals != new_antrepo_vals:
                self.cmb_ev_antrepo["values"] = new_antrepo_vals
                if not selected_antrepo or selected_antrepo not in new_antrepo_vals:
                    self.cmb_ev_antrepo.set("Tüm Antrepolar")
                    selected_antrepo = "Tüm Antrepolar"

        # Statü Combobox Güncelleme
        if hasattr(self, "cmb_ev_statu"):
            current_statu_vals = list(self.cmb_ev_statu["values"])
            new_statu_vals = ["Tüm Statüler"] + sorted(list(statuses))
            if current_statu_vals != new_statu_vals:
                self.cmb_ev_statu["values"] = new_statu_vals
                if not selected_statu or selected_statu not in new_statu_vals:
                    self.cmb_ev_statu.set("Tüm Statüler")
                    selected_statu = "Tüm Statüler"

        filtered = []
        for r in records:
            is_closed = self.is_evrim_record_closed(r)
            has_beyan = bool(r.get("beyanname_no") and len(str(r.get("beyanname_no", "")).strip()) > 3)

            # Aşama Filtresi (Hızlı Butonlar)
            if self.current_evrim_filter == "TESCILLI" and (is_closed or not has_beyan):
                continue
            elif self.current_evrim_filter == "PENDING" and (is_closed or has_beyan):
                continue
            elif self.current_evrim_filter == "CLOSED" and not is_closed:
                continue

            # Liman Filtresi
            if selected_port and selected_port != "Tüm Limanlar":
                if (r.get("liman") or "").strip() != selected_port:
                    continue

            # Antrepo Filtresi
            if selected_antrepo and selected_antrepo != "Tüm Antrepolar":
                r_lim = (r.get("liman") or "").strip()
                r_gum = (r.get("gumruk") or "").strip()
                if r_lim != selected_antrepo and r_gum != selected_antrepo:
                    continue

            # Statü Filtresi
            if selected_statu and selected_statu != "Tüm Statüler":
                if (r.get("statu") or "").strip() != selected_statu:
                    continue

            # Metin Arama Sorgusu
            if query:
                combined_text = " ".join([
                    str(r.get("dosya_no", "")),
                    str(r.get("hankook_ref", "")),
                    str(r.get("beyanname_no", "")),
                    str(r.get("statu", "")),
                    str(r.get("asama", "")),
                    str(r.get("liman", "")),
                    str(r.get("gumruk", "")),
                    str(r.get("rejim", "")),
                    str(r.get("yurt_disi_firma", ""))
                ]).upper()
                if query not in combined_text:
                    continue

            filtered.append(r)

        self.filtered_evrim_records = filtered
        if hasattr(self.app, "filtered_evrim_records"):
            self.app.filtered_evrim_records = filtered

        if hasattr(self, "lbl_ev_result_count"):
            self.lbl_ev_result_count.config(text=f"{len(filtered)} Kayıt Listeleniyor")

        self.render_evrim_tree()
        self.update_evrim_metrics()

    def update_evrim_metrics(self):
        records = self.evrim_records
        total = len(records)
        closed_cnt = sum(1 for r in records if self.is_evrim_record_closed(r))
        beyan_cnt = sum(1 for r in records if not self.is_evrim_record_closed(r) and (r.get("beyanname_no") and len(str(r.get("beyanname_no", "")).strip()) > 3))
        pending_cnt = total - closed_cnt - beyan_cnt
        if pending_cnt < 0:
            pending_cnt = 0

        ports = set((r.get("liman") or "").strip() for r in records if (r.get("liman") or "").strip())

        if hasattr(self, "card_ev_total") and "val_lbl" in self.card_ev_total:
            self.card_ev_total["val_lbl"].config(text=str(total))
        if hasattr(self, "card_ev_beyan") and "val_lbl" in self.card_ev_beyan:
            self.card_ev_beyan["val_lbl"].config(text=str(beyan_cnt))
        if hasattr(self, "card_ev_pending") and "val_lbl" in self.card_ev_pending:
            self.card_ev_pending["val_lbl"].config(text=str(pending_cnt))
        if hasattr(self, "card_ev_closed") and "val_lbl" in self.card_ev_closed:
            self.card_ev_closed["val_lbl"].config(text=str(closed_cnt))
        if hasattr(self, "card_ev_ports") and "val_lbl" in self.card_ev_ports:
            self.card_ev_ports["val_lbl"].config(text=str(len(ports)))

    def render_evrim_tree(self):
        if not hasattr(self, "tree_evrim"):
            return
        self.tree_evrim.delete(*self.tree_evrim.get_children())

        for idx, r in enumerate(self.filtered_evrim_records):
            is_closed = self.is_evrim_record_closed(r)
            has_beyan = bool(r.get("beyanname_no") and len(str(r.get("beyanname_no", "")).strip()) > 3)

            tag = "pending"
            if is_closed:
                tag = "closed"
            elif has_beyan:
                tag = "tescilli"

            self.tree_evrim.insert(
                "",
                tk.END,
                iid=str(idx),
                values=(
                    r.get("dosya_no", "—") or "—",
                    r.get("hankook_ref", "—") or "—",
                    r.get("beyanname_no", "—") or "—",
                    r.get("beyanname_tarihi", "—") or "—",
                    r.get("statu", "—") or "—",
                    f"{r.get('asama', '')} ({r.get('ilerleme', '')})" if r.get('ilerleme') else (r.get('asama', '—') or "—"),
                    r.get("liman", "—") or "—",
                    r.get("gumruk", "—") or "—",
                    r.get("teslim_tarihi", "—") or "—",
                    r.get("yurt_disi_firma", "—") or "—"
                ),
                tags=(tag,)
            )

    def refresh_table(self):
        self.apply_evrim_filters()

    def on_evrim_row_select(self, event=None):
        sel = self.tree_evrim.selection()
        if not sel:
            return

        if hasattr(self, "btn_ev_mail"):
            if len(sel) > 1:
                self.btn_ev_mail.config(text=f"📧 Toplu Vergi Maili Hazırla ({len(sel)} Dosya)", bg="#334155")
            else:
                self.btn_ev_mail.config(text="📧 Tek Tıkla Vergi Maili Hazırla", bg="#2563eb")

        first_idx = int(sel[0])
        if 0 <= first_idx < len(self.filtered_evrim_records):
            r = self.filtered_evrim_records[first_idx]
            self.selected_evrim_rec = r
            if hasattr(self.app, "selected_evrim_rec"):
                self.app.selected_evrim_rec = r

            is_closed = self.is_evrim_record_closed(r)
            has_beyan = bool(r.get("beyanname_no") and len(str(r.get("beyanname_no", "")).strip()) > 3)

            if len(sel) > 1:
                self.lbl_ev_status_badge.config(
                    text=f"⚡ {len(sel)} ADET DOSYA SEÇİLDİ (Toplu İşlem Hazır)",
                    bg="#dbeafe", fg="#1e40af"
                )
            elif is_closed:
                self.lbl_ev_status_badge.config(text="🔵 TESLİMATI YAPILMIŞ / KAPALI DOSYA", bg="#eff6ff", fg="#1e40af")
            elif has_beyan:
                self.lbl_ev_status_badge.config(text=f"🟢 TESCİL ALINDI: {r.get('beyanname_no')}", bg="#f0fdf4", fg="#166534")
            else:
                self.lbl_ev_status_badge.config(text=f"⏳ {r.get('statu', 'İşlem Aşamasında')}", bg="#fffbeb", fg="#92400e")

            # 11 Künye Alanını Güncelle
            self.ev_detail_labels["ref"].config(text=r.get("hankook_ref", "—") or "—")
            self.ev_detail_labels["dosya"].config(text=r.get("dosya_no", "—") or "—")
            self.ev_detail_labels["beyan_no"].config(text=r.get("beyanname_no", "—") or "—")
            self.ev_detail_labels["beyan_tarih"].config(text=r.get("beyanname_tarihi", "—") or "—")
            self.ev_detail_labels["gumruk"].config(text=r.get("gumruk", "—") or "—")
            self.ev_detail_labels["liman"].config(text=r.get("liman", "—") or "—")
            self.ev_detail_labels["rejim"].config(text=r.get("rejim", "—") or "—")
            self.ev_detail_labels["statu"].config(text=r.get("statu", "—") or "—")
            asama_txt = f"{r.get('asama', '—') or ''} {r.get('ilerleme', '')}".strip()
            self.ev_detail_labels["asama"].config(text=asama_txt or "—")
            self.ev_detail_labels["teslim"].config(text=r.get("teslim_tarihi", "—") or "—")
            self.ev_detail_labels["yurt_disi"].config(text=r.get("yurt_disi_firma", "—") or "—")

    def create_tax_mail_from_evrim(self):
        sel = self.tree_evrim.selection()
        if not sel:
            messagebox.showwarning("Uyarı", "Lütfen önce tablodan bir veya birden fazla dosya seçin!")
            return

        formatted_records = []
        for item_id in sel:
            try:
                idx = int(item_id)
                if 0 <= idx < len(self.filtered_evrim_records):
                    r = self.filtered_evrim_records[idx]
                    ref = (r.get("hankook_ref") or "").strip() or (r.get("dosya_no") or "").strip()
                    formatted_records.append({
                        "ref": ref,
                        "dosya": (r.get("dosya_no") or "").strip(),
                        "beyan_no": (r.get("beyanname_no") or "").strip()
                    })
            except Exception:
                continue

        if not formatted_records:
            messagebox.showwarning("Uyarı", "Seçili dosya bilgisi alınamadı.")
            return

        count = len(formatted_records)
        self.set_header_subtitle(f"⏳ {count} adet dosya için Vergi Talebi taslağı Outlook'ta hazırlanıyor...")

        def _w():
            if not self.engine or not hasattr(self.engine, "create_tax_mail_draft"):
                self.show_error("Hata", "Motor veya vergi maili oluşturucu bulunamadı.")
                return
            ok, msg = self.engine.create_tax_mail_draft(formatted_records)
            def _d():
                if ok:
                    self.set_header_subtitle(f"✅ {count} adet dosya için Vergi Talebi taslağı Outlook'ta açıldı!", fg="#16a34a")
                else:
                    self.set_header_subtitle("⚠️ Vergi maili taslağı oluşturulamadı.")
                    messagebox.showerror("Hata", msg)
            self.root.after(0, _d)
        threading.Thread(target=_w, daemon=True).start()
    def send_selected_to_rpa(self):
        """Seçili Evrim dosyalarını Otonom RPA İstasyonuna gönderir ve sekmeyi açar."""
        sel = list(self.tree_evrim.selection())
        if not sel:
            messagebox.showwarning("Uyarı", "Lütfen önce tablodan dosya seçin.")
            return

        items = []
        for item_id in sel:
            try:
                idx = int(item_id)
                if 0 <= idx < len(self.filtered_evrim_records):
                    r = self.filtered_evrim_records[idx]
                    items.append({
                        "dosya_no": r.get("dosya_no", ""),
                        "ref": r.get("hankook_ref", ""),
                        "beyanname_no": r.get("beyanname_no", ""),
                        "gemi": ""
                    })
            except Exception:
                continue

        rpa_view = getattr(self.app, "view_rpa", None) or (self.app.views.get("RPA_STATION") if hasattr(self.app, "views") else None)
        if rpa_view and hasattr(rpa_view, "add_external_files_to_queue"):
            rpa_view.add_external_files_to_queue(items)
            self.switch_tab("RPA_STATION")
            messagebox.showinfo("RPA İstasyonu", f"✅ {len(items)} adet dosya Otonom RPA İstasyonu kuyruğuna aktarıldı!")
        else:
            self.switch_tab("RPA_STATION")

    def copy_evrim_column(self, col_key, col_title, only_selected=True):
        """Evrim tablosunda belirtilen sütundaki verileri (seçili veya tüm tablodaki) alt alta panoya kopyalar."""
        if only_selected:
            sel = list(self.tree_evrim.selection())
            scope_name = f"Seçili {len(sel)} adet"
            records = []
            for item_id in sel:
                try:
                    idx = int(item_id)
                    if 0 <= idx < len(self.filtered_evrim_records):
                        records.append(self.filtered_evrim_records[idx])
                except Exception:
                    continue
        else:
            records = list(self.filtered_evrim_records or [])
            scope_name = f"Tablodaki tüm ({len(records)})"

        if not records:
            self.show_warning("Uyarı", "Kopyalanacak satır bulunamadı.")
            return

        vals = []
        for r in records:
            v = str(r.get(col_key) or "").strip()
            if v and v != "—":
                vals.append(v)

        if not vals:
            self.show_warning("Uyarı", f"Bu sütunda geçerli {col_title} bulunamadı.")
            return

        text = "\n".join(vals)
        self.copy_to_clipboard(text)
        messagebox.showinfo(
            "Kopyalandı",
            f"📋 {scope_name} {col_title} alt alta panoya kopyalandı!\n\n(Excel'e veya arama alanına doğrudan yapıştırabilirsiniz)"
        )

    def copy_selected_evrim_rows_as_table(self):
        """Seçili Evrim satırlarını Excel uyumlu tablo formatında panoya kopyalar."""
        sel = list(self.tree_evrim.selection())
        if not sel:
            self.show_warning("Uyarı", "Lütfen önce tablodan kopyalanacak satırları seçin.")
            return

        lines = ["Evrim Dosya No\tHankook Ref\tBeyanname No\tTescil Tarihi\tStatü Açıklama\tAşama / %\tLiman\tGümrük\tTeslim Tarihi\tYurt Dışı Firma"]
        for item_id in sel:
            try:
                idx = int(item_id)
                if 0 <= idx < len(self.filtered_evrim_records):
                    r = self.filtered_evrim_records[idx]
                    lines.append(
                        f"{r.get('dosya_no','')}\t{r.get('hankook_ref','')}\t{r.get('beyanname_no','')}\t{r.get('beyanname_tarihi','')}\t{r.get('statu','')}\t{r.get('asama','')}\t{r.get('liman','')}\t{r.get('gumruk','')}\t{r.get('teslim_tarihi','')}\t{r.get('yurt_disi_firma','')}"
                    )
            except Exception:
                continue

        text = "\n".join(lines)
        self.copy_to_clipboard(text)
        messagebox.showinfo(
            "Kopyalandı",
            f"📋 Seçili {len(sel)} satır Excel formatında panoya kopyalandı!\n\n(Excel'e doğrudan Ctrl+V ile yapıştırabilirsiniz)"
        )

    def copy_evrim_details(self):
        sel = self.tree_evrim.selection()
        if not sel:
            return

        lines = []
        for item_id in sel:
            try:
                idx = int(item_id)
                if 0 <= idx < len(self.filtered_evrim_records):
                    r = self.filtered_evrim_records[idx]
                    lines.append(f"Dosya No: {r.get('dosya_no')} | Ref: {r.get('hankook_ref')} | Beyan No: {r.get('beyanname_no')} | Liman: {r.get('liman')} | Statü: {r.get('statu')}")
            except Exception:
                continue

        if lines:
            text = "\n".join(lines)
            self.copy_to_clipboard(text)
            messagebox.showinfo("Kopyalandı", f"📋 {len(lines)} adet dosya künyesi panoya kopyalandı!")

    def export_evrim_to_excel(self):
        if not self.filtered_evrim_records:
            messagebox.showwarning("Uyarı", "Dışa aktarılacak kayıt bulunamadı.")
            return

        default_name = f"Evrim_Is_Takip_Raporu_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        save_path = filedialog.asksaveasfilename(
            title="Evrim Raporunu Excel Olarak Kaydet",
            initialfile=default_name,
            filetypes=[("Excel Dosyaları", "*.xlsx")]
        )
        if not save_path:
            return

        if not openpyxl:
            messagebox.showerror("Hata", "openpyxl kütüphanesi bulunamadı.")
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Evrim İş Takip"

            headers = [
                "Evrim Dosya No", "Hankook Referans No", "Beyanname No", "Beyanname Tarihi",
                "Statü Açıklama", "Aşama", "İlerleme %", "Bulunduğu Yer / Liman",
                "Gümrük Müdürlüğü", "Teslim Tarihi", "Rejim", "Yurt Dışı Firma", "Not"
            ]
            ws.append(headers)

            for cell in ws[1]:
                cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
                cell.fill = openpyxl.styles.PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

            for r in self.filtered_evrim_records:
                ws.append([
                    r.get("dosya_no", ""),
                    r.get("hankook_ref", ""),
                    r.get("beyanname_no", ""),
                    r.get("beyanname_tarihi", ""),
                    r.get("statu", ""),
                    r.get("asama", ""),
                    r.get("ilerleme", ""),
                    r.get("liman", ""),
                    r.get("gumruk", ""),
                    r.get("teslim_tarihi", ""),
                    r.get("rejim", ""),
                    r.get("yurt_disi_firma", ""),
                    r.get("not", "")
                ])

            # Otomatik sütun genişliği
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(save_path)
            if messagebox.askyesno("Başarılı", f"✅ {len(self.filtered_evrim_records)} adet dosya başarıyla Excel'e aktarıldı!\n\nDosyayı şimdi açmak ister misiniz?"):
                os.startfile(save_path)
        except Exception as e:
            messagebox.showerror("Hata", f"Excel oluşturulurken hata oluştu: {e}")

    def send_to_qnb_matcher(self):
        if not hasattr(self, "selected_evrim_rec") or not self.selected_evrim_rec:
            messagebox.showwarning("Uyarı", "Lütfen önce tablodan bir dosya seçin!")
            return

        r = self.selected_evrim_rec
        target_code = (r.get("beyanname_no") or "").strip() or (r.get("hankook_ref") or "").strip() or (r.get("dosya_no") or "").strip()

        if not target_code:
            messagebox.showwarning("Uyarı", "Seçili dosyaya ait Beyanname veya Referans numarası bulunamadı.")
            return

        self.copy_to_clipboard(target_code)
        self.switch_tab("QNB_PORTAL")

        qnb_proc = getattr(self.app, "qnb_process", None)

        def _deliver_ipc():
            for _ in range(5):
                if qnb_proc and qnb_proc.poll() is None and qnb_proc.stdin:
                    try:
                        msg = f"SET_BEYAN:{target_code}\n".encode("utf-8")
                        qnb_proc.stdin.write(msg)
                        qnb_proc.stdin.flush()
                    except Exception:
                        pass
                time.sleep(0.8)

        threading.Thread(target=_deliver_ipc, daemon=True).start()

        self.set_header_subtitle(
            f"✅ {target_code} numarası QNB Fatura Eşleştirici'ye aktarıldı!",
            fg="#16a34a"
        )

    def send_to_glp_query(self):
        if not hasattr(self, "selected_evrim_rec") or not self.selected_evrim_rec:
            messagebox.showwarning("Uyarı", "Lütfen önce tablodan bir dosya seçin!")
            return

        r = self.selected_evrim_rec
        ref = (r.get("hankook_ref") or "").strip() or (r.get("dosya_no") or "").strip()
        if not ref:
            return

        self.copy_to_clipboard(ref)
        self.switch_tab("GLP_PORTAL")
        self.set_header_subtitle(
            f"✅ Referans No ({ref}) panoya kopyalandı ve GLP Portalı açıldı!",
            fg="#16a34a"
        )

    def send_selected_evrim_to_talep(self):
        if not hasattr(self, "selected_evrim_rec") or not self.selected_evrim_rec:
            messagebox.showwarning("Uyarı", "Lütfen önce tablodan bir dosya seçin!")
            return

        if hasattr(self.app, "send_selected_evrim_to_talep"):
            self.app.send_selected_evrim_to_talep()
            return

        r = self.selected_evrim_rec
        dosya_no = (r.get("dosya_no") or "").strip()
        ref_no = (r.get("hankook_ref") or "").strip()
        liman = (r.get("liman") or "").strip()

        target_ref = dosya_no or ref_no
        self.copy_to_clipboard(target_ref)

        payload = {
            "dosyaNo": dosya_no,
            "ref": ref_no,
            "liman": liman,
            "aciklama": "ARDİYE FATURASI"
        }
        if hasattr(self.app, "send_to_talep_form"):
            self.app.send_to_talep_form(payload)
        else:
            self.switch_tab("TALEP")

    def send_to_gtb_tracking(self):
        """Seçili dosyaların tescilli beyanname numaralarını GTB Otonom Takip Havuzuna ekler."""
        sel = self.tree_evrim.selection()
        if not sel:
            messagebox.showwarning("Uyarı", "Lütfen önce tablodan bir veya birden fazla dosya seçin!")
            return

        tracker = getattr(self.app, "gtb_tracker", None)
        if not tracker:
            messagebox.showerror("Hata", "GTB Takip servisi devrede değil.")
            return

        added_count = 0
        already_tracked = 0
        no_beyan = 0

        for item_id in sel:
            try:
                idx = int(item_id)
                if 0 <= idx < len(self.filtered_evrim_records):
                    r = self.filtered_evrim_records[idx]
                    beyan = (r.get("beyanname_no") or "").strip()
                    clean_b = re.sub(r"[^A-Z0-9]", "", beyan.upper())
                    if not clean_b or len(clean_b) < 10:
                        no_beyan += 1
                        continue

                    dosya = (r.get("dosya_no") or "—").strip()
                    ref = (r.get("hankook_ref") or "—").strip()
                    firma = (r.get("yurt_disi_firma") or "WH-HANKOOK").strip()

                    res = tracker.add_to_tracking(clean_b, dosya_no=dosya, firma=firma, ref=ref)
                    if res.get("success"):
                        if res.get("created", True):
                            added_count += 1
                        else:
                            already_tracked += 1
            except Exception:
                continue

        # GTB Tablosunu yenile
        gtb_view = getattr(self.app, "view_gtb_takip", None) or (self.app.views.get("GTB_TAKIP") if hasattr(self.app, "views") else None)
        if gtb_view and hasattr(gtb_view, "refresh_gtb_takip_table"):
            gtb_view.refresh_gtb_takip_table()

        if added_count > 0:
            msg = f"✅ {added_count} adet beyanname GTB Takip Havuzuna eklendi!"
            if already_tracked > 0:
                msg += f"\n({already_tracked} beyanname zaten havuzda mevcuttu)"
            if no_beyan > 0:
                msg += f"\n({no_beyan} dosyada geçerli beyanname no bulunamadı)"
            self.set_header_subtitle(msg, fg="#16a34a")
            if messagebox.askyesno("GTB Takip", f"{msg}\n\nGTB Takip ekranına gitmek ister misiniz?"):
                self.switch_tab("GTB_TAKIP")
        elif already_tracked > 0:
            messagebox.showinfo("Bilgi", f"Seçili beyannameler ({already_tracked} adet) zaten GTB Takip Havuzunda mevcut.")
        else:
            messagebox.showwarning("Uyarı", "Seçili kayıtlarda geçerli (en az 10 haneli) beyanname numarası bulunamadı.")

    def trigger_dogruer_sync(self):
        self.switch_tab("DOGRUER_PORTAL")

    def sort_evrim_tree(self, col):
        """Tablo sütun başlıklarına tıklandığında akıllı sıralama yapar."""
        if not hasattr(self, "_sort_reverse_map"):
            self._sort_reverse_map = {}
        reverse = self._sort_reverse_map.get(col, False)

        def get_sort_key(rec):
            val = rec.get(col, "")
            if val is None:
                return ""
            val = str(val).strip()
            # Tarih formatı: GG.AA.YYYY
            if re.match(r"^\d{2}\.\d{2}\.\d{4}$", val):
                try:
                    parts = val.split(".")
                    return f"{parts[2]}{parts[1]}{parts[0]}"
                except Exception:
                    pass
            # Sayısal / Yüzde kontrolü
            try:
                clean_num = val.replace("%", "").replace(",", ".").strip()
                return float(clean_num)
            except ValueError:
                pass
            return val.lower()

        self.filtered_evrim_records.sort(key=get_sort_key, reverse=reverse)
        self._sort_reverse_map[col] = not reverse
        self.render_evrim_tree()
