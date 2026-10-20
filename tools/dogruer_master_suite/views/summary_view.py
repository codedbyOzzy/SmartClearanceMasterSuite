# -*- coding: utf-8 -*-
import os
import csv
import datetime
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

from .base_view import BaseView

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    openpyxl = None


class SummaryView(BaseView):
    """
    Günlük İş ve Operasyon Özeti (Executive Dashboard)
    - KPI Metrik Kartları
    - Operasyonel Aşama Dağılımı ve İlerleme Barları
    - Liman Bazlı Açık Dosya Dağılımı
    - Mail Sağlık Durumu ve Güvenlik İstihbaratı
    - Günlük Rapor Excel Dışa Aktarımı
    """

    def __init__(self, parent, app, *args, **kwargs):
        super().__init__(parent, app, *args, **kwargs)
        self.summary_bars = {}
        self._build_ui()

    def on_activate(self):
        self.update_summary_view()

    def create_metric_widget(self, parent, title, val, subtitle, accent_color, command):
        card = tk.Frame(parent, bg="#ffffff", padx=12, pady=8, relief=tk.FLAT, bd=0, cursor="hand2")
        card.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=3)

        top_strip = tk.Frame(card, bg=accent_color, height=3)
        top_strip.pack(fill=tk.X, pady=(0, 4))

        lbl_t = tk.Label(card, text=title, font=("Segoe UI", 8, "bold"), fg="#64748b", bg="#ffffff", cursor="hand2")
        lbl_t.pack(anchor="w")

        lbl_v = tk.Label(card, text=val, font=("Segoe UI", 15, "bold"), fg=accent_color, bg="#ffffff", cursor="hand2")
        lbl_v.pack(anchor="w")

        lbl_sub = tk.Label(card, text=subtitle, font=("Segoe UI", 8), fg="#94a3b8", bg="#ffffff", cursor="hand2")
        lbl_sub.pack(anchor="w")

        for w in (card, top_strip, lbl_t, lbl_v, lbl_sub):
            w.bind("<Button-1>", lambda e, cmd=command: cmd())

        return {"card": card, "val_lbl": lbl_v, "title": title}

    def _build_ui(self):
        # 1. ÜST KPI METRİK KARTLARI
        metric_bar = tk.Frame(self, bg="#f8fafc")
        metric_bar.pack(fill=tk.X, pady=(0, 12))

        self.card_s_total = self.create_metric_widget(
            metric_bar, "TOPLAM TAKİP EDİLEN", "0", "Excel Master Kayıt", "#0f172a",
            lambda: self.switch_tab("COCKPIT")
        )
        self.card_s_open = self.create_metric_widget(
            metric_bar, "AÇIK / CANLI DOSYALAR", "0", "Aktif Operasyon", "#1e3a8a",
            lambda: self.switch_tab("COCKPIT")
        )
        self.card_s_mail_wait = self.create_metric_widget(
            metric_bar, "VERGİ MAİLİ BEKLEYEN", "0", "Acil Gönderim", "#dc2626",
            lambda: self.switch_tab("COCKPIT")
        )
        self.card_s_closed = self.create_metric_widget(
            metric_bar, "KAPANDI / ARŞİV", "0", "Tamamlanan Dosyalar", "#16a34a",
            lambda: self.switch_tab("COCKPIT")
        )

        # 2. ORTA BÖLÜM: 2 SÜTUNLU DASHBOARD
        mid_frame = tk.Frame(self, bg="#f8fafc")
        mid_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

        # SOL KART: SÜREÇ & AŞAMA DAĞILIMI
        left_card = tk.Frame(mid_frame, bg="#ffffff", padx=16, pady=16, relief=tk.FLAT, bd=0)
        left_card.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 6))

        lbl_lc_title = tk.Label(
            left_card,
            text="⚡ Operasyonel Aşama Dağılımı & İlerleme",
            font=("Segoe UI", 11, "bold"),
            fg="#0f172a",
            bg="#ffffff"
        )
        lbl_lc_title.pack(anchor="w", pady=(0, 14))

        stages = [
            ("tescil", "🔵 Tescil & Muayene Aşamasındakiler", "#2563eb"),
            ("ordino", "🟠 Ordino & Tareks Aşamasındakiler", "#ea580c"),
            ("hazirlik", "🟡 Yeni Evrak / Hazırlık Sürecindekiler", "#eab308"),
            ("mail_sent", "🟢 Vergi Maili Gönderilmiş Dosyalar", "#16a34a"),
            ("mail_wait", "🔴 Acil Vergi Maili Bekleyenler", "#dc2626")
        ]

        for key, title, color in stages:
            b_box = tk.Frame(left_card, bg="#ffffff", pady=4)
            b_box.pack(fill=tk.X)

            h_row = tk.Frame(b_box, bg="#ffffff")
            h_row.pack(fill=tk.X, pady=(0, 3))

            lbl_t = tk.Label(h_row, text=title, font=("Segoe UI", 8, "bold"), fg="#334155", bg="#ffffff")
            lbl_t.pack(side=tk.LEFT)

            lbl_c = tk.Label(h_row, text="0 Dosya (%0)", font=("Segoe UI", 8, "bold"), fg=color, bg="#ffffff")
            lbl_c.pack(side=tk.RIGHT)

            p_track = tk.Frame(b_box, bg="#e2e8f0", height=8)
            p_track.pack(fill=tk.X)
            p_track.pack_propagate(False)

            p_fill = tk.Frame(p_track, bg=color, width=0, height=8)
            p_fill.pack(side=tk.LEFT, fill=tk.Y)

            self.summary_bars[key] = {
                "lbl": lbl_c,
                "track": p_track,
                "fill": p_fill,
                "color": color
            }

        # SAĞ KART: LİMAN DAĞILIMI & GÜVENLİK
        right_card = tk.Frame(mid_frame, bg="#ffffff", padx=16, pady=16, relief=tk.FLAT, bd=0)
        right_card.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(6, 0))

        lbl_rc_title = tk.Label(
            right_card,
            text="⚓ Liman Bazlı Açık Dosya Dağılımı & Güvenlik",
            font=("Segoe UI", 11, "bold"),
            fg="#0f172a",
            bg="#ffffff"
        )
        lbl_rc_title.pack(anchor="w", pady=(0, 10))

        self.lbl_s_mail_health = tk.Label(
            right_card,
            text="🔍 Vergi Mailleri Kontrol Ediliyor...",
            font=("Segoe UI", 9, "bold"),
            bg="#f0fdf4",
            fg="#166534",
            padx=12,
            pady=8,
            anchor="w"
        )
        self.lbl_s_mail_health.pack(fill=tk.X, pady=(0, 10))

        lbl_p_hdr = tk.Label(
            right_card,
            text="Liman / Antrepo Noktaları:",
            font=("Segoe UI", 8, "bold"),
            fg="#64748b",
            bg="#ffffff"
        )
        lbl_p_hdr.pack(anchor="w", pady=(0, 4))

        self.port_container = tk.Frame(right_card, bg="#ffffff")
        self.port_container.pack(fill=tk.BOTH, expand=True)

        btn_row = tk.Frame(right_card, bg="#ffffff", pady=8)
        btn_row.pack(fill=tk.X)

        btn_s_export = tk.Button(
            btn_row,
            text="📊 Günlük Rapor Excel İndir",
            font=("Segoe UI", 8, "bold"),
            bg="#2563eb",
            fg="#ffffff",
            activebackground="#1d4ed8",
            activeforeground="#ffffff",
            relief=tk.FLAT,
            padx=10,
            pady=5,
            cursor="hand2",
            command=self.export_summary_to_excel
        )
        btn_s_export.pack(side=tk.LEFT, padx=(0, 6))

        btn_s_excel = tk.Button(
            btn_row,
            text="📊 Excel Dosyasını Aç",
            font=("Segoe UI", 8, "bold"),
            bg="#107c41",
            fg="#ffffff",
            activebackground="#0b5e31",
            activeforeground="#ffffff",
            relief=tk.FLAT,
            padx=10,
            pady=5,
            cursor="hand2",
            command=self.open_excel_file
        )
        btn_s_excel.pack(side=tk.LEFT, padx=(0, 6))

        btn_s_shield = tk.Button(
            btn_row,
            text="🛡️ Güvenlik Kalkanı",
            font=("Segoe UI", 8, "bold"),
            bg="#1e3a8a",
            fg="#ffffff",
            activebackground="#1e40af",
            activeforeground="#ffffff",
            relief=tk.FLAT,
            padx=10,
            pady=5,
            cursor="hand2",
            command=lambda: self.switch_tab("SHIELD")
        )
        btn_s_shield.pack(side=tk.LEFT)

        # 3. ALT İSTİHBARAT BANDI
        self.bottom_insight = tk.Label(
            self,
            text="💡 Sistem Tavsiyesi: Vergi maili bekleyen acil dosyaları Operasyon Kokpiti veya Güvenlik Kalkanı üzerinden tek tıkla gönderebilirsiniz.",
            font=("Segoe UI", 9),
            bg="#f1f5f9",
            fg="#334155",
            padx=14,
            pady=8,
            relief=tk.SOLID,
            bd=1,
            anchor="w"
        )
        self.bottom_insight.pack(fill=tk.X)

    def switch_tab(self, tab_code):
        if hasattr(self.app, "switch_tab"):
            self.app.switch_tab(tab_code)

    def open_excel_file(self):
        p = self.config.get("excel_path", "")
        if p and os.path.exists(p):
            try:
                os.startfile(p)
            except Exception as e:
                messagebox.showerror("Hata", f"Excel dosyası açılamadı: {e}")
        else:
            messagebox.showwarning("Uyarı", f"Excel dosyası bulunamadı:\n{p}")

    def update_summary_view(self):
        records = getattr(self.engine, "records", [])
        total = len(records)
        open_c = sum(1 for r in records if not r.get("kapanis"))
        tescil_c = sum(1 for r in records if not r.get("kapanis") and (r.get("stage") in ["3_TESCILDE", "3_MUAYENE"]))
        mail_waiting = sum(1 for r in records if not r.get("kapanis") and r.get("is_tax_pending") and r.get("mail_badge") == "red")
        mail_sent = sum(1 for r in records if not r.get("kapanis") and r.get("mail_badge") == "green")
        ordino_c = sum(1 for r in records if not r.get("kapanis") and r.get("stage") == "2_ORDINO_TAREKS")
        hazirlik_c = open_c - tescil_c - ordino_c
        if hazirlik_c < 0:
            hazirlik_c = 0
        closed_c = sum(1 for r in records if r.get("kapanis"))

        # Metrik Kartları
        if hasattr(self, "card_s_total"):
            self.card_s_total["val_lbl"].config(text=str(total))
            self.card_s_open["val_lbl"].config(text=str(open_c))
            self.card_s_mail_wait["val_lbl"].config(text=str(mail_waiting))
            self.card_s_closed["val_lbl"].config(text=str(closed_c))

        # İlerleme Çubukları
        def set_bar(key, count, base):
            if key not in self.summary_bars:
                return
            item = self.summary_bars[key]
            pct = int((count / base) * 100) if base > 0 else 0
            item["lbl"].config(text=f"{count} Dosya (%{pct})")
            track_w = item["track"].winfo_width()
            if track_w <= 1:
                track_w = 280
            fill_w = int((pct / 100) * track_w) if base > 0 else 0
            item["fill"].config(width=fill_w)

        if open_c > 0:
            set_bar("tescil", tescil_c, open_c)
            set_bar("ordino", ordino_c, open_c)
            set_bar("hazirlik", hazirlik_c, open_c)
            set_bar("mail_sent", mail_sent, open_c)
            set_bar("mail_wait", mail_waiting, open_c)
        else:
            set_bar("tescil", 0, 0)
            set_bar("ordino", 0, 0)
            set_bar("hazirlik", 0, 0)
            set_bar("mail_sent", 0, 0)
            set_bar("mail_wait", 0, 0)

        # Mail Sağlık Durumu
        if hasattr(self, "lbl_s_mail_health"):
            if mail_waiting > 0:
                self.lbl_s_mail_health.config(
                    text=f"⚠️ DİKKAT: {mail_waiting} adet açık dosya için Vergi Talebi maili henüz gönderilmedi!",
                    bg="#fef2f2",
                    fg="#991b1b"
                )
            else:
                self.lbl_s_mail_health.config(
                    text="✅ HARİKA: Açık ve tescilli tüm dosyaların Vergi Talebi mailleri eksiksiz gönderildi.",
                    bg="#f0fdf4",
                    fg="#166534"
                )

        # Liman Dağılımı Listesi
        if hasattr(self, "port_container"):
            for w in self.port_container.winfo_children():
                w.destroy()

            port_counts = {}
            for r in records:
                if not r.get("kapanis"):
                    p = (r.get("liman") or "Belirtilmemiş").strip() or "Belirtilmemiş"
                    port_counts[p] = port_counts.get(p, 0) + 1

            sorted_ports = sorted(port_counts.items(), key=lambda x: x[1], reverse=True)
            if not sorted_ports:
                lbl_empty = tk.Label(self.port_container, text="Açık dosya kaydı bulunamadı.", font=("Segoe UI", 8), fg="#94a3b8", bg="#ffffff")
                lbl_empty.pack(anchor="w", pady=4)
            else:
                for p_name, p_cnt in sorted_ports[:6]:
                    p_row = tk.Frame(self.port_container, bg="#f8fafc", padx=8, pady=4, relief=tk.FLAT, bd=0)
                    p_row.pack(fill=tk.X, pady=2)

                    lbl_pn = tk.Label(p_row, text=f"📍 {p_name}", font=("Segoe UI", 8, "bold"), fg="#0f172a", bg="#f8fafc")
                    lbl_pn.pack(side=tk.LEFT)

                    lbl_pc = tk.Label(p_row, text=f"{p_cnt} Dosya", font=("Segoe UI", 8, "bold"), fg="#2563eb", bg="#f8fafc")
                    lbl_pc.pack(side=tk.RIGHT)

    def export_summary_to_excel(self):
        """Günlük operasyon özetini ve aktif dosyaları şık bir Excel raporu olarak dışa aktarır."""
        records = getattr(self.engine, "records", [])
        total = len(records)
        open_c = sum(1 for r in records if not r.get("kapanis"))
        tescil_c = sum(1 for r in records if not r.get("kapanis") and (r.get("stage") in ["3_TESCILDE", "3_MUAYENE"]))
        mail_waiting = sum(1 for r in records if not r.get("kapanis") and r.get("is_tax_pending") and r.get("mail_badge") == "red")
        mail_sent = sum(1 for r in records if not r.get("kapanis") and r.get("mail_badge") == "green")
        ordino_c = sum(1 for r in records if not r.get("kapanis") and r.get("stage") == "2_ORDINO_TAREKS")
        hazirlik_c = max(0, open_c - tescil_c - ordino_c)
        closed_c = sum(1 for r in records if r.get("kapanis"))

        now_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        def_filename = f"Gunluk_Operasyon_Raporu_{now_str}.xlsx"

        file_path = filedialog.asksaveasfilename(
            title="Günlük Raporu Excel Olarak Kaydet",
            initialfile=def_filename,
            defaultextension=".xlsx",
            filetypes=[("Excel Dosyası", "*.xlsx"), ("Tüm Dosyalar", "*.*")]
        )

        if not file_path:
            return

        try:
            if openpyxl:
                wb = openpyxl.Workbook()
                
                # --- SAYFA 1: OPERASYONEL ÖZET ---
                ws1 = wb.active
                ws1.title = "Operasyonel Özet"
                ws1.views.sheetView[0].showGridLines = True

                # Başlık
                ws1.merge_cells("A1:E1")
                ws1["A1"] = "DOĞRUER GÜMRÜK MÜŞAVİRLİĞİ - GÜNLÜK İŞ & OPERASYON RAPORU"
                ws1["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
                ws1["A1"].fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
                ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
                ws1.row_dimensions[1].height = 36

                # Tarih ve Kullanıcı
                u_name = getattr(self.app, "current_user", {}).get("full_name", "Kullanıcı")
                ws1["A2"] = f"Rapor Tarihi: {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')} | Hazırlayan: {u_name}"
                ws1["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
                ws1.row_dimensions[2].height = 20

                # KPI Tablosu
                thin_border = Border(
                    left=Side(style='thin', color='E2E8F0'),
                    right=Side(style='thin', color='E2E8F0'),
                    top=Side(style='thin', color='E2E8F0'),
                    bottom=Side(style='thin', color='E2E8F0')
                )

                ws1["A4"] = "METRİK"
                ws1["B4"] = "DEĞER"
                ws1["C4"] = "AÇIKLAMA"
                for col in ["A4", "B4", "C4"]:
                    ws1[col].font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
                    ws1[col].fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
                    ws1[col].alignment = Alignment(horizontal="center" if col == "B4" else "left", vertical="center")

                kpi_data = [
                    ("Toplam Takip Edilen Dosyalar", total, "Excel Master Kayıtları Toplamı"),
                    ("Açık / Canlı Dosyalar", open_c, "Aktif Gümrükleme Sürecindekiler"),
                    ("Vergi Maili Bekleyen Acil Dosyalar", mail_waiting, "Tescil Edilmiş / Mail Bekleyenler"),
                    ("Tamamlanan / Arşiv Dosyalar", closed_c, "Kapanış Yapılmış Dosyalar")
                ]

                r_idx = 5
                for m_title, m_val, m_desc in kpi_data:
                    ws1[f"A{r_idx}"] = m_title
                    ws1[f"B{r_idx}"] = m_val
                    ws1[f"C{r_idx}"] = m_desc
                    ws1[f"A{r_idx}"].font = Font(name="Segoe UI", size=9, bold=True)
                    ws1[f"B{r_idx}"].font = Font(name="Segoe UI", size=10, bold=True, color="1E3A8A")
                    ws1[f"B{r_idx}"].alignment = Alignment(horizontal="center")
                    ws1[f"C{r_idx}"].font = Font(name="Segoe UI", size=9, color="64748B")
                    for c_letter in ["A", "B", "C"]:
                        ws1[f"{c_letter}{r_idx}"].border = thin_border
                    r_idx += 1

                # Süreç Dağılımı
                r_idx += 2
                ws1[f"A{r_idx}"] = "SÜREÇ & AŞAMA DAĞILIMI"
                ws1[f"B{r_idx}"] = "DOSYA SAYISI"
                ws1[f"C{r_idx}"] = "ORAN"
                for col in [f"A{r_idx}", f"B{r_idx}", f"C{r_idx}"]:
                    ws1[col].font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
                    ws1[col].fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
                    ws1[col].alignment = Alignment(horizontal="center" if col != f"A{r_idx}" else "left", vertical="center")

                stage_rows = [
                    ("Tescil & Muayene Aşamasındakiler", tescil_c, f"%{int((tescil_c/open_c)*100) if open_c>0 else 0}"),
                    ("Ordino & Tareks Aşamasındakiler", ordino_c, f"%{int((ordino_c/open_c)*100) if open_c>0 else 0}"),
                    ("Yeni Evrak / Hazırlık Sürecindekiler", hazirlik_c, f"%{int((hazirlik_c/open_c)*100) if open_c>0 else 0}"),
                    ("Vergi Maili Gönderilmiş Dosyalar", mail_sent, f"%{int((mail_sent/open_c)*100) if open_c>0 else 0}"),
                    ("Acil Vergi Maili Bekleyenler", mail_waiting, f"%{int((mail_waiting/open_c)*100) if open_c>0 else 0}")
                ]

                r_idx += 1
                for s_title, s_cnt, s_pct in stage_rows:
                    ws1[f"A{r_idx}"] = s_title
                    ws1[f"B{r_idx}"] = s_cnt
                    ws1[f"C{r_idx}"] = s_pct
                    ws1[f"A{r_idx}"].font = Font(name="Segoe UI", size=9)
                    ws1[f"B{r_idx}"].font = Font(name="Segoe UI", size=9, bold=True)
                    ws1[f"B{r_idx}"].alignment = Alignment(horizontal="center")
                    ws1[f"C{r_idx}"].font = Font(name="Segoe UI", size=9)
                    ws1[f"C{r_idx}"].alignment = Alignment(horizontal="center")
                    for c_letter in ["A", "B", "C"]:
                        ws1[f"{c_letter}{r_idx}"].border = thin_border
                    r_idx += 1

                # Liman Dağılımı
                r_idx += 2
                ws1[f"A{r_idx}"] = "LİMAN / ANTREPO NOKTASI"
                ws1[f"B{r_idx}"] = "AÇIK DOSYA"
                ws1[f"C{r_idx}"] = "ORAN"
                for col in [f"A{r_idx}", f"B{r_idx}", f"C{r_idx}"]:
                    ws1[col].font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
                    ws1[col].fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
                    ws1[col].alignment = Alignment(horizontal="center" if col != f"A{r_idx}" else "left", vertical="center")

                port_counts = {}
                for r in records:
                    if not r.get("kapanis"):
                        p = (r.get("liman") or "Belirtilmemiş").strip() or "Belirtilmemiş"
                        port_counts[p] = port_counts.get(p, 0) + 1

                r_idx += 1
                for p_name, p_cnt in sorted(port_counts.items(), key=lambda x: x[1], reverse=True)[:10]:
                    ws1[f"A{r_idx}"] = p_name
                    ws1[f"B{r_idx}"] = p_cnt
                    ws1[f"C{r_idx}"] = f"%{int((p_cnt/open_c)*100) if open_c>0 else 0}"
                    ws1[f"A{r_idx}"].font = Font(name="Segoe UI", size=9)
                    ws1[f"B{r_idx}"].font = Font(name="Segoe UI", size=9, bold=True)
                    ws1[f"B{r_idx}"].alignment = Alignment(horizontal="center")
                    ws1[f"C{r_idx}"].font = Font(name="Segoe UI", size=9)
                    ws1[f"C{r_idx}"].alignment = Alignment(horizontal="center")
                    for c_letter in ["A", "B", "C"]:
                        ws1[f"{c_letter}{r_idx}"].border = thin_border
                    r_idx += 1

                ws1.column_dimensions["A"].width = 38
                ws1.column_dimensions["B"].width = 16
                ws1.column_dimensions["C"].width = 32

                # --- SAYFA 2: AKTİF DOSYALAR DETAY LİSTESİ ---
                ws2 = wb.create_sheet(title="Aktif Dosyalar")
                ws2.views.sheetView[0].showGridLines = True

                headers = [
                    "Dosya No", "Müşteri / Ref", "Konşimento (B/L)", "Beyanname No",
                    "Gümrük / Liman", "Süreç Aşaması", "Vergi Mail Durumu", "Kapanış"
                ]

                for col_num, h_text in enumerate(headers, 1):
                    cell = ws2.cell(row=1, column=col_num, value=h_text)
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                ws2.row_dimensions[1].height = 28

                row_cursor = 2
                for r in records:
                    if not r.get("kapanis"):
                        stage_txt = r.get("stage", "")
                        mail_stat = "Gönderildi" if r.get("mail_badge") == "green" else ("Bekliyor (Acil)" if r.get("is_tax_pending") else "Beklemiyor")

                        vals = [
                            str(r.get("dosya") or "—"),
                            str(r.get("ref") or "—"),
                            str(r.get("bl") or "—"),
                            str(r.get("beyan") or "—"),
                            str(r.get("liman") or "—"),
                            stage_txt,
                            mail_stat,
                            "Açık"
                        ]

                        for col_num, val in enumerate(vals, 1):
                            cell = ws2.cell(row=row_cursor, column=col_num, value=val)
                            cell.font = Font(name="Segoe UI", size=9)
                            cell.alignment = Alignment(horizontal="center" if col_num in [1, 3, 4, 7, 8] else "left", vertical="center")
                            cell.border = thin_border

                            if col_num == 7 and "Acil" in val:
                                cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                                cell.font = Font(name="Segoe UI", size=9, bold=True, color="991B1B")
                            elif col_num == 7 and "Gönderildi" in val:
                                cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                                cell.font = Font(name="Segoe UI", size=9, color="166534")

                        row_cursor += 1

                for col_letter, width in zip(["A", "B", "C", "D", "E", "F", "G", "H"], [16, 22, 20, 20, 22, 24, 18, 12]):
                    ws2.column_dimensions[col_letter].width = width

                wb.save(file_path)
            else:
                # Fallback to CSV
                with open(file_path, "w", encoding="utf-8-sig", newline="") as f:
                    writer = csv.writer(f)
                    writer.writerow(["DOĞRUER GÜMRÜK MÜŞAVİRLİĞİ - GÜNLÜK İŞ & OPERASYON RAPORU"])
                    writer.writerow(["Rapor Tarihi", datetime.datetime.now().strftime('%d.%m.%Y %H:%M')])
                    writer.writerow([])
                    writer.writerow(["Metrik", "Değer"])
                    writer.writerow(["Toplam Takip Edilen", total])
                    writer.writerow(["Açık Dosyalar", open_c])
                    writer.writerow(["Vergi Maili Bekleyen", mail_waiting])
                    writer.writerow(["Kapandı / Arşiv", closed_c])
                    writer.writerow([])
                    writer.writerow(["Dosya No", "Müşteri / Ref", "Konşimento (B/L)", "Beyanname No", "Gümrük / Liman", "Aşama", "Mail Durumu"])
                    for r in records:
                        if not r.get("kapanis"):
                            writer.writerow([
                                r.get("dosya", ""), r.get("ref", ""), r.get("bl", ""),
                                r.get("beyan", ""), r.get("liman", ""), r.get("stage", ""),
                                "Gönderildi" if r.get("mail_badge") == "green" else ("Bekliyor" if r.get("is_tax_pending") else "Normal")
                            ])

            if messagebox.askyesno("Başarılı", f"🎉 Günlük operasyon raporu başarıyla kaydedildi:\n\n{file_path}\n\nDosyayı şimdi açmak ister misiniz?"):
                try:
                    os.startfile(file_path)
                except Exception:
                    pass
        except Exception as e:
            messagebox.showerror("Hata", f"Excel raporu oluşturulamadı:\n{e}")

