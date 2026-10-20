# -*- coding: utf-8 -*-
"""
Doğruer Master Suite - Otonom RPA İstasyonu Görünümü
3 Robot Entegre (Beyanname, Excel Aktarım, PL4) ve Bağımsız Kalibrasyon Merkezi
"""

import os
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, scrolledtext
import threading
import time
from typing import Dict, Any, List

from tools.dogruer_master_suite.views.base_view import BaseView
from tools.dogruer_master_suite.engine.rpa_engine import (
    RPAEngine, DEFAULT_SABIT_DEGERLER,
    BEYANNAME_ALANLAR, AKTARIM_ALANLAR, PL4_ALANLAR,
    KOOR_BEYANNAME_P1, KOOR_BEYANNAME_P2, KOOR_AKTARIM, KOOR_PL4
)
from tools.dogruer_master_suite.engine.rpa_field_values import queue_item_from_record

try:
    import openpyxl
except ImportError:
    openpyxl = None


class RPAStationView(BaseView):
    """OTONOM RPA İSTASYONU - 3 ROBOTLU KONTROL & KALİBRASYON MERKEZİ"""

    def __init__(self, parent, app):
        super().__init__(parent, app)
        self.rpa = RPAEngine()
        self.floating_panel = None
        self.active_tab_name = "BEYANNAME"  # 'BEYANNAME', 'AKTARIM', 'PL4'

        self._setup_rpa_callbacks()
        self.setup_ui()

    def _setup_rpa_callbacks(self):
        self.rpa.on_status_change = self._on_rpa_status_change
        self.rpa.on_queue_update = self._on_rpa_queue_update
        self.rpa.on_step_progress = self._on_rpa_step_progress
        self.rpa.on_approval_needed = self._on_rpa_approval_needed
        self.rpa.on_finish = self._on_rpa_finish
        self.rpa.on_log = self._on_rpa_log

    def setup_ui(self):
        self.configure(bg="#f8fafc")

        # ---------------------------------------------------------------------
        # 1. ÜST SEGMENT SEÇİCİ (3 ROBOT ARASINDA GEÇİŞ)
        # ---------------------------------------------------------------------
        top_bar = tk.Frame(self, bg="#0f172a", padx=12, pady=8)
        top_bar.pack(fill=tk.X)

        lbl_title = tk.Label(
            top_bar, text="🤖 OTONOM RPA İSTASYONU", font=("Segoe UI", 11, "bold"),
            fg="#38bdf8", bg="#0f172a"
        )
        lbl_title.pack(side=tk.LEFT, padx=(0, 16))

        # Sekme Geçiş Butonları
        self.tab_buttons: Dict[str, tk.Button] = {}
        tabs = [
            ("BEYANNAME", "📑 1. Beyanname Doldurma Robotu"),
            ("AKTARIM", "📊 2. Evrim Excel Aktarım Robotu (F8)"),
            ("PL4", "🏷️ 3. Toplu PL4 Ön İzin Robotu")
        ]

        for code, text in tabs:
            btn = tk.Button(
                top_bar, text=text, font=("Segoe UI", 9, "bold"),
                bg="#1e293b", fg="#94a3b8", activebackground="#334155", activeforeground="#ffffff",
                relief=tk.FLAT, padx=12, pady=4, cursor="hand2",
                command=lambda c=code: self.switch_robot_tab(c)
            )
            btn.pack(side=tk.LEFT, padx=3)
            self.tab_buttons[code] = btn

        # Acil Durum Butonu (Başlık Barında Sabit)
        btn_panic = tk.Button(
            top_bar, text="🚨 ACİL DURDUR (ESC / F12)", font=("Segoe UI", 9, "bold"),
            bg="#dc2626", fg="#ffffff", activebackground="#b91c1c", activeforeground="#ffffff",
            relief=tk.FLAT, padx=12, pady=4, cursor="hand2", command=self.trigger_emergency_stop
        )
        btn_panic.pack(side=tk.RIGHT)

        btn_probe = tk.Button(
            top_bar, text="🔎 Evrim'i Algıla", font=("Segoe UI", 9, "bold"),
            bg="#0f766e", fg="#ffffff", activebackground="#115e59", activeforeground="#ffffff",
            relief=tk.FLAT, padx=10, pady=4, cursor="hand2", command=self.probe_evrim_window
        )
        btn_probe.pack(side=tk.RIGHT, padx=(0, 8))

        # ---------------------------------------------------------------------
        # 2. ROBOT PANELLERİ CONTAINER'I
        # ---------------------------------------------------------------------
        self.container = tk.Frame(self, bg="#f8fafc")
        self.container.pack(fill=tk.BOTH, expand=True, padx=14, pady=8)

        # Robot 1: Beyanname
        self.panel_beyan = tk.Frame(self.container, bg="#f8fafc")
        self._build_beyanname_panel(self.panel_beyan)

        # Robot 2: Aktarım
        self.panel_aktarim = tk.Frame(self.container, bg="#f8fafc")
        self._build_aktarim_panel(self.panel_aktarim)

        # Robot 3: PL4
        self.panel_pl4 = tk.Frame(self.container, bg="#f8fafc")
        self._build_pl4_panel(self.panel_pl4)

        # ---------------------------------------------------------------------
        # 3. ALT CANLI İŞLEM GÜNLÜĞÜ (KONSOL)
        # ---------------------------------------------------------------------
        log_frame = tk.LabelFrame(self, text=" 📜 Canlı Robot İşlem & Kalibrasyon Günlüğü ", font=("Segoe UI", 8, "bold"), bg="#f8fafc", padx=8, pady=4)
        log_frame.pack(fill=tk.X, padx=14, pady=(0, 8))

        self.txt_console = scrolledtext.ScrolledText(
            log_frame, height=4, bg="#0f172a", fg="#4ade80", font=("Consolas", 9), relief=tk.FLAT
        )
        self.txt_console.pack(fill=tk.BOTH, expand=True)

        # Varsayılan sekme
        self.switch_robot_tab("BEYANNAME")

    def switch_robot_tab(self, code: str):
        self.active_tab_name = code
        for c, btn in self.tab_buttons.items():
            if c == code:
                btn.config(bg="#7c3aed", fg="#ffffff")
            else:
                btn.config(bg="#1e293b", fg="#94a3b8")

        self.panel_beyan.pack_forget()
        self.panel_aktarim.pack_forget()
        self.panel_pl4.pack_forget()

        if code == "BEYANNAME":
            self.panel_beyan.pack(fill=tk.BOTH, expand=True)
            self._log(f"📑 Beyanname Doldurma Robotu aktif. Profil: {self.rpa.active_beyan_profile}")
        elif code == "AKTARIM":
            self.panel_aktarim.pack(fill=tk.BOTH, expand=True)
            self._log("📊 Evrim Excel Aktarım Robotu aktif. F8 tetiklemeli mod.")
            self.refresh_excel_folder()
        else:
            self.panel_pl4.pack(fill=tk.BOTH, expand=True)
            self._log("🏷️ Toplu PL4 Ön İzin Ekleme Robotu aktif.")

    # =========================================================================
    # 1. BEYANNAME DOLDURMA ROBOTU PANELİ
    # =========================================================================
    def _build_beyanname_panel(self, parent):
        # Üst Metrik Kartları
        cards = tk.Frame(parent, bg="#f8fafc", pady=4)
        cards.pack(fill=tk.X)

        self.c1_v, self.c1_s = self._make_card(cards, "ROBOT DURUMU", "HAZIR (BEKLEMEDE)", "Kuyruk işlem bekliyor", "#0284c7")
        self.c2_v, self.c2_s = self._make_card(cards, "ÇALIŞMA MODU", "👁️ YARI OTONOM", "Space ile dosya onayı", "#7c3aed")
        self.c3_v, self.c3_s = self._make_card(cards, "KUYRUK DURUMU", "0 / 0 Dosya", "Kuyruk boş", "#059669")
        self.c4_v, self.c4_s = self._make_card(cards, "KALİBRASYON DURUMU", f"{len(self.rpa.koor_beyan)} / 29 Alan", "Koordinatlar hazır", "#d97706")

        # Kontrol Barı
        bar = tk.Frame(parent, bg="#ffffff", padx=10, pady=6, relief=tk.GROOVE, bd=1)
        bar.pack(fill=tk.X, pady=(4, 6))

        tk.Label(bar, text="Mod:", font=("Segoe UI", 9, "bold"), bg="#ffffff").pack(side=tk.LEFT, padx=(0, 4))
        self.var_beyan_mode = tk.StringVar(value="SEMI_AUTONOMOUS")
        tk.Radiobutton(bar, text="👁️ Yarı Otonom (Space)", variable=self.var_beyan_mode, value="SEMI_AUTONOMOUS", bg="#ffffff", command=self._on_beyan_mode_change).pack(side=tk.LEFT, padx=4)
        tk.Radiobutton(bar, text="🚀 Tam Otonom", variable=self.var_beyan_mode, value="FULL_AUTONOMOUS", font=("Segoe UI", 9, "bold"), fg="#7c3aed", bg="#ffffff", command=self._on_beyan_mode_change).pack(side=tk.LEFT, padx=4)

        self.var_beyan_test = tk.BooleanVar(value=False)
        tk.Checkbutton(bar, text="Güvenli Test Modu (Yazmaz / Tamam Basmaz)", variable=self.var_beyan_test, fg="#2563eb", bg="#ffffff", font=("Segoe UI", 8, "bold")).pack(side=tk.LEFT, padx=8)

        tk.Label(bar, text="Profil:", font=("Segoe UI", 9, "bold"), bg="#ffffff").pack(side=tk.LEFT, padx=(12, 4))
        self.cb_beyan_prof = ttk.Combobox(bar, values=["Profil 1 (Ana Ekran)", "Profil 2 (İkinci Monitör)"], state="readonly", width=18)
        self.cb_beyan_prof.set("Profil 1 (Ana Ekran)")
        self.cb_beyan_prof.pack(side=tk.LEFT)
        self.cb_beyan_prof.bind("<<ComboboxSelected>>", lambda e: self.rpa.set_beyan_profile(self.cb_beyan_prof.get()))

        # Aksiyonlar
        self.btn_beyan_start = tk.Button(bar, text="🚀 Robotu Başlat", font=("Segoe UI", 9, "bold"), bg="#7c3aed", fg="#ffffff", padx=12, pady=4, relief=tk.FLAT, cursor="hand2", command=self.start_beyanname_process)
        self.btn_beyan_start.pack(side=tk.RIGHT, padx=3)
        self.btn_beyan_pause = tk.Button(bar, text="⏸️ Duraklat", font=("Segoe UI", 9, "bold"), bg="#d97706", fg="#ffffff", padx=10, pady=4, relief=tk.FLAT, cursor="hand2", command=self.rpa.pause, state=tk.DISABLED)
        self.btn_beyan_pause.pack(side=tk.RIGHT, padx=3)
        self.btn_beyan_resume = tk.Button(bar, text="▶️ Devam Et", font=("Segoe UI", 9, "bold"), bg="#059669", fg="#ffffff", padx=10, pady=4, relief=tk.FLAT, cursor="hand2", command=self.rpa.resume, state=tk.DISABLED)
        self.btn_beyan_resume.pack(side=tk.RIGHT, padx=3)

        # Kalibrasyon Butonları
        btn_calib = tk.Button(bar, text="🎯 Kalibrasyon Başlat (F9)", font=("Segoe UI", 8, "bold"), bg="#ea580c", fg="#ffffff", padx=8, pady=4, relief=tk.FLAT, cursor="hand2", command=lambda: self.start_calibration_for("BEYANNAME"))
        btn_calib.pack(side=tk.RIGHT, padx=6)
        btn_test = tk.Button(bar, text="👀 Koordinatları Test Et", font=("Segoe UI", 8), bg="#0284c7", fg="#ffffff", padx=8, pady=4, relief=tk.FLAT, cursor="hand2", command=lambda: self.test_coordinates_for("BEYANNAME"))
        btn_test.pack(side=tk.RIGHT, padx=3)

        # Gövde PanedWindow
        paned = tk.PanedWindow(parent, orient=tk.HORIZONTAL, bg="#e2e8f0", sashwidth=4, bd=0)
        paned.pack(fill=tk.BOTH, expand=True)

        # Sol: Kuyruk Tablosu
        left = tk.Frame(paned, bg="#ffffff", padx=8, pady=6)
        paned.add(left, minsize=520)

        q_tool = tk.Frame(left, bg="#ffffff")
        q_tool.pack(fill=tk.X, pady=(0, 4))
        tk.Label(q_tool, text="📋 İŞLENECEK DOSYA KUYRUĞU", font=("Segoe UI", 9, "bold"), bg="#ffffff", fg="#1e293b").pack(side=tk.LEFT)
        tk.Button(q_tool, text="➕ Dosya Ekle", font=("Segoe UI", 8), bg="#f1f5f9", padx=6, relief=tk.GROOVE, cursor="hand2", command=self.show_add_file_dialog).pack(side=tk.RIGHT, padx=2)
        tk.Button(q_tool, text="📂 Excel'den Al", font=("Segoe UI", 8), bg="#f1f5f9", padx=6, relief=tk.GROOVE, cursor="hand2", command=self.load_from_excel_file).pack(side=tk.RIGHT, padx=2)
        tk.Button(q_tool, text="🧹 Temizle", font=("Segoe UI", 8), bg="#fee2e2", fg="#b91c1c", padx=6, relief=tk.GROOVE, cursor="hand2", command=self.clear_queue).pack(side=tk.RIGHT, padx=2)

        tree_f = tk.Frame(left, bg="#ffffff")
        tree_f.pack(fill=tk.BOTH, expand=True)
        cols = ("sira", "dosya_no", "ref", "beyan_no", "gemi", "liman", "status", "message")
        self.tree_q = ttk.Treeview(tree_f, columns=cols, show="headings", height=10)
        self.tree_q.heading("sira", text="#")
        self.tree_q.heading("dosya_no", text="Dosya No")
        self.tree_q.heading("ref", text="Referans No")
        self.tree_q.heading("beyan_no", text="Beyanname No")
        self.tree_q.heading("gemi", text="Taşıt / Gemi")
        self.tree_q.heading("liman", text="Liman")
        self.tree_q.heading("status", text="Durum")
        self.tree_q.heading("message", text="Sonuç / Not")

        self.tree_q.column("sira", width=30, anchor="center")
        self.tree_q.column("dosya_no", width=90, anchor="center")
        self.tree_q.column("ref", width=100, anchor="center")
        self.tree_q.column("beyan_no", width=120, anchor="center")
        self.tree_q.column("gemi", width=110)
        self.tree_q.column("liman", width=70, anchor="center")
        self.tree_q.column("status", width=90, anchor="center")
        self.tree_q.column("message", width=140)

        sb = ttk.Scrollbar(tree_f, orient=tk.VERTICAL, command=self.tree_q.yview)
        self.tree_q.configure(yscrollcommand=sb.set)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree_q.pack(fill=tk.BOTH, expand=True)

        self.tree_q.tag_configure("TAMAMLANDI", background="#dcfce7", foreground="#15803d")
        self.tree_q.tag_configure("YAZILIYOR", background="#fef3c7", foreground="#b45309")
        self.tree_q.tag_configure("ONAY_BEKLIYOR", background="#ede9fe", foreground="#6d28d9")
        self.tree_q.tag_configure("DURDURULDU", background="#fee2e2", foreground="#b91c1c")
        self.tree_q.tag_configure("HATA", background="#fee2e2", foreground="#b91c1c")

        # Sağ: Canlı Parametre Düzenleyici
        right = tk.Frame(paned, bg="#ffffff", padx=10, pady=6, width=320)
        paned.add(right, minsize=300)

        tk.Label(right, text="✏️ CANLI PARAMETRE DÜZENLEYİCİ", font=("Segoe UI", 9, "bold"), bg="#ffffff", fg="#7c3aed").pack(anchor="w")
        tk.Label(right, text="Kalan dosyaların Evrim değerlerini anında değiştirin:", font=("Segoe UI", 8), bg="#ffffff", fg="#64748b").pack(anchor="w", pady=(0, 4))

        self.param_entries: Dict[str, tk.Entry] = {}
        fields = [
            ("tasit_kimligi", "Taşıt / Gemi:", "MAERSK EL BOSQUE"),
            ("teslim_sekli", "Teslim Şekli (20):", "CIP"),
            ("teslim_sehri", "Teslim Şehri:", "KOCAELİ"),
            ("liman_kutu", "Liman Kodu:", "DP"),
            ("diger_tutar", "Diğer Gider Tutarı:", "3000"),
            ("diger_aciklama_metin", "Diğer Açıklama:", "Y.ICI"),
            ("liman_gider_tutar", "Liman Gider Tutarı:", "5000"),
            ("liman_gider_aciklama_metin", "Liman Açıklama:", "."),
            ("ek_bilgi_metin", "44 No'lu Kutu Notu:", "Tareks ve Kons. icin odenen damga vergisi tutari yurt ici diger gidere beyan edildi."),
        ]

        for k, lbl_t, def_v in fields:
            row = tk.Frame(right, bg="#ffffff")
            row.pack(fill=tk.X, pady=1)
            tk.Label(row, text=lbl_t, font=("Segoe UI", 8, "bold"), bg="#ffffff", fg="#334155", width=16, anchor="w").pack(side=tk.LEFT)
            ent = tk.Entry(row, font=("Segoe UI", 9), relief=tk.SOLID, bd=1)
            ent.insert(0, def_v)
            ent.pack(side=tk.LEFT, fill=tk.X, expand=True)
            self.param_entries[k] = ent

        btn_app = tk.Button(
            right, text="💾 Parametreleri Kalan Dosyalara Uygula", font=("Segoe UI", 8, "bold"),
            bg="#0284c7", fg="#ffffff", pady=5, relief=tk.FLAT, cursor="hand2", command=self.apply_params_to_remaining
        )
        btn_app.pack(fill=tk.X, pady=(8, 2))

    # =========================================================================
    # 2. EVRİM EXCEL AKTARIM ROBOTU PANELİ
    # =========================================================================
    def _build_aktarim_panel(self, parent):
        cards = tk.Frame(parent, bg="#f8fafc", pady=4)
        cards.pack(fill=tk.X)

        self.ak1_v, self.ak1_s = self._make_card(cards, "AKTARIM STATÜSÜ", "BEKLEMEDE", "Excel dosyası hazır", "#0284c7")
        self.ak2_v, self.ak2_s = self._make_card(cards, "KLASÖRDEKİ EXCELLER", "0 Dosya", "Taranıyor...", "#059669")
        self.ak3_v, self.ak3_s = self._make_card(cards, "F8 MODU", "F8 İLE ONAYLI", "Her dosyada F8 bekler", "#7c3aed")
        self.ak4_v, self.ak4_s = self._make_card(cards, "KALİBRASYON", f"{len(self.rpa.koor_aktarim)} / 14 Alan", "aktarim_koordinatlar.json", "#ea580c")

        # Klasör ve Kontrol Barı
        bar = tk.Frame(parent, bg="#ffffff", padx=10, pady=8, relief=tk.GROOVE, bd=1)
        bar.pack(fill=tk.X, pady=(4, 6))

        tk.Label(bar, text="📁 Excel Klasörü:", font=("Segoe UI", 9, "bold"), bg="#ffffff").pack(side=tk.LEFT, padx=(0, 4))
        self.ent_excel_klasor = tk.Entry(bar, font=("Segoe UI", 9), relief=tk.SOLID, bd=1, width=45)
        default_k = r"C:\Users\ozan.gemici\OneDrive - DOGRUER GUMRUK MUSAVIRLIGI A.S\Belgeler\OTONOMAKTARIM"
        self.ent_excel_klasor.insert(0, default_k)
        self.ent_excel_klasor.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 6))

        tk.Button(bar, text="Seç...", bg="#f1f5f9", padx=8, relief=tk.GROOVE, cursor="hand2", command=self.select_excel_folder).pack(side=tk.LEFT, padx=2)
        tk.Button(bar, text="🔄 Yenile", bg="#f1f5f9", padx=8, relief=tk.GROOVE, cursor="hand2", command=self.refresh_excel_folder).pack(side=tk.LEFT, padx=(2, 10))

        self.var_aktarim_f8 = tk.BooleanVar(value=True)
        tk.Checkbutton(bar, text="Her Dosyada [F8] Bekle", variable=self.var_aktarim_f8, bg="#ffffff", font=("Segoe UI", 8, "bold")).pack(side=tk.LEFT, padx=4)

        self.var_aktarim_test = tk.BooleanVar(value=False)
        tk.Checkbutton(bar, text="Güvenli Test Modu (Aktar Basmaz)", variable=self.var_aktarim_test, fg="#2563eb", bg="#ffffff", font=("Segoe UI", 8, "bold")).pack(side=tk.LEFT, padx=4)

        # Kalibrasyon Butonları
        tk.Button(bar, text="🎯 Aktarım Kalibrasyonu (F9)", font=("Segoe UI", 8, "bold"), bg="#ea580c", fg="#ffffff", padx=8, pady=4, relief=tk.FLAT, cursor="hand2", command=lambda: self.start_calibration_for("AKTARIM")).pack(side=tk.RIGHT, padx=4)
        tk.Button(bar, text="👀 Test Et", font=("Segoe UI", 8), bg="#0284c7", fg="#ffffff", padx=8, pady=4, relief=tk.FLAT, cursor="hand2", command=lambda: self.test_coordinates_for("AKTARIM")).pack(side=tk.RIGHT, padx=2)
        self.btn_aktarim_start = tk.Button(bar, text="🚀 Aktarımı Başlat", font=("Segoe UI", 9, "bold"), bg="#059669", fg="#ffffff", padx=12, pady=4, relief=tk.FLAT, cursor="hand2", command=self.start_excel_aktarim_process)
        self.btn_aktarim_start.pack(side=tk.RIGHT, padx=6)

        # PanedWindow (Sol: Excel Listesi / Sağ: Sabit Değerler)
        paned = tk.PanedWindow(parent, orient=tk.HORIZONTAL, bg="#e2e8f0", sashwidth=4, bd=0)
        paned.pack(fill=tk.BOTH, expand=True)

        left = tk.Frame(paned, bg="#ffffff", padx=8, pady=6)
        paned.add(left, minsize=480)

        tk.Label(left, text="📑 KLASÖRDE BULUNAN EXCEL DOSYALARI (Seçip Başlatabilirsiniz):", font=("Segoe UI", 9, "bold"), bg="#ffffff", fg="#1e293b").pack(anchor="w", pady=(0, 4))
        lb_frame = tk.Frame(left, bg="#ffffff")
        lb_frame.pack(fill=tk.BOTH, expand=True)

        sb_ak = tk.Scrollbar(lb_frame)
        sb_ak.pack(side=tk.RIGHT, fill=tk.Y)
        self.lb_excel = tk.Listbox(
            lb_frame, yscrollcommand=sb_ak.set, font=("Consolas", 10),
            selectbackground="#7c3aed", selectforeground="white", activestyle="none"
        )
        self.lb_excel.pack(fill=tk.BOTH, expand=True)
        sb_ak.config(command=self.lb_excel.yview)

        right = tk.Frame(paned, bg="#ffffff", padx=10, pady=6, width=320)
        paned.add(right, minsize=300)

        tk.Label(right, text="⚙️ EVRİM EXCEL AKTARIM SABİTLERİ", font=("Segoe UI", 9, "bold"), bg="#ffffff", fg="#059669").pack(anchor="w")
        tk.Label(right, text="Evrim Aktarım Formu açıldığında yazılan sabitler:", font=("Segoe UI", 8), bg="#ffffff", fg="#64748b").pack(anchor="w", pady=(0, 6))

        sabits = [
            ("Döviz Kodu:", "TL"),
            ("ATR Kodu:", "KORE"),
            ("Alıcı Kodu:", "9122"),
            ("Gümrük Kodu:", "18"),
            ("İşlem Tipi:", "Excelden Beyannameye"),
        ]
        for l_txt, v_txt in sabits:
            r = tk.Frame(right, bg="#ffffff")
            r.pack(fill=tk.X, pady=2)
            tk.Label(r, text=l_txt, font=("Segoe UI", 8, "bold"), bg="#ffffff", width=14, anchor="w").pack(side=tk.LEFT)
            tk.Label(r, text=v_txt, font=("Segoe UI", 8), bg="#f1f5f9", padx=6, pady=2, anchor="w", relief=tk.GROOVE).pack(side=tk.LEFT, fill=tk.X, expand=True)

        # Açıklama Kutusu
        tk.Label(
            right,
            text="\n💡 Çalışma Mantığı:\n1. Evrim'de Beyanname modülünü açın.\n2. Robotu Başlat'a basın.\n3. Sarı kutuya tıkladıktan sonra klavyeden [ F8 ] tuşuna basarak her dosyayı tek tıkla içeri aktarın.",
            font=("Segoe UI", 8), bg="#ffffff", fg="#334155", justify="left"
        ).pack(anchor="w", pady=10)

    # =========================================================================
    # 3. TOPLU PL4 ÖN İZİN EKLEME ROBOTU PANELİ
    # =========================================================================
    def _build_pl4_panel(self, parent):
        cards = tk.Frame(parent, bg="#f8fafc", pady=4)
        cards.pack(fill=tk.X)

        self.pl1_v, self.pl1_s = self._make_card(cards, "PL4 ROBOT DURUMU", "HAZIR", "Dosya listesi bekleniyor", "#0284c7")
        self.pl2_v, self.pl2_s = self._make_card(cards, "İŞLENECEK DOSYALAR", "0 Dosya", "Kuyruk boş", "#059669")
        self.pl3_v, self.pl3_s = self._make_card(cards, "KOD DEĞERİ", "PL4", "Ön İzin Kodu", "#7c3aed")
        self.pl4_v, self.pl4_s = self._make_card(cards, "KALİBRASYON", f"{len(self.rpa.koor_pl4)} / 5 Alan", "pl4_koordinatlar.json", "#ea580c")

        # Kontrol Barı
        bar = tk.Frame(parent, bg="#ffffff", padx=10, pady=8, relief=tk.GROOVE, bd=1)
        bar.pack(fill=tk.X, pady=(4, 6))

        tk.Button(bar, text="➕ Manuel Dosya Ekle", font=("Segoe UI", 8, "bold"), bg="#f1f5f9", padx=8, pady=4, relief=tk.GROOVE, cursor="hand2", command=self.show_pl4_add_dialog).pack(side=tk.LEFT, padx=3)
        tk.Button(bar, text="📂 Excel'den Yükle", font=("Segoe UI", 8), bg="#f1f5f9", padx=8, pady=4, relief=tk.GROOVE, cursor="hand2", command=self.load_pl4_from_excel).pack(side=tk.LEFT, padx=3)
        tk.Button(bar, text="🧹 Listeyi Temizle", font=("Segoe UI", 8), bg="#fee2e2", fg="#b91c1c", padx=8, pady=4, relief=tk.GROOVE, cursor="hand2", command=self.clear_pl4_queue).pack(side=tk.LEFT, padx=3)

        self.var_pl4_test = tk.BooleanVar(value=False)
        tk.Checkbutton(bar, text="Güvenli Test Modu (PL4 Yazmaz)", variable=self.var_pl4_test, fg="#2563eb", bg="#ffffff", font=("Segoe UI", 8, "bold")).pack(side=tk.LEFT, padx=8)

        tk.Button(bar, text="🎯 PL4 Kalibrasyonu (F9)", font=("Segoe UI", 8, "bold"), bg="#ea580c", fg="#ffffff", padx=8, pady=4, relief=tk.FLAT, cursor="hand2", command=lambda: self.start_calibration_for("PL4")).pack(side=tk.RIGHT, padx=4)
        tk.Button(bar, text="👀 Test Et", font=("Segoe UI", 8), bg="#0284c7", fg="#ffffff", padx=8, pady=4, relief=tk.FLAT, cursor="hand2", command=lambda: self.test_coordinates_for("PL4")).pack(side=tk.RIGHT, padx=2)
        self.btn_pl4_start = tk.Button(bar, text="🚀 PL4 Robotunu Başlat", font=("Segoe UI", 9, "bold"), bg="#7c3aed", fg="#ffffff", padx=12, pady=4, relief=tk.FLAT, cursor="hand2", command=self.start_pl4_process)
        self.btn_pl4_start.pack(side=tk.RIGHT, padx=6)

        # Gövde
        body = tk.Frame(parent, bg="#ffffff", padx=10, pady=8, relief=tk.GROOVE, bd=1)
        body.pack(fill=tk.BOTH, expand=True)

        tk.Label(body, text="🏷️ PL4 EKLENECEK DOSYA NUMARALARI LİSTESİ:", font=("Segoe UI", 9, "bold"), bg="#ffffff", fg="#1e293b").pack(anchor="w", pady=(0, 4))
        lb_f = tk.Frame(body, bg="#ffffff")
        lb_f.pack(fill=tk.BOTH, expand=True)

        sb_pl = tk.Scrollbar(lb_f)
        sb_pl.pack(side=tk.RIGHT, fill=tk.Y)
        self.lb_pl4 = tk.Listbox(
            lb_f, yscrollcommand=sb_pl.set, font=("Consolas", 10),
            selectbackground="#0284c7", selectforeground="white"
        )
        self.lb_pl4.pack(fill=tk.BOTH, expand=True)
        sb_pl.config(command=self.lb_pl4.yview)

    # -------------------------------------------------------------------------
    # Kart Üretici Yardımcısı
    # -------------------------------------------------------------------------
    def _make_card(self, parent, title, val, sub, color):
        card = tk.Frame(parent, bg="#ffffff", padx=10, pady=6, relief=tk.GROOVE, bd=1)
        card.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=2)

        strip = tk.Frame(card, bg=color, height=3)
        strip.pack(fill=tk.X, pady=(0, 2))

        tk.Label(card, text=title, font=("Segoe UI", 8, "bold"), fg="#64748b", bg="#ffffff").pack(anchor="w")
        v = tk.Label(card, text=val, font=("Segoe UI", 11, "bold"), fg=color, bg="#ffffff")
        v.pack(anchor="w")
        s = tk.Label(card, text=sub, font=("Segoe UI", 8), fg="#94a3b8", bg="#ffffff")
        s.pack(anchor="w")
        return v, s

    def _log(self, msg: str):
        self.txt_console.insert(tk.END, f"[{time.strftime('%H:%M:%S')}] {msg}\n")
        self.txt_console.see(tk.END)

    def _on_rpa_log(self, msg: str):
        try:
            self.root.after(0, lambda: self._log(msg))
        except Exception:
            pass

    # -------------------------------------------------------------------------
    # KALİBRASYON ÇAĞIRICILARI
    # -------------------------------------------------------------------------
    def start_calibration_for(self, robot_type: str):
        ans = messagebox.askyesno(
            "Kalibrasyon Başlatılsın mı?",
            f"[{robot_type}] Robotu için kalibrasyon sihirbazı başlatılacak.\n\n"
            "Evrim İthalat/İhracat açık olsun (zorla başka monitöre taşınmaz).\n"
            "Her kutu için imleci üzerine götürüp [ F9 ] basın.\n"
            "Kayıt pencere-görelidir; PC kapanınca 3 monitör sırası değişse de isabet korunur.\n"
            "İptal: ESC  |  Robot çalışırken acil durdurma: ESC, F12 veya Ctrl+Alt+F12.\n\nHazır mısınız?"
        )
        if not ans:
            return

        def _step(alan_adi, idx, total):
            pass

        def _done(success):
            def _u():
                if success:
                    messagebox.showinfo("Başarılı", f"[{robot_type}] Kalibrasyonu tamamlandı ve kaydedildi!")
                    self.rpa.load_all_coordinates()
                    self._refresh_calibration_cards()
                else:
                    messagebox.showwarning("İptal", "Kalibrasyon iptal edildi.")
            self.root.after(0, _u)

        self.rpa.run_calibration_wizard(robot_type, _step, _done)

    def probe_evrim_window(self):
        ok = self.rpa.probe_evrim()
        if ok:
            messagebox.showinfo("Evrim Algılandı", "Evrim penceresi bulundu. Ayrıntı robot günlüğündedir.")
        else:
            messagebox.showerror(
                "Evrim Yok",
                "Evrim penceresi bulunamadı.\n\n"
                "ithalsql.exe / ihrsql.exe açık olsun.\n"
                "Simge durumundaysa görev çubuğundan geri yükleyin."
            )

    def test_coordinates_for(self, robot_type: str):
        self.rpa.test_robot_coordinates(robot_type, lambda s: messagebox.showinfo("Test Bitti", f"[{robot_type}] Testi tamamlandı."))

    def _refresh_calibration_cards(self):
        self.c4_v.config(text=f"{len(self.rpa.koor_beyan)} / 29 Alan")
        self.ak4_v.config(text=f"{len(self.rpa.koor_aktarim)} / 14 Alan")
        self.pl4_v.config(text=f"{len(self.rpa.koor_pl4)} / 5 Alan")

    # -------------------------------------------------------------------------
    # BEYANNAME YÖNETİMİ
    # -------------------------------------------------------------------------
    def _on_beyan_mode_change(self):
        m = self.var_beyan_mode.get()
        self.rpa.mode = m
        if m == "FULL_AUTONOMOUS":
            self.c2_v.config(text="🚀 TAM OTONOM", fg="#7c3aed")
            self.c2_s.config(text="Duraksız seri doldurur")
        else:
            self.c2_v.config(text="👁️ YARI OTONOM", fg="#2563eb")
            self.c2_s.config(text="Space tuşu ile onay bekler")

    def apply_params_to_remaining(self):
        new_p = {k: ent.get().strip() for k, ent in self.param_entries.items()}
        self.rpa.update_remaining_params(new_p)
        self._refresh_tree_view()
        messagebox.showinfo("Başarılı", "Parametreler sıradaki tüm dosyalara uygulandı!")

    def add_external_files_to_queue(self, file_records: List[Dict[str, Any]]):
        current_list = list(self.rpa.queue)
        fb = self.param_entries["tasit_kimligi"].get()
        for f in file_records:
            current_list.append(queue_item_from_record(f, fallback_gemi=fb))
        self.rpa.set_queue(current_list)
        self._refresh_tree_view()

    def show_add_file_dialog(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("Dosya Numaraları Ekle")
        dlg.geometry("400x300")
        dlg.transient(self.root)
        dlg.grab_set()

        tk.Label(dlg, text="Dosya Numaralarını Yapıştırın (Her satıra bir adet):", font=("Segoe UI", 9, "bold")).pack(anchor="w", padx=12, pady=(10, 4))
        txt = tk.Text(dlg, font=("Segoe UI", 10), height=8)
        txt.pack(fill=tk.BOTH, expand=True, padx=12, pady=4)

        def _add():
            raw = txt.get("1.0", tk.END).strip().splitlines()
            items = [{"dosya_no": l.strip(), "ref": "", "beyanname_no": ""} for l in raw if l.strip()]
            if items:
                self.add_external_files_to_queue(items)
            dlg.destroy()

        b_box = tk.Frame(dlg)
        b_box.pack(fill=tk.X, padx=12, pady=8)
        tk.Button(b_box, text="Ekle", bg="#0284c7", fg="#ffffff", font=("Segoe UI", 9, "bold"), padx=12, pady=4, command=_add).pack(side=tk.RIGHT)
        tk.Button(b_box, text="Vazgeç", padx=10, pady=4, command=dlg.destroy).pack(side=tk.RIGHT, padx=6)

    def load_from_excel_file(self):
        path = filedialog.askopenfilename(title="Excel Seç", filetypes=[("Excel", "*.xlsx *.xls")])
        if not path or not openpyxl:
            return
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            items = []
            rows = list(ws.iter_rows(values_only=True))
            start = 0
            header_map = {}
            if rows:
                headers = [str(c or "").strip().lower() for c in rows[0]]
                if any(h in headers for h in ("dosya", "dosya no", "dosya_no", "ref", "gemi", "liman")):
                    start = 1
                    for i, h in enumerate(headers):
                        header_map[h] = i
            for row in rows[start:]:
                if not row or not row[0]:
                    continue
                v = str(row[0]).strip()
                if v.lower() in ("dosya", "dosya no", "dosya_no"):
                    continue
                rec = {"dosya_no": v}
                def col(*names, idx=None):
                    for n in names:
                        if n in header_map and header_map[n] < len(row) and row[header_map[n]] is not None:
                            return str(row[header_map[n]]).strip()
                    if idx is not None and idx < len(row) and row[idx] is not None:
                        return str(row[idx]).strip()
                    return ""
                rec["ref"] = col("ref", "hankook_ref", idx=1)
                rec["beyanname_no"] = col("beyanname", "beyanname_no", "beyan", idx=2)
                rec["gemi"] = col("gemi", "tasit", idx=3)
                rec["liman"] = col("liman", idx=4)
                rec["fatura_bedeli"] = col("fatura", "fatura_bedeli", "vergi")
                items.append(rec)
            if items:
                self.rpa.set_queue(items)
                self._refresh_tree_view()
                messagebox.showinfo("Başarılı", f"{len(items)} dosya kuyruğa alındı!")
        except Exception as e:
            messagebox.showerror("Hata", f"Excel okunamadı: {e}")

    def clear_queue(self):
        self.rpa.set_queue([])
        self._refresh_tree_view()

    def _refresh_tree_view(self):
        for r in self.tree_q.get_children():
            self.tree_q.delete(r)

        total = len(self.rpa.queue)
        comp = len([x for x in self.rpa.queue if x["status"] == "TAMAMLANDI"])
        self.c3_v.config(text=f"{comp} / {total} Dosya")
        self.c3_s.config(text=f"Kalan: {total - comp} Dosya")

        for idx, item in enumerate(self.rpa.queue):
            status = item.get("status", "BEKLIYOR")
            self.tree_q.insert(
                "", tk.END, iid=f"q_{idx}",
                values=(
                    idx + 1, item.get("dosya_no", ""), item.get("ref", "—"),
                    item.get("beyanname_no", "—"), item.get("gemi", ""), item.get("liman", ""),
                    status, item.get("message", "")
                ),
                tags=(status,)
            )

    def start_beyanname_process(self):
        if not self.rpa.queue:
            messagebox.showwarning("Uyarı", "Kuyrukta dosya yok!")
            return
        new_p = {k: ent.get().strip() for k, ent in self.param_entries.items()}
        self.rpa.active_params.update(new_p)
        self.btn_beyan_start.config(state=tk.DISABLED)
        self.btn_beyan_pause.config(state=tk.NORMAL)
        self._show_floating_panic_button()
        self.rpa.start_beyanname_process(dry_run=self.var_beyan_test.get())

    # -------------------------------------------------------------------------
    # EXCEL AKTARIM YÖNETİMİ
    # -------------------------------------------------------------------------
    def select_excel_folder(self):
        k = filedialog.askdirectory(initialdir=self.ent_excel_klasor.get())
        if k:
            self.ent_excel_klasor.delete(0, tk.END)
            self.ent_excel_klasor.insert(0, k)
            self.refresh_excel_folder()

    def refresh_excel_folder(self):
        k = self.ent_excel_klasor.get().strip()
        self.lb_excel.delete(0, tk.END)
        self.excel_files = []
        if not os.path.exists(k):
            self._log(f"Excel klasörü bulunamadı: {k}")
            self.ak2_v.config(text="0 Dosya")
            return

        files = sorted([f for f in os.listdir(k) if f.lower().endswith(".xlsx") and not f.startswith("~")])
        self.excel_files = [os.path.join(k, f) for f in files]
        for idx, f in enumerate(files, 1):
            self.lb_excel.insert(tk.END, f"  {idx:>3}.  {f}")

        self.ak2_v.config(text=f"{len(files)} Dosya")
        self.ak2_s.config(text="Aktarıma hazır")
        self._log(f"Excel klasörü tarandı: {len(files)} dosya bulundu.")

    def start_excel_aktarim_process(self):
        if not getattr(self, "excel_files", []):
            messagebox.showwarning("Uyarı", "Klasörde .xlsx dosyası bulunamadı!")
            return

        sel = self.lb_excel.curselection()
        start_i = sel[0] if sel else 0
        to_process = self.excel_files[start_i:]

        test_mode = self.var_aktarim_test.get()
        wait_f8 = self.var_aktarim_f8.get()

        self._show_floating_panic_button()
        self.rpa.start_aktarim_process(to_process, test_modu=test_mode, f8_bekle=wait_f8)

    # -------------------------------------------------------------------------
    # PL4 ROBOTU YÖNETİMİ
    # -------------------------------------------------------------------------
    def show_pl4_add_dialog(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("PL4 Dosya Numaraları Ekle")
        dlg.geometry("400x300")
        dlg.transient(self.root)
        dlg.grab_set()

        tk.Label(dlg, text="PL4 Eklenecek Dosya Numaraları:", font=("Segoe UI", 9, "bold")).pack(anchor="w", padx=12, pady=(10, 4))
        txt = tk.Text(dlg, font=("Segoe UI", 10), height=8)
        txt.pack(fill=tk.BOTH, expand=True, padx=12, pady=4)

        def _add():
            raw = txt.get("1.0", tk.END).strip().splitlines()
            for l in raw:
                v = l.strip()
                if v:
                    self.lb_pl4.insert(tk.END, v)
            self.pl2_v.config(text=f"{self.lb_pl4.size()} Dosya")
            dlg.destroy()

        b_box = tk.Frame(dlg)
        b_box.pack(fill=tk.X, padx=12, pady=8)
        tk.Button(b_box, text="Ekle", bg="#7c3aed", fg="#ffffff", font=("Segoe UI", 9, "bold"), padx=12, pady=4, command=_add).pack(side=tk.RIGHT)
        tk.Button(b_box, text="Vazgeç", padx=10, pady=4, command=dlg.destroy).pack(side=tk.RIGHT, padx=6)

    def load_pl4_from_excel(self):
        path = filedialog.askopenfilename(title="Excel Seç", filetypes=[("Excel", "*.xlsx *.xls")])
        if not path or not openpyxl:
            return
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            cnt = 0
            for row in ws.iter_rows(values_only=True):
                if row and row[0]:
                    v = str(row[0]).strip()
                    if v.lower() not in ("dosya", "dosya no", "dosya_no"):
                        self.lb_pl4.insert(tk.END, v)
                        cnt += 1
            self.pl2_v.config(text=f"{self.lb_pl4.size()} Dosya")
            messagebox.showinfo("Başarılı", f"{cnt} dosya PL4 listesine eklendi!")
        except Exception as e:
            messagebox.showerror("Hata", f"Excel okunamadı: {e}")

    def clear_pl4_queue(self):
        self.lb_pl4.delete(0, tk.END)
        self.pl2_v.config(text="0 Dosya")

    def start_pl4_process(self):
        items = list(self.lb_pl4.get(0, tk.END))
        if not items:
            messagebox.showwarning("Uyarı", "PL4 kuyruğunda dosya yok!")
            return
        self._show_floating_panic_button()
        self.rpa.start_pl4_process(items, test_modu=self.var_pl4_test.get())

    # -------------------------------------------------------------------------
    # YÜZEN PANİK & ACİL DURUM FRENİ
    # -------------------------------------------------------------------------
    def _show_floating_panic_button(self):
        if self.floating_panel and self.floating_panel.winfo_exists():
            return
        self.floating_panel = tk.Toplevel(self.root)
        self.floating_panel.title("RPA PANİK")
        self.floating_panel.geometry("320x90+50+50")
        self.floating_panel.attributes("-topmost", True)
        self.floating_panel.overrideredirect(True)
        self.floating_panel.configure(bg="#0f172a", bd=2, relief=tk.RAISED)

        head = tk.Frame(self.floating_panel, bg="#1e293b", pady=2)
        head.pack(fill=tk.X)
        tk.Label(head, text=f"🤖 {self.rpa.active_robot} ROBOTU AKTİF", font=("Segoe UI", 8, "bold"), fg="#38bdf8", bg="#1e293b").pack(side=tk.LEFT, padx=6)

        btn_box = tk.Frame(self.floating_panel, bg="#0f172a", pady=8)
        btn_box.pack(fill=tk.BOTH, expand=True, padx=8)

        self.btn_fl_ok = tk.Button(
            btn_box, text="✅ ONAYLA (Space)", font=("Segoe UI", 9, "bold"),
            bg="#7c3aed", fg="#ffffff", relief=tk.FLAT, cursor="hand2", command=self.rpa.approve_and_continue
        )
        self.btn_fl_ok.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 4))

        tk.Button(
            btn_box, text="🚨 DURDUR (ESC)", font=("Segoe UI", 9, "bold"),
            bg="#dc2626", fg="#ffffff", relief=tk.FLAT, cursor="hand2", command=self.trigger_emergency_stop
        ).pack(side=tk.RIGHT, fill=tk.BOTH, expand=True)

    def _hide_floating_panic_button(self):
        if self.floating_panel and self.floating_panel.winfo_exists():
            try:
                self.floating_panel.destroy()
            except Exception:
                pass
            self.floating_panel = None

    def trigger_emergency_stop(self):
        self.rpa.trigger_emergency_stop()
        self._hide_floating_panic_button()
        self.btn_beyan_start.config(state=tk.NORMAL)
        self.btn_beyan_pause.config(state=tk.DISABLED)
        self.btn_beyan_resume.config(state=tk.DISABLED)

    # -------------------------------------------------------------------------
    # CALLBACK GÜNCELLEMELERİ
    # -------------------------------------------------------------------------
    def _on_rpa_status_change(self, status: str, msg: str):
        def _u():
            if status == "RUNNING":
                self.c1_v.config(text="⚙️ ÇALIŞIYOR", fg="#7c3aed")
            elif status == "WAITING_APPROVAL":
                self.c1_v.config(text="👁️ ONAY BEKLİYOR", fg="#b45309")
            elif status == "PAUSED":
                self.c1_v.config(text="⏸️ DURAKLATILDI", fg="#d97706")
            elif status == "STOPPED":
                self.c1_v.config(text="🛑 DURDURULDU", fg="#dc2626")
            else:
                self.c1_v.config(text="HAZIR", fg="#0284c7")
            self.c1_s.config(text=msg[:35])
            self._log(msg)
        try:
            self.root.after(0, _u)
        except Exception:
            pass

    def _on_rpa_queue_update(self, idx: int, item: Dict[str, Any]):
        def _u():
            iid = f"q_{idx}"
            if self.tree_q.exists(iid):
                st = item.get("status", "BEKLIYOR")
                self.tree_q.item(
                    iid,
                    values=(
                        idx + 1, item.get("dosya_no", ""), item.get("ref", "—"),
                        item.get("beyanname_no", "—"), item.get("gemi", ""), item.get("liman", ""),
                        st, item.get("message", "")
                    ),
                    tags=(st,)
                )
                self.tree_q.see(iid)

            tot = len(self.rpa.queue)
            comp = len([x for x in self.rpa.queue if x["status"] == "TAMAMLANDI"])
            self.c3_v.config(text=f"{comp} / {tot} Dosya")
            self.c3_s.config(text=f"Kalan: {tot - comp} Dosya")
        try:
            self.root.after(0, _u)
        except Exception:
            pass

    def _on_rpa_step_progress(self, alan_adi: str, cur: int, tot: int):
        pass

    def _on_rpa_approval_needed(self, item: Dict[str, Any]):
        def _bind():
            try:
                self.root.bind("<space>", lambda e: self.rpa.approve_and_continue())
                self.root.bind("<Return>", lambda e: self.rpa.approve_and_continue())
            except Exception:
                pass
        try:
            self.root.after(0, _bind)
        except Exception:
            pass

    def _on_rpa_finish(self, success: bool, msg: str):
        def _u():
            self._hide_floating_panic_button()
            self.btn_beyan_start.config(state=tk.NORMAL)
            self.btn_beyan_pause.config(state=tk.DISABLED)
            self.btn_beyan_resume.config(state=tk.DISABLED)
            if success:
                messagebox.showinfo("Tamamlandı", msg)
            else:
                messagebox.showwarning("Bilgi", msg)
        try:
            self.root.after(0, _u)
        except Exception:
            pass
