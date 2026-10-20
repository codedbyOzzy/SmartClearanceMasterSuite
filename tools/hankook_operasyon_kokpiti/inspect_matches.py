import openpyxl
import re
import win32com.client
import pythoncom

excel_path = r"C:\Users\ozan.gemici\OneDrive - DOGRUER GUMRUK MUSAVIRLIGI A.Ş\Resimler\HANKOOK İŞ TAKİP RAPORU.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb["GÜNCEL RAPOR"]

reg_files = []
for r in range(2, ws.max_row + 1):
    dosya = str(ws.cell(r, 1).value or "").strip()
    ref = str(ws.cell(r, 3).value or "").strip()
    bl = str(ws.cell(r, 6).value or "").strip()
    tescil = str(ws.cell(r, 12).value or "").strip()
    kapanis = str(ws.cell(r, 14).value or "").strip()
    vergi = str(ws.cell(r, 15).value or "").strip()
    if (dosya or ref) and tescil and not kapanis:
        reg_files.append({"row": r, "dosya": dosya, "ref": ref, "bl": bl, "tescil": tescil, "vergi": vergi})

pythoncom.CoInitialize()
outlook = win32com.client.Dispatch("Outlook.Application")
ns = outlook.GetNamespace("MAPI")
sent = ns.GetDefaultFolder(5)
items = sent.Items
items.Sort("[SentOn]", True)

sent_list = []
for i in range(min(items.Count, 350)):
    it = items[i]
    subj = str(it.Subject or "")
    sent_list.append({
        "subj": subj,
        "to": str(it.To or ""),
        "sent_on": str(it.SentOn),
        "atts": [a.FileName for a in it.Attachments if not re.match(r"^image\d+", a.FileName, re.I)]
    })

print(f"Total Tescilli Acik Dosyalar: {len(reg_files)}")
for rf in reg_files:
    ref_c = re.sub(r"[^A-Z0-9]", "", rf["ref"].upper())
    dosya_c = re.sub(r"[^A-Z0-9]", "", rf["dosya"].upper())
    bl_c = re.sub(r"[^A-Z0-9]", "", rf["bl"].upper())

    vergi_matches = []
    other_matches = []
    for m in sent_list:
        subj_c = re.sub(r"[^A-Z0-9]", "", m["subj"].upper())
        att_c = re.sub(r"[^A-Z0-9]", "", " ".join(m["atts"]).upper())

        is_vergi_mail = "VERG" in m["subj"].upper() or "VERGI" in m["subj"].upper()

        if (ref_c and (ref_c in subj_c or ref_c in att_c)) or \
           (dosya_c and len(dosya_c) >= 6 and (dosya_c in subj_c or dosya_c in att_c)) or \
           (bl_c and len(bl_c) >= 7 and (bl_c in subj_c or bl_c in att_c)):
            if is_vergi_mail:
                vergi_matches.append(m)
            else:
                other_matches.append(m)

    d_name = rf["dosya"]
    r_name = rf["ref"]
    t_name = rf["tescil"][:10]
    v_name = rf["vergi"]
    print(f"Dosya: {d_name} | Ref: {r_name} | Tescil: {t_name} | Vergi Odeme: {v_name}")
    print(f"   -> Vergi Talebi Mailleri: {len(vergi_matches)} | Diger Mailler: {len(other_matches)}")
    if vergi_matches:
        print(f"      Vergi Maili: {vergi_matches[0]['subj']} ({vergi_matches[0]['sent_on'][:16]})")
    elif other_matches:
        print(f"      Diger Mail: {other_matches[0]['subj']} ({other_matches[0]['sent_on'][:16]})")
