import openpyxl
import re
import win32com.client
import pythoncom

excel_path = r"C:\Users\ozan.gemici\OneDrive - DOGRUER GUMRUK MUSAVIRLIGI A.Ş\Resimler\HANKOOK İŞ TAKİP RAPORU.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb["GÜNCEL RAPOR"]

reg_22 = []
for r in range(2, ws.max_row + 1):
    dosya = str(ws.cell(r, 1).value or "").strip()
    ref = str(ws.cell(r, 3).value or "").strip()
    tescil = str(ws.cell(r, 12).value or "").strip()
    kapanis = str(ws.cell(r, 14).value or "").strip()
    vergi = str(ws.cell(r, 15).value or "").strip()
    aciklama = str(ws.cell(r, 17).value or "").strip()

    if (dosya or ref) and not kapanis:
        is_tescil = bool(tescil) or ("TESCIL" in aciklama.upper()) or ("TESCİL" in aciklama.upper())
        is_unpaid = not vergi

        if is_tescil and is_unpaid:
            reg_22.append({"row": r, "dosya": dosya, "ref": ref, "tescil": tescil, "aciklama": aciklama})

pythoncom.CoInitialize()
outlook = win32com.client.Dispatch("Outlook.Application")
ns = outlook.GetNamespace("MAPI")
sent = ns.GetDefaultFolder(5)
items = sent.Items
items.Sort("[SentOn]", True)

sent_list = []
for i in range(min(items.Count, 350)):
    it = items[i]
    subj = str(it.Subject or "")
    sent_list.append({
        "subj": subj,
        "to": str(it.To or ""),
        "sent_on": str(it.SentOn),
        "atts": [a.FileName for a in it.Attachments if not re.match(r"^image\d+", a.FileName, re.I)]
    })

atildi_count = 0
bekliyor_count = 0

for rf in reg_22:
    ref_c = re.sub(r"[^A-Z0-9]", "", rf["ref"].upper())
    dosya_c = re.sub(r"[^A-Z0-9]", "", rf["dosya"].upper())

    found = False
    for m in sent_list:
        subj_c = re.sub(r"[^A-Z0-9]", "", m["subj"].upper())
        att_c = re.sub(r"[^A-Z0-9]", "", " ".join(m["atts"]).upper())

        is_vergi_mail = "VERG" in m["subj"].upper() or "VERGI" in m["subj"].upper()

        if is_vergi_mail and ((ref_c and (ref_c in subj_c or ref_c in att_c)) or (dosya_c and len(dosya_c) >= 6 and (dosya_c in subj_c or dosya_c in att_c))):
            found = True
            break

    if found:
        atildi_count += 1
        print(f"ATILDI: {rf['dosya']} | Ref: {rf['ref']} | Tescil: {rf['tescil'][:10]} | Aciklama: {rf['aciklama']}")
    else:
        bekliyor_count += 1
        print(f"BEKLIYOR: {rf['dosya']} | Ref: {rf['ref']} | Tescil: {rf['tescil'][:10]} | Aciklama: {rf['aciklama']}")

print(f"\nOZET: Toplam {len(reg_22)} Tescilli Dosyadan -> {atildi_count} Tanesine Mail Atilmis, {bekliyor_count} Tanesi Mail Bekliyor!")
