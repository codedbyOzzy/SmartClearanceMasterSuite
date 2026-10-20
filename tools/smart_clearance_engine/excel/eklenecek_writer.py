# -*- coding: utf-8 -*-
"""
eklenecek_writer.py — Master Data Kontrol'de bulunan "farklar"i (master'da
hic olmayan urunler + tanimi uyusmayan urunler), kullanicinin "mastera
eklenecek urunler" sablonunun yapisini BOZMADAN, o sablonun temiz bir
kopyasina yazar.

Iki ayri durum icin iki ayri fonksiyon var (build_eklenecek_workbook /
build_tanim_uyusmuyor_workbook) cunku bilinen veri farkli:
- Master'da hic olmayan (missing_code) urunler icin GTIP/Turkce Tanim/
  Anlasma Kodu gibi alanlar BILEREK bos birakilir — bunlar yeni bir urun
  icin gercek gumruk siniflandirmasi gerektirir, sistem tahmin ETMEZ.
- Tanimi uyusmayan (description_mismatch) urunler icin kod zaten master'da
  KAYITLI oldugundan GTIP/Turkce Tanim/Anlasma Kodu master'dan (guvenilir,
  mevcut veri) aynen kopyalanir — burada da TAHMIN yapilmaz, sadece
  zaten bilinen master verisi tasinir.
"""
from copy import copy

import openpyxl

from tools.smart_clearance_engine.excel import eklenecek_schema as S

DATA_SHEET = "Sayfa1"


def build_eklenecek_workbook(template_path, missing_rows):
    """missing_rows: MasterCheckRow listesi (status=missing_code olanlar)."""
    def values_for(item):
        return {
            S.COL_FIRMA_EVRIM_NO: S.FIRMA_EVRIM_NO,
            S.COL_IHRACATCI_EVRIM_NO: S.IHRACATCI_EVRIM_NO,
            S.COL_YABANCI_TANIM: item.invoice_description,
            S.COL_URUN_KODU: item.product_code,
            S.COL_MENSEI: S.ORIGIN_DISPLAY_NAME.get(item.origin, None),
            S.COL_ILISKI: S.ILISKI_SABIT,
            S.COL_ROYALTI: S.ROYALTI_SABIT,
            S.COL_DFIF: S.DFIF_SABIT,
            S.COL_AMBARGO: S.AMBARGO_SABIT,
            S.COL_KDV_ORAN: S.KDV_ORAN_SABIT,
            S.COL_TIP: S.TIP_SABIT,
        }
    return _build_workbook(template_path, missing_rows, values_for)


def build_tanim_uyusmuyor_workbook(template_path, mismatch_rows):
    """mismatch_rows: MasterCheckRow listesi (status=description_mismatch olanlar).
    Her satirin mense'si de KENDI faturasindan gelir (item.origin), missing_code
    durumuyla ayni mantik — bkz. build_eklenecek_workbook."""
    def values_for(item):
        return {
            S.COL_FIRMA_EVRIM_NO: S.FIRMA_EVRIM_NO,
            S.COL_IHRACATCI_EVRIM_NO: S.IHRACATCI_EVRIM_NO,
            S.COL_YABANCI_TANIM: item.invoice_description,
            S.COL_TURKCE_TANIM: item.master_turkish_description,
            S.COL_URUN_KODU: item.product_code,
            S.COL_GTIP: item.master_gtip,
            S.COL_MENSEI: S.ORIGIN_DISPLAY_NAME.get(item.origin, None),
            S.COL_ANLASMA_KODU: item.master_agreement_code,
            S.COL_ILISKI: S.ILISKI_SABIT,
            S.COL_ROYALTI: S.ROYALTI_SABIT,
            S.COL_DFIF: S.DFIF_SABIT,
            S.COL_AMBARGO: S.AMBARGO_SABIT,
            S.COL_KDV_ORAN: S.KDV_ORAN_SABIT,
            S.COL_TIP: S.TIP_SABIT,
        }
    return _build_workbook(template_path, mismatch_rows, values_for)


def _build_workbook(template_path, rows, values_for):
    """Sablonun temiz bir kopyasini acar, her satir icin values_for(item)
    dondurdugu degerleri yazar (stil/format kaynak satirdan kopyalanir).
    Donen: (workbook, yazilan_satir_sayisi)"""
    wb = openpyxl.load_workbook(template_path)
    ws = wb[DATA_SHEET]
    style_row = ws[S.FIRST_DATA_ROW] if ws.max_row >= S.FIRST_DATA_ROW else ws[2]

    _clear_existing_data_rows(ws)

    total_cols = S.TOTAL_COLS
    if len(style_row) > 0:
        total_cols = min(S.TOTAL_COLS, len(style_row))

    row_idx = S.FIRST_DATA_ROW
    for item in rows:
        values = values_for(item)
        for col_idx in range(1, total_cols + 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=values.get(col_idx))
            if col_idx - 1 < len(style_row):
                src = style_row[col_idx - 1]
                if src.font:
                    cell.font = copy(src.font)
                if src.number_format:
                    cell.number_format = src.number_format
        row_idx += 1

    return wb, row_idx - S.FIRST_DATA_ROW


def _clear_existing_data_rows(ws):
    if ws.max_row >= S.FIRST_DATA_ROW:
        ws.delete_rows(S.FIRST_DATA_ROW, ws.max_row - S.FIRST_DATA_ROW + 1)
