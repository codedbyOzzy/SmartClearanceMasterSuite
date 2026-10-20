# -*- coding: utf-8 -*-
"""
muhasebe_writer.py — Muhasebe Takip panelinin renkli Excel çıktısını
üretir.
"""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

SHEET_NAME = "Eşleştirme Portalı"

HEADERS = ["Sıra No", "Liman", "Dosya Numarası", "PO No", "Taşıyıcı", "Konşimento No", "Eşleşme Durumu", "Açıklama"]

_HEADER_FILL = PatternFill(start_color="FF1E1B4B", end_color="FF1E1B4B", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFFFF", bold=True)

_MATCH_ROW_FILL = PatternFill(start_color="FFF0FDF4", end_color="FFF0FDF4", fill_type="solid")
_MATCH_BADGE_FILL = PatternFill(start_color="FFDCFCE7", end_color="FFDCFCE7", fill_type="solid")
_MATCH_BADGE_FONT = Font(color="FF166534", bold=True)

_UNMATCH_ROW_FILL = PatternFill(start_color="FFFEF2F2", end_color="FFFEF2F2", fill_type="solid")
_UNMATCH_STATUS_FILL = PatternFill(start_color="FFFEE2E2", end_color="FFFEE2E2", fill_type="solid")
_UNMATCH_STATUS_FONT = Font(color="FF991B1B", bold=True)
_UNMATCH_ACIKLAMA_FILL = PatternFill(start_color="FFFEF3C7", end_color="FFFEF3C7", fill_type="solid")
_UNMATCH_ACIKLAMA_FONT = Font(color="FF92400E", bold=True)

_DOSYA_NO_FONT = Font(color="FF1E1B4B", bold=True)

_COL_WIDTHS = [8, 14, 16, 14, 24, 20, 22, 16]


def build_muhasebe_workbook(satirlar: list) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    for col_idx, baslik in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=baslik)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row_idx, s in enumerate(satirlar, start=2):
        gonderildi = s["gonderildi"]
        row_fill = _MATCH_ROW_FILL if gonderildi else _UNMATCH_ROW_FILL
        values = [
            s["sira_no"], s["liman"], s["dosya_no"], s["po_no"],
            s["tasiyici"], s["konsimento_no"], s["eslesme_durumu"], s["aciklama"],
        ]
        for col_idx, deger in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=deger)
            cell.fill = row_fill
            if col_idx == 3:  # Dosya Numarası — anahtar sütun, vurgulu
                cell.font = _DOSYA_NO_FONT

        # Eşleşme Durumu (7) ve Açıklama (8) rozet renkleri
        durum_cell = ws.cell(row=row_idx, column=7)
        aciklama_cell = ws.cell(row=row_idx, column=8)
        if gonderildi:
            durum_cell.fill = _MATCH_BADGE_FILL
            durum_cell.font = _MATCH_BADGE_FONT
            aciklama_cell.fill = _MATCH_BADGE_FILL
            aciklama_cell.font = _MATCH_BADGE_FONT
        else:
            durum_cell.fill = _UNMATCH_STATUS_FILL
            durum_cell.font = _UNMATCH_STATUS_FONT
            aciklama_cell.fill = _UNMATCH_ACIKLAMA_FILL
            aciklama_cell.font = _UNMATCH_ACIKLAMA_FONT

    for col_idx, width in enumerate(_COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    return wb
