# -*- coding: utf-8 -*-
"""
==============================================================================
🚢 MSC CANLI KONTEYNER TAKİP & VARIŞ (ETA) RADARI
Firma: Doğruer Gümrük Müşavirliği A.Ş.
Modül: Bağımsız Yardımcı Araç (Stand-Alone Tool)
Açıklama: MSC sisteminden canlı konteyner, gemi, aktarma ve kesin ETA sorgular.
==============================================================================
"""

import os
import sys
import json
import time
import argparse
import threading
from datetime import datetime

# Windows encoding fix
if sys.platform == "win32":
    import codecs
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Curl cffi import
try:
    from curl_cffi import requests
except ImportError:
    print("[HATA] 'curl_cffi' kütüphanesi bulunamadı. Lütfen 'pip install curl_cffi' çalıştırın.")
    sys.exit(1)

# Openpyxl for excel support
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    openpyxl = None

# Tkinter for GUI
try:
    import tkinter as tk
    from tkinter import ttk, messagebox, filedialog
except ImportError:
    tk = None


class MSCLiveTracker:
    """MSC Canlı Konteyner ve Konşimento Takip Motoru."""

    BASE_URL = "https://www.msc.com/en/track-a-shipment"
    API_URL = "https://www.msc.com/api/feature/tools/TrackingInfo"

    def __init__(self):
        self.session = requests.Session(impersonate="chrome124")
        self._initialized = False
        self.verification_token = ""

    def init_session(self):
        """MSC ana sayfasına bağlanarak gerekli oturum çerezlerini alır."""
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9,tr;q=0.8"
        }
        try:
            r = self.session.get(self.BASE_URL, headers=headers, timeout=15)
            if r.status_code == 200:
                self._initialized = True
                try:
                    import re
                    m = re.search(r'name=["\']__RequestVerificationToken["\']\s+type=["\']hidden["\']\s+value=["\']([^"\']+)["\']', r.text)
                    if not m:
                        m = re.search(r'value=["\']([^"\']+)["\']\s+name=["\']__RequestVerificationToken["\']', r.text)
                    if m:
                        self.verification_token = m.group(1)
                except Exception:
                    pass
                return True
        except Exception as e:
            print(f"[UYARI] MSC oturum başlatılamadı: {e}")
        return False

    def track(self, tracking_number, mode="0"):
        """
        Konteyner veya Konşimento numarasını sorgular.
        mode="0" -> Container / Bill of Lading
        mode="1" -> Booking Number
        """
        if not self._initialized:
            self.init_session()

        clean_number = str(tracking_number or "").strip().upper()
        if not clean_number:
            return {"success": False, "error": "Boş takip numarası"}

        payload = {
            "trackingNumber": clean_number,
            "trackingMode": str(mode)
        }

        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/json",
            "X-Requested-With": "XMLHttpRequest",
            "Origin": "https://www.msc.com",
            "Referer": self.BASE_URL
        }
        if self.verification_token:
            headers["__RequestVerificationToken"] = self.verification_token

        try:
            resp = self.session.post(self.API_URL, json=payload, headers=headers, timeout=15)
            if resp.status_code != 200:
                return {"success": False, "error": f"HTTP {resp.status_code}"}

            data = resp.json()
            if not data.get("IsSuccess") or not data.get("Data"):
                msg = data.get("Message") or data.get("errorMessage") or "Kayıt bulunamadı veya geçersiz numara"
                return {"success": False, "error": msg}

            return self._parse_msc_response(clean_number, data["Data"])

        except Exception as e:
            return {"success": False, "error": str(e)}

    def _parse_msc_response(self, query_number, raw_data):
        """MSC ham JSON yanıtını temiz ve standart bir formata çevirir."""
        bols = raw_data.get("BillOfLadings", [])
        parsed_containers = []

        today = datetime.now().date()

        for b in bols:
            bl_no = b.get("BillOfLadingNumber") or "—"
            gen_info = b.get("GeneralTrackingInfo") or {}
            pol = gen_info.get("PortOfLoad") or "—"
            pod = gen_info.get("PortOfDischarge") or "—"
            final_eta = gen_info.get("FinalPodEtaDate") or ""

            containers = b.get("ContainersInfo") or []
            for c in containers:
                cntr_no = c.get("ContainerNumber") or query_number
                cntr_type = c.get("ContainerType") or "—"
                cntr_pod_eta = c.get("PodEtaDate") or final_eta or ""
                latest_move = c.get("LatestMove") or "—"
                delivered = c.get("Delivered", False)

                events = c.get("Events") or []
                vessel_name = "—"
                voyage_no = "—"
                eta_date_str = cntr_pod_eta
                eta_location = pod
                transshipments = []

                # Olayları incele
                for ev in events:
                    desc = str(ev.get("Description") or "")
                    ev_date = ev.get("Date") or ""
                    ev_loc = ev.get("Location") or ""
                    ev_detail = ev.get("Detail") or []

                    if "Estimated Time of Arrival" in desc or "ETA" in desc:
                        if ev_date:
                            eta_date_str = ev_date
                        if ev_loc:
                            eta_location = ev_loc
                        if ev_detail and len(ev_detail) >= 1:
                            vessel_name = ev_detail[0]
                        if ev_detail and len(ev_detail) >= 2:
                            voyage_no = ev_detail[1]

                    # Aktarma limanları
                    intermediary = ev.get("IntermediaryPortCalls") or []
                    for port_call in intermediary:
                        loc = port_call.get("LocationName")
                        if loc and loc not in transshipments:
                            transshipments.append(loc)

                    if vessel_name == "—" and ev_detail:
                        if len(ev_detail) >= 1 and not ev_detail[0].startswith("LADEN") and not ev_detail[0].startswith("EMPTY"):
                            vessel_name = ev_detail[0]

                # Gün ve durum hesabı
                kalan_gun_str = "—"
                fark_gun = None
                status_tag = "normal"

                if eta_date_str:
                    clean_dt_str = eta_date_str.split(" ")[0].strip()
                    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y"):
                        try:
                            eta_dt = datetime.strptime(clean_dt_str, fmt).date()
                            fark_gun = (eta_dt - today).days
                            if delivered:
                                kalan_gun_str = "Teslim Edildi"
                                status_tag = "delivered"
                            elif fark_gun == 0:
                                kalan_gun_str = "Bugün Varış"
                                status_tag = "today"
                            elif fark_gun > 0:
                                kalan_gun_str = f"{fark_gun} Gün Kaldı"
                                status_tag = "soon" if fark_gun <= 3 else "normal"
                            else:
                                kalan_gun_str = f"Gecikti ({abs(fark_gun)} gün)"
                                status_tag = "delayed"
                            break
                        except Exception:
                            pass

                parsed_containers.append({
                    "container_no": cntr_no,
                    "bl_no": bl_no,
                    "container_type": cntr_type,
                    "vessel": vessel_name,
                    "voyage": voyage_no,
                    "pol": pol,
                    "pod": pod,
                    "eta_location": eta_location,
                    "live_eta": eta_date_str,
                    "fark_gun": fark_gun,
                    "kalan_gun": kalan_gun_str,
                    "status_tag": status_tag,
                    "latest_move": latest_move,
                    "transshipments": " -> ".join(transshipments) if transshipments else "Yok (Direkt)",
                    "delivered": delivered,
                    "events_count": len(events)
                })

        return {
            "success": True,
            "query": query_number,
            "total_containers": len(parsed_containers),
            "containers": parsed_containers
        }


# ==============================================================================
# GRAFİK ARAYÜZ (TKINTER GUI)
# ==============================================================================

class MSCTrackerGUI:
    """MSC Canlı Takip Kullanıcı Arayüzü."""

    DEFAULT_EXCEL = r"C:\Users\ozan.gemici\OneDrive - DOGRUER GUMRUK MUSAVIRLIGI A.Ş\Belgeler\HANKOOK_IS_TAKIP_RAPORU_OZEL.xlsx"

    def __init__(self, root):
        self.root = root
        self.root.title("🚢 Doğruer — MSC Canlı Konteyner & ETA Takip Radarı")
        self.root.geometry("1150x680")
        self.root.minsize(950, 550)
        self.root.configure(bg="#0f172a")

        self.tracker = MSCLiveTracker()
        self.results_data = []

        self._setup_styles()
        self._setup_ui()

    def _setup_styles(self):
        style = ttk.Style()
        style.theme_use("clam")

        # Treeview styling
        style.configure(
            "Custom.Treeview",
            background="#ffffff",
            foreground="#0f172a",
            fieldbackground="#ffffff",
            font=("Segoe UI", 9),
            rowheight=28
        )
        style.configure(
            "Custom.Treeview.Heading",
            background="#1e293b",
            foreground="#ffffff",
            font=("Segoe UI", 9, "bold"),
            relief="flat"
        )
        style.map("Custom.Treeview.Heading", background=[("active", "#334155")])

    def _setup_ui(self):
        # Üst Başlık Banner
        header = tk.Frame(self.root, bg="#1e293b", padx=20, pady=14)
        header.pack(fill=tk.X)

        lbl_title = tk.Label(
            header, text="🚢 MSC CANLI KONTEYNER & ETA VARIŞ RADARI",
            font=("Segoe UI", 15, "bold"), fg="#38bdf8", bg="#1e293b"
        )
        lbl_title.pack(anchor="w")

        lbl_sub = tk.Label(
            header,
            text="MSC resmi altyapısından canlı konteyner konumu, aktarma gemisi ve tahliye limanı kesin ETA verilerini çeker.",
            font=("Segoe UI", 9), fg="#94a3b8", bg="#1e293b"
        )
        lbl_sub.pack(anchor="w", pady=(3, 0))

        # Kontrol Paneli
        ctrl = tk.Frame(self.root, bg="#0f172a", padx=16, pady=12)
        ctrl.pack(fill=tk.X)

        # Tekil Arama Kutusu
        f_search = tk.LabelFrame(ctrl, text="  Tekil Konteyner / B/L Sorgula  ", bg="#0f172a", fg="#38bdf8", font=("Segoe UI", 9, "bold"), padx=10, pady=8)
        f_search.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 10))

        self.txt_single = tk.Entry(f_search, font=("Segoe UI", 11, "bold"), width=18, bd=1, relief=tk.SOLID)
        self.txt_single.pack(side=tk.LEFT, padx=(0, 8), ipady=3)
        self.txt_single.insert(0, "MSNU8608960")
        self.txt_single.bind("<Return>", lambda e: self.search_single())

        btn_search = tk.Button(
            f_search, text="🔍 Canlı Sorgula", font=("Segoe UI", 9, "bold"),
            bg="#0284c7", fg="#ffffff", activebackground="#0369a1", activeforeground="#ffffff",
            relief=tk.FLAT, padx=14, pady=4, cursor="hand2", command=self.search_single
        )
        btn_search.pack(side=tk.LEFT)

        # Toplu Excel İşlemleri
        f_excel = tk.LabelFrame(ctrl, text="  Toplu Excel Dosyasından Tara  ", bg="#0f172a", fg="#10b981", font=("Segoe UI", 9, "bold"), padx=10, pady=8)
        f_excel.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        btn_excel_scan = tk.Button(
            f_excel, text="📂 Excel'den MSC Dosyalarını Tara", font=("Segoe UI", 9, "bold"),
            bg="#10b981", fg="#ffffff", activebackground="#059669", activeforeground="#ffffff",
            relief=tk.FLAT, padx=12, pady=4, cursor="hand2", command=self.scan_from_excel
        )
        btn_excel_scan.pack(side=tk.LEFT, padx=(0, 8))

        btn_export = tk.Button(
            f_excel, text="📊 Sonuçları Excel İndir", font=("Segoe UI", 9, "bold"),
            bg="#6366f1", fg="#ffffff", activebackground="#4f46e5", activeforeground="#ffffff",
            relief=tk.FLAT, padx=12, pady=4, cursor="hand2", command=self.export_results_excel
        )
        btn_export.pack(side=tk.LEFT)

        self.lbl_status = tk.Label(
            f_excel, text="Hazır", font=("Segoe UI", 9, "italic"),
            fg="#94a3b8", bg="#0f172a"
        )
        self.lbl_status.pack(side=tk.LEFT, padx=(12, 0))

        # Tablo Alanı
        t_frame = tk.Frame(self.root, bg="#ffffff")
        t_frame.pack(fill=tk.BOTH, expand=True, padx=16, pady=(0, 10))

        cols = ("cntr", "bl", "gemi", "pol", "pod", "eta", "kalan", "aktarma", "son_hareket")
        self.tree = ttk.Treeview(t_frame, columns=cols, show="headings", style="Custom.Treeview")

        self.tree.heading("cntr", text="KONTEYNER NO")
        self.tree.heading("bl", text="MSC B/L NO")
        self.tree.heading("gemi", text="GÜNCEL GEMİ / SEFER")
        self.tree.heading("pol", text="YÜKLEME LİMANI")
        self.tree.heading("pod", text="VARIŞ LİMANI")
        self.tree.heading("eta", text="CANLI ETA")
        self.tree.heading("kalan", text="DURUM / KALAN GÜN")
        self.tree.heading("aktarma", text="AKTARMA LİMANLARI")
        self.tree.heading("son_hareket", text="SON HAREKET")

        self.tree.column("cntr", width=120, anchor="center")
        self.tree.column("bl", width=120, anchor="center")
        self.tree.column("gemi", width=150, anchor="w")
        self.tree.column("pol", width=110, anchor="center")
        self.tree.column("pod", width=120, anchor="center")
        self.tree.column("eta", width=100, anchor="center")
        self.tree.column("kalan", width=130, anchor="center")
        self.tree.column("aktarma", width=150, anchor="w")
        self.tree.column("son_hareket", width=120, anchor="center")

        sb = ttk.Scrollbar(t_frame, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree.pack(fill=tk.BOTH, expand=True)

        # Renk Etiketleri
        self.tree.tag_configure("today", background="#f0fdf4", foreground="#16a34a")     # Yeşil
        self.tree.tag_configure("delayed", background="#fef2f2", foreground="#dc2626")   # Kırmızı
        self.tree.tag_configure("soon", background="#fffbeb", foreground="#d97706")      # Turuncu
        self.tree.tag_configure("normal", background="#ffffff", foreground="#0f172a")
        self.tree.tag_configure("delivered", background="#f1f5f9", foreground="#64748b")

        # Alt Bilgi Çubuğu
        status_bar = tk.Frame(self.root, bg="#1e293b", padx=16, pady=6)
        status_bar.pack(fill=tk.X, side=tk.BOTTOM)

        self.lbl_summary = tk.Label(
            status_bar, text="Kayıt sayısı: 0", font=("Segoe UI", 9),
            fg="#94a3b8", bg="#1e293b"
        )
        self.lbl_summary.pack(side=tk.LEFT)

        lbl_cred = tk.Label(
            status_bar, text="Doğruer SmartClearance — Doğrudan Canlı MSC API Entegrasyonu",
            font=("Segoe UI", 8), fg="#64748b", bg="#1e293b"
        )
        lbl_cred.pack(side=tk.RIGHT)

    def search_single(self):
        query = self.txt_single.get().strip().upper()
        if not query:
            messagebox.showwarning("Uyarı", "Lütfen bir konteyner veya B/L numarası yazın.")
            return

        self.lbl_status.config(text=f"🔄 {query} sorgulanıyor...", fg="#38bdf8")

        def _work():
            res = self.tracker.track(query)
            self.root.after(0, lambda: self._handle_single_result(res))

        threading.Thread(target=_work, daemon=True).start()

    def _handle_single_result(self, res):
        if not res.get("success"):
            self.lbl_status.config(text=f"❌ Hata: {res.get('error')}", fg="#ef4444")
            messagebox.showerror("Sorgu Hatası", f"MSC Sorgusu Başarısız:\n{res.get('error')}")
            return

        containers = res.get("containers", [])
        self.results_data.extend(containers)
        self._populate_table()
        self.lbl_status.config(text=f"✅ {len(containers)} kayıt getirildi", fg="#10b981")

    def scan_from_excel(self):
        target_path = self.DEFAULT_EXCEL
        if not os.path.exists(target_path):
            target_path = filedialog.askopenfilename(
                title="İş Takip Excel Dosyasını Seçin",
                filetypes=[("Excel Dosyaları", "*.xlsx *.xls")]
            )

        if not target_path or not os.path.exists(target_path):
            return

        if not openpyxl:
            messagebox.showerror("Hata", "openpyxl kütüphanesi eksik.")
            return

        self.lbl_status.config(text="📂 Excel okunuyor...", fg="#38bdf8")

        def _work():
            try:
                wb = openpyxl.load_workbook(target_path, data_only=True)
                ws = wb.active
                for name in wb.sheetnames:
                    if "TAKİP" in name.upper() or "RAPOR" in name.upper():
                        ws = wb[name]
                        break

                candidates = []
                header_row = 1
                for r in range(1, 15):
                    row_vals = [str(ws.cell(r, c).value or "").upper() for c in range(1, 25)]
                    if any("KONTEYNER" in v or "ACENTA" in v or "DOSYA" in v for v in row_vals):
                        header_row = r
                        break

                for r in range(header_row + 1, ws.max_row + 1):
                    row_str = " ".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1)).upper()
                    if "MSC" in row_str or "GLOVIS" in row_str or "MEDITERRANEAN" in row_str:
                        for c in range(1, ws.max_column + 1):
                            val = str(ws.cell(r, c).value or "").strip().upper()
                            if len(val) == 11 and val[:4].isalpha() and val[4:].isdigit():
                                if val not in candidates:
                                    candidates.append(val)

                if not candidates:
                    self.root.after(0, lambda: messagebox.showinfo("Bilgi", "Excel'de MSC ile ilişkili konteyner bulunamadı."))
                    return

                self.root.after(0, lambda: self.lbl_status.config(text=f"🔄 {len(candidates)} konteyner MSC'den taranıyor...", fg="#38bdf8"))

                scanned = []
                for idx, cntr in enumerate(candidates, 1):
                    self.root.after(0, lambda c=cntr, i=idx, tot=len(candidates): self.lbl_status.config(
                        text=f"[{i}/{tot}] {c} sorgulanıyor..."
                    ))
                    res = self.tracker.track(cntr)
                    if res.get("success"):
                        scanned.extend(res.get("containers", []))
                    time.sleep(0.5)

                self.results_data = scanned
                self.root.after(0, lambda: self._on_bulk_done(len(scanned), len(candidates)))

            except Exception as e:
                self.root.after(0, lambda err=str(e): messagebox.showerror("Hata", f"Excel okuma hatası:\n{err}"))

        threading.Thread(target=_work, daemon=True).start()

    def _on_bulk_done(self, count, total_candidates):
        self._populate_table()
        self.lbl_status.config(text=f"✅ Tarama tamamlandı ({count}/{total_candidates} yüke ulaşıldı)", fg="#10b981")
        messagebox.showinfo("Tamamlandı", f"MSC Canlı Tarama Tamamlandı!\nToplam bulunan: {count} konteyner.")

    def _populate_table(self):
        for r in self.tree.get_children():
            self.tree.delete(r)

        for item in self.results_data:
            gemi_info = f"{item['vessel']} ({item['voyage']})" if item['voyage'] != "—" else item['vessel']
            self.tree.insert(
                "", tk.END,
                values=(
                    item["container_no"],
                    item["bl_no"],
                    gemi_info,
                    item["pol"],
                    item["pod"],
                    item["live_eta"],
                    item["kalan_gun"],
                    item["transshipments"],
                    item["latest_move"]
                ),
                tags=(item["status_tag"],)
            )

        self.lbl_summary.config(text=f"Toplam Takip Edilen Konteyner: {len(self.results_data)}")

    def export_results_excel(self):
        if not self.results_data:
            messagebox.showwarning("Uyarı", "Dışa aktarılacak veri bulunmuyor.")
            return

        if not openpyxl:
            messagebox.showerror("Hata", "openpyxl kütüphanesi yüklü değil.")
            return

        save_path = filedialog.asksaveasfilename(
            title="Excel Olarak Kaydet",
            defaultextension=".xlsx",
            filetypes=[("Excel Çalışma Kitabı", "*.xlsx")],
            initialfile=f"MSC_Canli_ETA_Raporu_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        )
        if not save_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "MSC Canlı ETA Radarı"

            ws.append(["DOĞRUER SMARTCLEARANCE — MSC CANLI KONTEYNER & ETA TAKİP RAPORU"])
            ws.append([f"Rapor Tarihi: {datetime.now().strftime('%d.%m.%Y %H:%M')} | Toplam Yük: {len(self.results_data)}"])
            ws.append([])

            headers = ["KONTEYNER NO", "MSC B/L NO", "GÜNCEL GEMİ", "SEFER", "YÜKLEME LİMANI", "TAHLİYE LİMANI", "CANLI ETA", "DURUM / KALAN GÜN", "AKTARMA LİMANLARI", "SON HAREKET"]
            ws.append(headers)

            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(4, col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for item in self.results_data:
                ws.append([
                    item["container_no"],
                    item["bl_no"],
                    item["vessel"],
                    item["voyage"],
                    item["pol"],
                    item["pod"],
                    item["live_eta"],
                    item["kalan_gun"],
                    item["transshipments"],
                    item["latest_move"]
                ])

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(save_path)
            messagebox.showinfo("Başarılı", f"Rapor başarıyla kaydedildi:\n{save_path}")

        except Exception as e:
            messagebox.showerror("Hata", f"Excel kaydedilemedi:\n{e}")


# ==============================================================================
# KOMUT SATIRI (CLI) ÇALIŞTIRMA
# ==============================================================================

def run_cli(args):
    tracker = MSCLiveTracker()
    print("=" * 70)
    print("🚢 MSC CANLI KONTEYNER TAKİP SİSTEMİ (CLI)")
    print("=" * 70)

    if args.container:
        print(f"\nSorgulanıyor: {args.container}...")
        res = tracker.track(args.container)
        if not res.get("success"):
            print(f"[HATA] {res.get('error')}")
            return

        for c in res.get("containers", []):
            print("\n" + "-" * 50)
            print(f"📦 Konteyner No     : {c['container_no']} ({c['container_type']})")
            print(f"📄 MSC B/L No       : {c['bl_no']}")
            print(f"🚢 Gemi / Sefer     : {c['vessel']} ({c['voyage']})")
            print(f"📍 Yükleme Limanı   : {c['pol']}")
            print(f"🎯 Tahliye Limanı   : {c['pod']}")
            print(f"📅 Canlı ETA        : {c['live_eta']} -> {c['kalan_gun']}")
            print(f"🔄 Aktarma Limanları: {c['transshipments']}")
            print(f"📍 Son Hareket      : {c['latest_move']}")
            print("-" * 50)


def main():
    parser = argparse.ArgumentParser(description="MSC Canlı Konteyner & ETA Takip Radarı")
    parser.add_argument("container", nargs="?", help="Sorgulanacak MSC Konteyner veya B/L No")
    parser.add_argument("--cli", action="store_true", help="Grafik arayüz olmadan terminalde çalıştır")
    args = parser.parse_args()

    if args.container or args.cli:
        run_cli(args)
    else:
        if tk:
            root = tk.Tk()
            app = MSCTrackerGUI(root)
            root.mainloop()
        else:
            print("Tkinter bulunamadı, CLI modunda çalıştırılıyor.")
            run_cli(args)


if __name__ == "__main__":
    main()
