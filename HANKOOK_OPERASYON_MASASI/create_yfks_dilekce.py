import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_font(run, name='Arial', size=11, bold=False, color=RGBColor(0, 0, 0)):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color

# 1. Header (Centered)
p_tc = doc.add_paragraph()
p_tc.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_tc.paragraph_format.space_after = Pt(4)
p_tc.paragraph_format.line_spacing = 1.15

r = p_tc.add_run('T.C.\nTİCARET BAKANLIĞI\nDOĞU MARMARA GÜMRÜK VE DIŞ TİCARET BÖLGE MÜDÜRLÜĞÜ\nDERİNCE GÜMRÜK MÜDÜRLÜĞÜ\'NE')
set_font(r, name='Arial', size=11, bold=True)

# 2. Date (Right Aligned)
p_date = doc.add_paragraph()
p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_date.paragraph_format.space_before = Pt(16)
p_date.paragraph_format.space_after = Pt(20)
r_date = p_date.add_run('11.09.2026')
set_font(r_date, name='Arial', size=11, bold=True)

# 3. Konu & İlgi
p_konu = doc.add_paragraph()
p_konu.paragraph_format.space_after = Pt(4)
p_konu.paragraph_format.line_spacing = 1.15
r1 = p_konu.add_run('Konu: ')
set_font(r1, name='Arial', size=11, bold=True)
r2 = p_konu.add_run('Yabancı Firma Kayıt Sistemi')
set_font(r2, name='Arial', size=11, bold=False)

p_ilgi = doc.add_paragraph()
p_ilgi.paragraph_format.space_after = Pt(20)
p_ilgi.paragraph_format.line_spacing = 1.15
r3 = p_ilgi.add_run('İlgi: ')
set_font(r3, name='Arial', size=11, bold=True)
r4 = p_ilgi.add_run('12.06.2019 tarihli ve 2019/24 sayılı Genelge')
set_font(r4, name='Arial', size=11, bold=False)

# 4. Body Paragraph
p_body = doc.add_paragraph()
p_body.paragraph_format.space_after = Pt(24)
p_body.paragraph_format.line_spacing = 1.25
r_body = p_body.add_run('İlgi de yer alan tarih ve sayılı genelgeye istinaden aşağıda belirttiğimiz firmanın Yabancı Firma Kayıt Sistemi\'ne kaydı gerçekleştirilmiş olup sistem onayının yapılmasını emir ve müsaadelerinize arz ederiz.')
set_font(r_body, name='Arial', size=11, bold=False)

# 5. Table
table = doc.add_table(rows=2, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

col_widths = [Inches(1.8), Inches(3.2), Inches(1.5)]
headers = ['YFKS NUMARASI', 'FİRMA ÜNVANI', 'FİRMA ÜLKESİ']
values = ['26SG001042025', 'HANKOOK TIRE SINGAPORE PTE. LTD', 'SİNGAPUR']

# Header
hdr_cells = table.rows[0].cells
for i, name in enumerate(headers):
    hdr_cells[i].text = name
    p = hdr_cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.runs[0], name='Arial', size=10, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EBF1F5"/>')
    hdr_cells[i]._tc.get_or_add_tcPr().append(shd)

# Data
val_cells = table.rows[1].cells
for i, val in enumerate(values):
    val_cells[i].text = val
    p = val_cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.runs[0], name='Arial', size=10, bold=False)

# Set widths & borders
for row in table.rows:
    for i, w in enumerate(col_widths):
        row.cells[i].width = w
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="999999"/><w:left w:val="single" w:sz="4" w:space="0" w:color="999999"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="999999"/><w:right w:val="single" w:sz="4" w:space="0" w:color="999999"/></w:tcBorders>')
        row.cells[i]._tc.get_or_add_tcPr().append(tcBorders)

# 6. Sign-off (Right Aligned)
p_sign = doc.add_paragraph()
p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_sign.paragraph_format.space_before = Pt(36)
p_sign.paragraph_format.line_spacing = 1.2
r_s1 = p_sign.add_run('Saygılarımızla,\n')
set_font(r_s1, name='Arial', size=11, bold=False)
r_s2 = p_sign.add_run('HANKOOK LASTİKLERİ A.Ş.\n')
set_font(r_s2, name='Arial', size=11, bold=True)
r_s3 = p_sign.add_run('Doğruer Gümrük Müşavirliği A.Ş.\n')
set_font(r_s3, name='Arial', size=10, bold=False)
r_s4 = p_sign.add_run('(Yetkili Temsilci)')
set_font(r_s4, name='Arial', size=9, bold=False, color=RGBColor(0x66, 0x66, 0x66))

out_path1 = r'C:\Users\ozan.gemici\OneDrive - DOGRUER GUMRUK MUSAVIRLIGI A.Ş\Belgeler\HANKOOK_ONCELIKLI_DOSYALAR\SINGAPUR_SEVKIYATLARI\YFKS_Dilekcesi_HANKOOK_SINGAPORE.docx'
out_path2 = r'C:\Users\ozan.gemici\Downloads\YFKS_Dilekcesi_HANKOOK_SINGAPORE.docx'

doc.save(out_path1)
doc.save(out_path2)
print('Saved to:', out_path1)
print('Saved to:', out_path2)
