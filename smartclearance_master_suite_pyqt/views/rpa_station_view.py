# -*- coding: utf-8 -*-
"""
rpa_station_view.py — Otonom RPA İstasyonu Görünümü (PyQt6)
3 Robot Entegre (Beyanname, Excel Aktarım, PL4) ve Bağımsız Kalibrasyon Merkezi.
"""

import os
import sys
import time
from typing import Dict, Any, List, Optional

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit,
    QPushButton, QTableWidget, QTableWidgetItem, QHeaderView,
    QFrame, QMessageBox, QFileDialog, QStackedWidget, QRadioButton,
    QButtonGroup, QCheckBox, QComboBox, QTextEdit, QSplitter,
    QListWidget, QDialog, QAbstractItemView
)
from PyQt6.QtCore import Qt, QTimer
from PyQt6.QtGui import QColor, QFont, QCursor, QKeySequence, QShortcut

# Lib/site-packages'i sys.path'e dahil et
_root = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
_lib_pkgs = os.path.join(_root, "Lib", "site-packages")
if os.path.exists(_lib_pkgs) and _lib_pkgs not in sys.path:
    sys.path.insert(0, _lib_pkgs)

from .base_view import BaseView
from ..components.metric_card import MetricCard

from tools.dogruer_master_suite.engine.rpa_engine import (
    RPAEngine, DEFAULT_SABIT_DEGERLER,
    BEYANNAME_ALANLAR, AKTARIM_ALANLAR, PL4_ALANLAR,
    KOOR_BEYANNAME_P1, KOOR_BEYANNAME_P2, KOOR_AKTARIM, KOOR_PL4
)
from tools.dogruer_master_suite.engine.rpa_field_values import queue_item_from_record

try:
    import openpyxl
except ImportError:
    openpyxl = None


class RPAStationView(BaseView):
    """OTONOM RPA İSTASYONU - 3 ROBOTLU KONTROL & KALİBRASYON MERKEZİ (PyQt6)"""

    def __init__(self, parent=None, main_window=None):
        super().__init__(parent, main_window)
        self.rpa = RPAEngine()
        self.floating_panel: Optional[QDialog] = None
        self.active_tab_name = "BEYANNAME"
        self.excel_files: List[str] = []

        self._setup_rpa_callbacks()
        self._build_ui()

    def _setup_rpa_callbacks(self):
        self.rpa.on_status_change = self._on_rpa_status_change
        self.rpa.on_queue_update = self._on_rpa_queue_update
        self.rpa.on_step_progress = self._on_rpa_step_progress
        self.rpa.on_approval_needed = self._on_rpa_approval_needed
        self.rpa.on_finish = self._on_rpa_finish
        self.rpa.on_log = self._on_rpa_log

    # ------------------------------------------------------------------
    # UI SETUP
    # ------------------------------------------------------------------
    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(14, 10, 14, 10)
        layout.setSpacing(8)

        # 1. Üst Segment Seçici Barı
        top_bar = QFrame(self)
        top_bar.setStyleSheet("background-color: #0f172a; border-radius: 8px; padding: 4px;")
        tb_layout = QHBoxLayout(top_bar)
        tb_layout.setContentsMargins(10, 6, 10, 6)
        tb_layout.setSpacing(6)

        lbl_title = QLabel("🤖 OTONOM RPA İSTASYONU", top_bar)
        lbl_title.setStyleSheet("font-size: 13px; font-weight: 800; color: #38bdf8; border: none;")
        tb_layout.addWidget(lbl_title)

        self.tab_buttons: Dict[str, QPushButton] = {}
        tabs = [
            ("BEYANNAME", "📑 1. Beyanname Doldurma Robotu"),
            ("AKTARIM", "📊 2. Evrim Excel Aktarım Robotu (F8)"),
            ("PL4", "🏷️ 3. Toplu PL4 Ön İzin Robotu")
        ]
        for code, text in tabs:
            btn = QPushButton(text, top_bar)
            btn.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
            btn.clicked.connect(lambda ch, c=code: self.switch_robot_tab(c))
            tb_layout.addWidget(btn)
            self.tab_buttons[code] = btn

        tb_layout.addStretch()

        btn_probe = QPushButton("🔎 Evrim'i Algıla", top_bar)
        btn_probe.setStyleSheet("background-color: #0f766e; color: #ffffff; font-weight: 700; padding: 6px 12px; border-radius: 6px;")
        btn_probe.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn_probe.clicked.connect(self.probe_evrim_window)
        tb_layout.addWidget(btn_probe)

        btn_panic = QPushButton("🚨 ACİL DURDUR (ESC / F12)", top_bar)
        btn_panic.setStyleSheet("background-color: #dc2626; color: #ffffff; font-weight: 800; padding: 6px 14px; border-radius: 6px;")
        btn_panic.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn_panic.clicked.connect(self.trigger_emergency_stop)
        tb_layout.addWidget(btn_panic)

        layout.addWidget(top_bar)

        # 2. Robot Panelleri Container'ı (QStackedWidget)
        self.stack = QStackedWidget(self)

        # Panel 1: Beyanname
        self.panel_beyan = QWidget()
        self._build_beyanname_panel(self.panel_beyan)
        self.stack.addWidget(self.panel_beyan)

        # Panel 2: Aktarım
        self.panel_aktarim = QWidget()
        self._build_aktarim_panel(self.panel_aktarim)
        self.stack.addWidget(self.panel_aktarim)

        # Panel 3: PL4
        self.panel_pl4 = QWidget()
        self._build_pl4_panel(self.panel_pl4)
        self.stack.addWidget(self.panel_pl4)

        layout.addWidget(self.stack, stretch=1)

        # 3. Canlı İşlem Günlüğü (Konsol)
        log_frame = QFrame(self)
        log_frame.setStyleSheet("background-color: #111827; border: 1px solid #1e293b; border-radius: 8px; padding: 4px;")
        lf_layout = QVBoxLayout(log_frame)
        lf_layout.setContentsMargins(8, 4, 8, 4)
        lf_layout.setSpacing(4)

        lbl_log = QLabel("📜 Canlı Robot İşlem & Kalibrasyon Günlüğü:", log_frame)
        lbl_log.setStyleSheet("font-size: 11px; font-weight: 700; color: #38bdf8;")
        lf_layout.addWidget(lbl_log)

        self.txt_console = QTextEdit(log_frame)
        self.txt_console.setReadOnly(True)
        self.txt_console.setFixedHeight(95)
        self.txt_console.setStyleSheet("background-color: #0f172a; color: #4ade80; font-family: Consolas, monospace; font-size: 11px; border-radius: 6px; padding: 6px;")
        lf_layout.addWidget(self.txt_console)
        layout.addWidget(log_frame)

        # Kısayollar: ESC and F12
        shortcut_esc = QShortcut(QKeySequence(Qt.Key.Key_Escape), self)
        shortcut_esc.activated.connect(self.trigger_emergency_stop)
        shortcut_f12 = QShortcut(QKeySequence(Qt.Key.Key_F12), self)
        shortcut_f12.activated.connect(self.trigger_emergency_stop)

        self.switch_robot_tab("BEYANNAME")

    def switch_robot_tab(self, code: str):
        self.active_tab_name = code
        for c, btn in self.tab_buttons.items():
            if c == code:
                btn.setStyleSheet("background-color: #7c3aed; color: #ffffff; font-weight: 700; padding: 6px 12px; border-radius: 6px;")
            else:
                btn.setStyleSheet("background-color: #1e293b; color: #94a3b8; font-weight: 600; padding: 6px 12px; border-radius: 6px;")

        if code == "BEYANNAME":
            self.stack.setCurrentWidget(self.panel_beyan)
            self._log(f"📑 Beyanname Doldurma Robotu aktif. Profil: {self.rpa.active_beyan_profile}")
        elif code == "AKTARIM":
            self.stack.setCurrentWidget(self.panel_aktarim)
            self._log("📊 Evrim Excel Aktarım Robotu aktif. F8 tetiklemeli mod.")
            self.refresh_excel_folder()
        else:
            self.stack.setCurrentWidget(self.panel_pl4)
            self._log("🏷️ Toplu PL4 Ön İzin Ekleme Robotu aktif.")

    # ------------------------------------------------------------------
    # 1. BEYANNAME PANELİ
    # ------------------------------------------------------------------
    def _build_beyanname_panel(self, parent: QWidget):
        p_layout = QVBoxLayout(parent)
        p_layout.setContentsMargins(0, 0, 0, 0)
        p_layout.setSpacing(8)

        # Kartlar
        cards = QHBoxLayout()
        cards.setSpacing(8)
        self.c1 = MetricCard("ROBOT DURUMU", "HAZIR (BEKLEMEDE)", "Kuyruk işlem bekliyor", "#0284c7", parent)
        self.c2 = MetricCard("ÇALIŞMA MODU", "👁️ YARI OTONOM", "Space ile dosya onayı", "#7c3aed", parent)
        self.c3 = MetricCard("KUYRUK DURUMU", "0 / 0 Dosya", "Kuyruk boş", "#059669", parent)
        self.c4 = MetricCard("KALİBRASYON DURUMU", f"{len(self.rpa.koor_beyan)} / 29 Alan", "Koordinatlar hazır", "#d97706", parent)
        cards.addWidget(self.c1)
        cards.addWidget(self.c2)
        cards.addWidget(self.c3)
        cards.addWidget(self.c4)
        p_layout.addLayout(cards)

        # Kontrol Barı
        bar = QFrame(parent)
        bar.setStyleSheet("background-color: #111827; border: 1px solid #1e293b; border-radius: 8px; padding: 6px;")
        b_layout = QHBoxLayout(bar)
        b_layout.setContentsMargins(10, 4, 10, 4)
        b_layout.setSpacing(8)

        lbl_m = QLabel("Mod:", bar)
        lbl_m.setStyleSheet("font-weight: 700; color: #38bdf8;")
        b_layout.addWidget(lbl_m)

        self.rb_semi = QRadioButton("👁️ Yarı Otonom (Space)", bar)
        self.rb_semi.setStyleSheet("color: #e2e8f0; font-weight: 600;")
        self.rb_semi.setChecked(True)
        self.rb_semi.toggled.connect(self._on_beyan_mode_change)
        b_layout.addWidget(self.rb_semi)

        self.rb_full = QRadioButton("🚀 Tam Otonom", bar)
        self.rb_full.setStyleSheet("font-weight: 700; color: #a78bfa;")
        self.rb_full.toggled.connect(self._on_beyan_mode_change)
        b_layout.addWidget(self.rb_full)

        self.chk_beyan_test = QCheckBox("Güvenli Test Modu (Yazmaz / Tamam Basmaz)", bar)
        self.chk_beyan_test.setStyleSheet("color: #e2e8f0;")
        self.chk_beyan_test.setStyleSheet("font-weight: 700; color: #2563eb;")
        b_layout.addWidget(self.chk_beyan_test)

        lbl_p = QLabel("Profil:", bar)
        lbl_p.setStyleSheet("font-weight: 700; color: #94a3b8; background: transparent;")
        b_layout.addWidget(lbl_p)

        self.cb_beyan_prof = QComboBox(bar)
        self.cb_beyan_prof.addItems(["Profil 1 (Ana Ekran)", "Profil 2 (İkinci Monitör)"])
        self.cb_beyan_prof.currentTextChanged.connect(lambda txt: self.rpa.set_beyan_profile(txt))
        b_layout.addWidget(self.cb_beyan_prof)

        b_layout.addStretch()

        btn_calib = QPushButton("🎯 Kalibrasyon Başlat (F9)", bar)
        btn_calib.setStyleSheet("background-color: #ea580c; color: #ffffff; font-weight: 700; padding: 6px 10px; border-radius: 6px;")
        btn_calib.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn_calib.clicked.connect(lambda: self.start_calibration_for("BEYANNAME"))
        b_layout.addWidget(btn_calib)

        btn_test = QPushButton("👀 Test Et", bar)
        btn_test.setStyleSheet("background-color: #0284c7; color: #ffffff; font-weight: 600; padding: 6px 10px; border-radius: 6px;")
        btn_test.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn_test.clicked.connect(lambda: self.test_coordinates_for("BEYANNAME"))
        b_layout.addWidget(btn_test)

        self.btn_beyan_start = QPushButton("🚀 Robotu Başlat", bar)
        self.btn_beyan_start.setStyleSheet("background-color: #7c3aed; color: #ffffff; font-weight: 700; padding: 6px 14px; border-radius: 6px;")
        self.btn_beyan_start.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        self.btn_beyan_start.clicked.connect(self.start_beyanname_process)
        b_layout.addWidget(self.btn_beyan_start)

        self.btn_beyan_pause = QPushButton("⏸️ Duraklat", bar)
        self.btn_beyan_pause.setStyleSheet("background-color: #d97706; color: #ffffff; font-weight: 700; padding: 6px 12px; border-radius: 6px;")
        self.btn_beyan_pause.setEnabled(False)
        self.btn_beyan_pause.clicked.connect(self.rpa.pause)
        b_layout.addWidget(self.btn_beyan_pause)

        self.btn_beyan_resume = QPushButton("▶️ Devam Et", bar)
        self.btn_beyan_resume.setStyleSheet("background-color: #059669; color: #ffffff; font-weight: 700; padding: 6px 12px; border-radius: 6px;")
        self.btn_beyan_resume.setEnabled(False)
        self.btn_beyan_resume.clicked.connect(self.rpa.resume)
        b_layout.addWidget(self.btn_beyan_resume)

        p_layout.addWidget(bar)

        # Splitter: Kuyruk Tablosu vs Canlı Parametre Düzenleyici
        splitter = QSplitter(Qt.Orientation.Horizontal, parent)

        # Sol: Kuyruk Tablosu
        left_w = QFrame(splitter)
        left_w.setStyleSheet("background-color: #111827; border: 1px solid #1e293b; border-radius: 8px; padding: 6px;")
        lw_layout = QVBoxLayout(left_w)
        lw_layout.setContentsMargins(6, 6, 6, 6)

        q_head = QHBoxLayout()
        lbl_qh = QLabel("📋 İŞLENECEK DOSYA KUYRUĞU", left_w)
        lbl_qh.setStyleSheet("font-weight: 800; color: #38bdf8;")
        q_head.addWidget(lbl_qh)
        q_head.addStretch()

        btn_add_f = QPushButton("➕ Dosya Ekle", left_w)
        btn_add_f.clicked.connect(self.show_add_file_dialog)
        q_head.addWidget(btn_add_f)

        btn_excel_q = QPushButton("📂 Excel'den Al", left_w)
        btn_excel_q.clicked.connect(self.load_from_excel_file)
        q_head.addWidget(btn_excel_q)

        btn_clear_q = QPushButton("🧹 Temizle", left_w)
        btn_clear_q.setStyleSheet("color: #dc2626;")
        btn_clear_q.clicked.connect(self.clear_queue)
        q_head.addWidget(btn_clear_q)
        lw_layout.addLayout(q_head)

        self.tree_q = QTableWidget(left_w)
        self.tree_q.setColumnCount(8)
        self.tree_q.setHorizontalHeaderLabels([
            "#", "Dosya No", "Referans No", "Beyanname No",
            "Taşıt / Gemi", "Liman", "Durum", "Sonuç / Not"
        ])
        self.tree_q.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.tree_q.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.tree_q.verticalHeader().setVisible(False)
        self.tree_q.horizontalHeader().setStretchLastSection(True)
        lw_layout.addWidget(self.tree_q)
        splitter.addWidget(left_w)

        # Sağ: Parametre Düzenleyici
        right_w = QFrame(splitter)
        right_w.setStyleSheet("background-color: #111827; border: 1px solid #1e293b; border-radius: 8px; padding: 8px;")
        right_w.setFixedWidth(330)
        rw_layout = QVBoxLayout(right_w)
        rw_layout.setContentsMargins(8, 8, 8, 8)
        rw_layout.setSpacing(4)

        lbl_rw = QLabel("✏️ CANLI PARAMETRE DÜZENLEYİCİ", right_w)
        lbl_rw.setStyleSheet("font-weight: 800; color: #a78bfa;")
        rw_layout.addWidget(lbl_rw)

        lbl_rw_sub = QLabel("Kalan dosyaların Evrim değerlerini anında değiştirin:", right_w)
        lbl_rw_sub.setStyleSheet("font-size: 10px; color: #94a3b8;")
        rw_layout.addWidget(lbl_rw_sub)

        self.param_entries: Dict[str, QLineEdit] = {}
        fields = [
            ("tasit_kimligi", "Taşıt / Gemi:", "MAERSK EL BOSQUE"),
            ("teslim_sekli", "Teslim Şekli (20):", "CIP"),
            ("teslim_sehri", "Teslim Şehri:", "KOCAELİ"),
            ("liman_kutu", "Liman Kodu:", "DP"),
            ("diger_tutar", "Diğer Gider Tutarı:", "3000"),
            ("diger_aciklama_metin", "Diğer Açıklama:", "Y.ICI"),
            ("liman_gider_tutar", "Liman Gider Tutarı:", "5000"),
            ("liman_gider_aciklama_metin", "Liman Açıklama:", "."),
            ("ek_bilgi_metin", "44 No'lu Kutu Notu:", "Tareks ve Kons. damga vergisi yurt ici gidere beyan edildi."),
        ]

        for k, lbl_t, def_v in fields:
            row = QHBoxLayout()
            lbl_f = QLabel(lbl_t, right_w)
            lbl_f.setStyleSheet("font-size: 11px; font-weight: 600; color: #334155;")
            lbl_f.setFixedWidth(130)
            row.addWidget(lbl_f)
            ent = QLineEdit(right_w)
            ent.setText(def_v)
            ent.setStyleSheet("padding: 3px 6px; border: 1px solid #cbd5e1; border-radius: 4px;")
            row.addWidget(ent)
            self.param_entries[k] = ent
            rw_layout.addLayout(row)

        btn_app = QPushButton("💾 Parametreleri Kalan Dosyalara Uygula", right_w)
        btn_app.setStyleSheet("background-color: #0284c7; color: #ffffff; font-weight: 700; padding: 8px; border-radius: 6px; margin-top: 6px;")
        btn_app.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn_app.clicked.connect(self.apply_params_to_remaining)
        rw_layout.addWidget(btn_app)
        rw_layout.addStretch()

        splitter.addWidget(right_w)
        splitter.setStretchFactor(0, 3)
        splitter.setStretchFactor(1, 1)

        p_layout.addWidget(splitter, stretch=1)

    # ------------------------------------------------------------------
    # 2. EVRİM EXCEL AKTARIM PANELİ
    # ------------------------------------------------------------------
    def _build_aktarim_panel(self, parent: QWidget):
        p_layout = QVBoxLayout(parent)
        p_layout.setContentsMargins(0, 0, 0, 0)
        p_layout.setSpacing(8)

        cards = QHBoxLayout()
        cards.setSpacing(8)
        self.ak1 = MetricCard("AKTARIM STATÜSÜ", "BEKLEMEDE", "Excel dosyası hazır", "#0284c7", parent)
        self.ak2 = MetricCard("KLASÖRDEKİ EXCELLER", "0 Dosya", "Taranıyor...", "#059669", parent)
        self.ak3 = MetricCard("F8 MODU", "F8 İLE ONAYLI", "Her dosyada F8 bekler", "#7c3aed", parent)
        self.ak4 = MetricCard("KALİBRASYON", f"{len(self.rpa.koor_aktarim)} / 14 Alan", "aktarim_koordinatlar.json", "#ea580c", parent)
        cards.addWidget(self.ak1)
        cards.addWidget(self.ak2)
        cards.addWidget(self.ak3)
        cards.addWidget(self.ak4)
        p_layout.addLayout(cards)

        bar = QFrame(parent)
        bar.setStyleSheet("background-color: #111827; border: 1px solid #1e293b; border-radius: 8px; padding: 6px;")
        b_layout = QHBoxLayout(bar)
        b_layout.setContentsMargins(10, 4, 10, 4)
        b_layout.setSpacing(8)

        lbl_k = QLabel("📁 Excel Klasörü:", bar)
        lbl_k.setStyleSheet("font-weight: 700; color: #38bdf8;")
        b_layout.addWidget(lbl_k)

        self.ent_excel_klasor = QLineEdit(bar)
        default_k = r"C:\Users\ozan.gemici\OneDrive - DOGRUER GUMRUK MUSAVIRLIGI A.S\Belgeler\OTONOMAKTARIM"
        self.ent_excel_klasor.setText(default_k)
        self.ent_excel_klasor.setStyleSheet("padding: 4px 8px; border: 1px solid #334155; border-radius: 6px; background-color: #0b0f19; color: #f8fafc;")
        b_layout.addWidget(self.ent_excel_klasor, stretch=1)

        btn_browse_k = QPushButton("Seç...", bar)
        btn_browse_k.clicked.connect(self.select_excel_folder)
        b_layout.addWidget(btn_browse_k)

        btn_refresh_k = QPushButton("🔄 Yenile", bar)
        btn_refresh_k.clicked.connect(self.refresh_excel_folder)
        b_layout.addWidget(btn_refresh_k)

        self.chk_aktarim_f8 = QCheckBox("Her Dosyada [F8] Bekle", bar)
        self.chk_aktarim_f8.setChecked(True)
        self.chk_aktarim_f8.setStyleSheet("font-weight: 700;")
        b_layout.addWidget(self.chk_aktarim_f8)

        self.chk_aktarim_test = QCheckBox("Güvenli Test Modu (Aktar Basmaz)", bar)
        self.chk_aktarim_test.setStyleSheet("font-weight: 700; color: #2563eb;")
        b_layout.addWidget(self.chk_aktarim_test)

        b_layout.addStretch()

        btn_calib_ak = QPushButton("🎯 Kalibrasyon (F9)", bar)
        btn_calib_ak.setStyleSheet("background-color: #ea580c; color: #ffffff; font-weight: 700; padding: 6px 10px; border-radius: 6px;")
        btn_calib_ak.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn_calib_ak.clicked.connect(lambda: self.start_calibration_for("AKTARIM"))
        b_layout.addWidget(btn_calib_ak)

        btn_test_ak = QPushButton("👀 Test Et", bar)
        btn_test_ak.setStyleSheet("background-color: #0284c7; color: #ffffff; font-weight: 600; padding: 6px 10px; border-radius: 6px;")
        btn_test_ak.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn_test_ak.clicked.connect(lambda: self.test_coordinates_for("AKTARIM"))
        b_layout.addWidget(btn_test_ak)

        self.btn_aktarim_start = QPushButton("🚀 Aktarımı Başlat", bar)
        self.btn_aktarim_start.setStyleSheet("background-color: #059669; color: #ffffff; font-weight: 700; padding: 6px 14px; border-radius: 6px;")
        self.btn_aktarim_start.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        self.btn_aktarim_start.clicked.connect(self.start_excel_aktarim_process)
        b_layout.addWidget(self.btn_aktarim_start)

        p_layout.addWidget(bar)

        # Splitter (Excel Listesi vs Bilgiler)
        splitter = QSplitter(Qt.Orientation.Horizontal, parent)

        left_w = QFrame(splitter)
        left_w.setStyleSheet("background-color: #111827; border: 1px solid #1e293b; border-radius: 8px; padding: 8px;")
        lw_layout = QVBoxLayout(left_w)
        lbl_ex = QLabel("📑 KLASÖRDE BULUNAN EXCEL DOSYALARI (Seçip Başlatabilirsiniz):", left_w)
        lbl_ex.setStyleSheet("font-weight: 800; color: #38bdf8;")
        lw_layout.addWidget(lbl_ex)

        self.lb_excel = QListWidget(left_w)
        self.lb_excel.setStyleSheet("background-color: #0b0f19; color: #f8fafc; border: 1px solid #334155; border-radius: 6px; font-family: Consolas, monospace; font-size: 11px;")
        lw_layout.addWidget(self.lb_excel)
        splitter.addWidget(left_w)

        right_w = QFrame(splitter)
        right_w.setStyleSheet("background-color: #111827; border: 1px solid #1e293b; border-radius: 8px; padding: 12px;")
        right_w.setFixedWidth(340)
        rw_layout = QVBoxLayout(right_w)
        rw_layout.setSpacing(6)

        lbl_s = QLabel("⚙️ EVRİM EXCEL AKTARIM SABİTLERİ", right_w)
        lbl_s.setStyleSheet("font-weight: 800; color: #34d399;")
        rw_layout.addWidget(lbl_s)

        sabits = [
            ("Döviz Kodu:", "TL"),
            ("ATR Kodu:", "KORE"),
            ("Alıcı Kodu:", "9122"),
            ("Gümrük Kodu:", "18"),
            ("İşlem Tipi:", "Excelden Beyannameye"),
        ]
        for l_txt, v_txt in sabits:
            r = QHBoxLayout()
            lbl = QLabel(l_txt, right_w)
            lbl.setStyleSheet("font-weight: 700; color: #334155;")
            lbl.setFixedWidth(110)
            r.addWidget(lbl)
            val = QLabel(v_txt, right_w)
            val.setStyleSheet("background-color: #f1f5f9; padding: 4px 8px; border-radius: 4px; font-weight: 600;")
            r.addWidget(val)
            rw_layout.addLayout(r)

        lbl_info = QLabel(
            "\n💡 Çalışma Mantığı:\n"
            "1. Evrim'de Beyanname modülünü açın.\n"
            "2. Robotu Başlat'a basın.\n"
            "3. Sarı kutuya tıkladıktan sonra klavyeden [ F8 ] tuşuna basarak her dosyayı tek tıkla içeri aktarın.",
            right_w
        )
        lbl_info.setStyleSheet("font-size: 11px; color: #475569; background-color: #f8fafc; padding: 10px; border-radius: 6px;")
        rw_layout.addWidget(lbl_info)
        rw_layout.addStretch()

        splitter.addWidget(right_w)
        splitter.setStretchFactor(0, 3)
        splitter.setStretchFactor(1, 1)

        p_layout.addWidget(splitter, stretch=1)

    # ------------------------------------------------------------------
    # 3. PL4 PANELİ
    # ------------------------------------------------------------------
    def _build_pl4_panel(self, parent: QWidget):
        p_layout = QVBoxLayout(parent)
        p_layout.setContentsMargins(0, 0, 0, 0)
        p_layout.setSpacing(8)

        cards = QHBoxLayout()
        cards.setSpacing(8)
        self.pl1 = MetricCard("PL4 ROBOT DURUMU", "HAZIR", "Dosya listesi bekleniyor", "#0284c7", parent)
        self.pl2 = MetricCard("İŞLENECEK DOSYALAR", "0 Dosya", "Kuyruk boş", "#059669", parent)
        self.pl3 = MetricCard("KOD DEĞERİ", "PL4", "Ön İzin Kodu", "#7c3aed", parent)
        self.pl4 = MetricCard("KALİBRASYON", f"{len(self.rpa.koor_pl4)} / 5 Alan", "pl4_koordinatlar.json", "#ea580c", parent)
        cards.addWidget(self.pl1)
        cards.addWidget(self.pl2)
        cards.addWidget(self.pl3)
        cards.addWidget(self.pl4)
        p_layout.addLayout(cards)

        bar = QFrame(parent)
        bar.setStyleSheet("background-color: #111827; border: 1px solid #1e293b; border-radius: 8px; padding: 6px;")
        b_layout = QHBoxLayout(bar)
        b_layout.setContentsMargins(10, 4, 10, 4)
        b_layout.setSpacing(8)

        btn_add = QPushButton("➕ Manuel Dosya Ekle", bar)
        btn_add.clicked.connect(self.show_pl4_add_dialog)
        b_layout.addWidget(btn_add)

        btn_xl = QPushButton("📂 Excel'den Yükle", bar)
        btn_xl.clicked.connect(self.load_pl4_from_excel)
        b_layout.addWidget(btn_xl)

        btn_clr = QPushButton("🧹 Listeyi Temizle", bar)
        btn_clr.setStyleSheet("color: #dc2626;")
        btn_clr.clicked.connect(self.clear_pl4_queue)
        b_layout.addWidget(btn_clr)

        self.chk_pl4_test = QCheckBox("Güvenli Test Modu (PL4 Yazmaz)", bar)
        self.chk_pl4_test.setStyleSheet("font-weight: 700; color: #38bdf8;")
        b_layout.addWidget(self.chk_pl4_test)

        b_layout.addStretch()

        btn_calib_pl = QPushButton("🎯 Kalibrasyon (F9)", bar)
        btn_calib_pl.setStyleSheet("background-color: #ea580c; color: #ffffff; font-weight: 700; padding: 6px 10px; border-radius: 6px;")
        btn_calib_pl.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn_calib_pl.clicked.connect(lambda: self.start_calibration_for("PL4"))
        b_layout.addWidget(btn_calib_pl)

        btn_test_pl = QPushButton("👀 Test Et", bar)
        btn_test_pl.setStyleSheet("background-color: #0284c7; color: #ffffff; font-weight: 600; padding: 6px 10px; border-radius: 6px;")
        btn_test_pl.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn_test_pl.clicked.connect(lambda: self.test_coordinates_for("PL4"))
        b_layout.addWidget(btn_test_pl)

        self.btn_pl4_start = QPushButton("🚀 PL4 Robotunu Başlat", bar)
        self.btn_pl4_start.setStyleSheet("background-color: #7c3aed; color: #ffffff; font-weight: 700; padding: 6px 14px; border-radius: 6px;")
        self.btn_pl4_start.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        self.btn_pl4_start.clicked.connect(self.start_pl4_process)
        b_layout.addWidget(self.btn_pl4_start)

        p_layout.addWidget(bar)

        body = QFrame(parent)
        body.setStyleSheet("background-color: #111827; border: 1px solid #1e293b; border-radius: 8px; padding: 8px;")
        bd_layout = QVBoxLayout(body)
        lbl_pl_title = QLabel("🏷️ PL4 EKLENECEK DOSYA NUMARALARI LİSTESİ:", body)
        lbl_pl_title.setStyleSheet("font-weight: 800; color: #38bdf8;")
        bd_layout.addWidget(lbl_pl_title)

        self.lb_pl4 = QListWidget(body)
        self.lb_pl4.setStyleSheet("background-color: #0b0f19; color: #f8fafc; border: 1px solid #334155; border-radius: 6px; font-family: Consolas, monospace; font-size: 11px;")
        bd_layout.addWidget(self.lb_pl4)

        p_layout.addWidget(body, stretch=1)

    # ------------------------------------------------------------------
    # LOG & HELPERS
    # ------------------------------------------------------------------
    def _log(self, msg: str):
        t_str = time.strftime('%H:%M:%S')
        self.txt_console.append(f"[{t_str}] {msg}")
        self.txt_console.verticalScrollBar().setValue(self.txt_console.verticalScrollBar().maximum())

    def _on_rpa_log(self, msg: str):
        QTimer.singleShot(0, lambda: self._log(msg))

    def _refresh_calibration_cards(self):
        self.c4.set_value(f"{len(self.rpa.koor_beyan)} / 29 Alan")
        self.ak4.set_value(f"{len(self.rpa.koor_aktarim)} / 14 Alan")
        self.pl4.set_value(f"{len(self.rpa.koor_pl4)} / 5 Alan")

    # ------------------------------------------------------------------
    # CALIBRATION & PROBING
    # ------------------------------------------------------------------
    def probe_evrim_window(self):
        ok = self.rpa.probe_evrim()
        if ok:
            self.show_info("Evrim Algılandı", "Evrim penceresi bulundu. Ayrıntı robot günlüğündedir.")
        else:
            self.show_error(
                "Evrim Yok",
                "Evrim penceresi bulunamadı.\n\n"
                "ithalsql.exe / ihrsql.exe açık olsun.\n"
                "Simge durumundaysa görev çubuğundan geri yükleyin."
            )

    def start_calibration_for(self, robot_type: str):
        reply = QMessageBox.question(
            self,
            "Kalibrasyon Başlatılsın mı?",
            f"[{robot_type}] Robotu için kalibrasyon sihirbazı başlatılacak.\n\n"
            "Evrim İthalat/İhracat açık olsun.\n"
            "Her kutu için imleci üzerine götürüp [ F9 ] basın.\n"
            "İptal: ESC  |  Robot çalışırken acil durdurma: ESC veya F12.\n\nHazır mısınız?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if reply != QMessageBox.StandardButton.Yes:
            return

        def _step(alan_adi, idx, total):
            pass

        def _done(success):
            def _u():
                if success:
                    self.show_info("Başarılı", f"[{robot_type}] Kalibrasyonu tamamlandı ve kaydedildi!")
                    self.rpa.load_all_coordinates()
                    self._refresh_calibration_cards()
                else:
                    self.show_warning("İptal", "Kalibrasyon iptal edildi.")
            QTimer.singleShot(0, _u)

        self.rpa.run_calibration_wizard(robot_type, _step, _done)

    def test_coordinates_for(self, robot_type: str):
        self.rpa.test_robot_coordinates(robot_type, lambda s: self.show_info("Test Bitti", f"[{robot_type}] Testi tamamlandı."))

    # ------------------------------------------------------------------
    # BEYANNAME ROBOTU İŞLEMLERİ
    # ------------------------------------------------------------------
    def _on_beyan_mode_change(self):
        if self.rb_full.isChecked():
            self.rpa.mode = "FULL_AUTONOMOUS"
            self.c2.set_value("🚀 TAM OTONOM")
        else:
            self.rpa.mode = "SEMI_AUTONOMOUS"
            self.c2.set_value("👁️ YARI OTONOM")

    def apply_params_to_remaining(self):
        new_p = {k: ent.text().strip() for k, ent in self.param_entries.items()}
        self.rpa.update_remaining_params(new_p)
        self._refresh_tree_view()
        self.show_info("Başarılı", "Parametreler sıradaki tüm dosyalara uygulandı!")

    def add_external_files_to_queue(self, file_records: List[Dict[str, Any]]):
        current_list = list(self.rpa.queue)
        fb = self.param_entries["tasit_kimligi"].text().strip()
        for f in file_records:
            current_list.append(queue_item_from_record(f, fallback_gemi=fb))
        self.rpa.set_queue(current_list)
        self._refresh_tree_view()

    def show_add_file_dialog(self):
        dlg = QDialog(self)
        dlg.setWindowTitle("Dosya Numaraları Ekle")
        dlg.resize(400, 300)
        d_layout = QVBoxLayout(dlg)
        d_layout.addWidget(QLabel("Dosya Numaralarını Yapıştırın (Her satıra bir adet):", dlg))

        txt = QTextEdit(dlg)
        d_layout.addWidget(txt, stretch=1)

        b_box = QHBoxLayout()
        b_box.addStretch()
        btn_cancel = QPushButton("Vazgeç", dlg)
        btn_cancel.clicked.connect(dlg.reject)
        b_box.addWidget(btn_cancel)

        btn_add = QPushButton("Ekle", dlg)
        btn_add.setStyleSheet("background-color: #0284c7; color: #ffffff; font-weight: 700;")
        def _add():
            raw = txt.toPlainText().strip().splitlines()
            items = [{"dosya_no": l.strip(), "ref": "", "beyanname_no": ""} for l in raw if l.strip()]
            if items:
                self.add_external_files_to_queue(items)
            dlg.accept()
        btn_add.clicked.connect(_add)
        b_box.addWidget(btn_add)
        d_layout.addLayout(b_box)
        dlg.exec()

    def load_from_excel_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Excel Seç", "", "Excel (*.xlsx *.xls)")
        if not path or not openpyxl:
            return
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            items = []
            rows = list(ws.iter_rows(values_only=True))
            start = 0
            header_map = {}
            if rows:
                headers = [str(c or "").strip().lower() for c in rows[0]]
                if any(h in headers for h in ("dosya", "dosya no", "dosya_no", "ref", "gemi", "liman")):
                    start = 1
                    for i, h in enumerate(headers):
                        header_map[h] = i
            for row in rows[start:]:
                if not row or not row[0]:
                    continue
                v = str(row[0]).strip()
                if v.lower() in ("dosya", "dosya no", "dosya_no"):
                    continue
                rec = {"dosya_no": v}
                def col(*names, idx=None):
                    for n in names:
                        if n in header_map and header_map[n] < len(row) and row[header_map[n]] is not None:
                            return str(row[header_map[n]]).strip()
                    if idx is not None and idx < len(row) and row[idx] is not None:
                        return str(row[idx]).strip()
                    return ""
                rec["ref"] = col("ref", "hankook_ref", idx=1)
                rec["beyanname_no"] = col("beyanname", "beyanname_no", "beyan", idx=2)
                rec["gemi"] = col("gemi", "tasit", idx=3)
                rec["liman"] = col("liman", idx=4)
                rec["fatura_bedeli"] = col("fatura", "fatura_bedeli", "vergi")
                items.append(rec)
            if items:
                self.rpa.set_queue(items)
                self._refresh_tree_view()
                self.show_info("Başarılı", f"{len(items)} dosya kuyruğa alındı!")
        except Exception as e:
            self.show_error("Hata", f"Excel okunamadı: {e}")

    def clear_queue(self):
        self.rpa.set_queue([])
        self._refresh_tree_view()

    def _refresh_tree_view(self):
        total = len(self.rpa.queue)
        comp = len([x for x in self.rpa.queue if x["status"] == "TAMAMLANDI"])
        self.c3.set_value(f"{comp} / {total} Dosya")

        self.tree_q.setRowCount(total)
        for idx, item in enumerate(self.rpa.queue):
            status = item.get("status", "BEKLIYOR")
            values = [
                str(idx + 1), item.get("dosya_no", ""), item.get("ref", "—"),
                item.get("beyanname_no", "—"), item.get("gemi", ""), item.get("liman", ""),
                status, item.get("message", "")
            ]

            if status == "TAMAMLANDI":
                bg = QColor("#dcfce7"); fg = QColor("#15803d")
            elif status == "YAZILIYOR":
                bg = QColor("#fef3c7"); fg = QColor("#b45309")
            elif status == "ONAY_BEKLIYOR":
                bg = QColor("#ede9fe"); fg = QColor("#6d28d9")
            elif status in ("DURDURULDU", "HATA"):
                bg = QColor("#fee2e2"); fg = QColor("#b91c1c")
            else:
                bg = QColor("#ffffff"); fg = QColor("#1e293b")

            for col, val in enumerate(values):
                cell = QTableWidgetItem(str(val))
                cell.setBackground(bg)
                cell.setForeground(fg)
                if col in [0, 1, 2, 3, 5, 6]:
                    cell.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.tree_q.setItem(idx, col, cell)

    def start_beyanname_process(self):
        if not self.rpa.queue:
            self.show_warning("Uyarı", "Kuyrukta dosya yok!")
            return
        new_p = {k: ent.text().strip() for k, ent in self.param_entries.items()}
        self.rpa.active_params.update(new_p)
        self.btn_beyan_start.setEnabled(False)
        self.btn_beyan_pause.setEnabled(True)
        self._show_floating_panic_button()
        self.rpa.start_beyanname_process(dry_run=self.chk_beyan_test.isChecked())

    # ------------------------------------------------------------------
    # EXCEL AKTARIM İŞLEMLERİ
    # ------------------------------------------------------------------
    def select_excel_folder(self):
        k = QFileDialog.getExistingDirectory(self, "Excel Klasörü Seç", self.ent_excel_klasor.text())
        if k:
            self.ent_excel_klasor.setText(k)
            self.refresh_excel_folder()

    def refresh_excel_folder(self):
        k = self.ent_excel_klasor.text().strip()
        self.lb_excel.clear()
        self.excel_files = []
        if not os.path.exists(k):
            self._log(f"Excel klasörü bulunamadı: {k}")
            self.ak2.set_value("0 Dosya")
            return

        files = sorted([f for f in os.listdir(k) if f.lower().endswith(".xlsx") and not f.startswith("~")])
        self.excel_files = [os.path.join(k, f) for f in files]
        for idx, f in enumerate(files, 1):
            self.lb_excel.addItem(f"  {idx:>3}.  {f}")

        self.ak2.set_value(f"{len(files)} Dosya")
        self._log(f"Excel klasörü tarandı: {len(files)} dosya bulundu.")

    def start_excel_aktarim_process(self):
        if not self.excel_files:
            self.show_warning("Uyarı", "Klasörde .xlsx dosyası bulunamadı!")
            return

        sel_row = self.lb_excel.currentRow()
        start_i = sel_row if sel_row >= 0 else 0
        to_process = self.excel_files[start_i:]

        test_mode = self.chk_aktarim_test.isChecked()
        wait_f8 = self.chk_aktarim_f8.isChecked()

        self._show_floating_panic_button()
        self.rpa.start_aktarim_process(to_process, test_modu=test_mode, f8_bekle=wait_f8)

    # ------------------------------------------------------------------
    # PL4 İŞLEMLERİ
    # ------------------------------------------------------------------
    def show_pl4_add_dialog(self):
        dlg = QDialog(self)
        dlg.setWindowTitle("PL4 Dosya Numaraları Ekle")
        dlg.resize(400, 300)
        d_layout = QVBoxLayout(dlg)
        d_layout.addWidget(QLabel("PL4 Eklenecek Dosya Numaraları:", dlg))

        txt = QTextEdit(dlg)
        d_layout.addWidget(txt, stretch=1)

        b_box = QHBoxLayout()
        b_box.addStretch()
        btn_cancel = QPushButton("Vazgeç", dlg)
        btn_cancel.clicked.connect(dlg.reject)
        b_box.addWidget(btn_cancel)

        btn_add = QPushButton("Ekle", dlg)
        btn_add.setStyleSheet("background-color: #7c3aed; color: #ffffff; font-weight: 700;")
        def _add():
            raw = txt.toPlainText().strip().splitlines()
            for l in raw:
                v = l.strip()
                if v:
                    self.lb_pl4.addItem(v)
            self.pl2.set_value(f"{self.lb_pl4.count()} Dosya")
            dlg.accept()
        btn_add.clicked.connect(_add)
        b_box.addWidget(btn_add)
        d_layout.addLayout(b_box)
        dlg.exec()

    def load_pl4_from_excel(self):
        path, _ = QFileDialog.getOpenFileName(self, "Excel Seç", "", "Excel (*.xlsx *.xls)")
        if not path or not openpyxl:
            return
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            cnt = 0
            for row in ws.iter_rows(values_only=True):
                if row and row[0]:
                    v = str(row[0]).strip()
                    if v.lower() not in ("dosya", "dosya no", "dosya_no"):
                        self.lb_pl4.addItem(v)
                        cnt += 1
            self.pl2.set_value(f"{self.lb_pl4.count()} Dosya")
            self.show_info("Başarılı", f"{cnt} dosya PL4 listesine eklendi!")
        except Exception as e:
            self.show_error("Hata", f"Excel okunamadı: {e}")

    def clear_pl4_queue(self):
        self.lb_pl4.clear()
        self.pl2.set_value("0 Dosya")

    def start_pl4_process(self):
        items = [self.lb_pl4.item(i).text() for i in range(self.lb_pl4.count())]
        if not items:
            self.show_warning("Uyarı", "PL4 kuyruğunda dosya yok!")
            return
        self._show_floating_panic_button()
        self.rpa.start_pl4_process(items, test_modu=self.chk_pl4_test.isChecked())

    # ------------------------------------------------------------------
    # YÜZEN PANİK / ONAY PENCERESİ
    # ------------------------------------------------------------------
    def _show_floating_panic_button(self):
        if self.floating_panel:
            return
        self.floating_panel = QDialog(self, Qt.WindowType.WindowStaysOnTopHint | Qt.WindowType.FramelessWindowHint)
        self.floating_panel.resize(320, 90)
        self.floating_panel.move(60, 60)
        self.floating_panel.setStyleSheet("background-color: #0f172a; border: 2px solid #38bdf8; border-radius: 8px;")

        fl_layout = QVBoxLayout(self.floating_panel)
        fl_layout.setContentsMargins(10, 6, 10, 8)

        lbl = QLabel(f"🤖 {self.rpa.active_robot} ROBOTU AKTİF", self.floating_panel)
        lbl.setStyleSheet("color: #38bdf8; font-weight: 800; font-size: 11px; border: none;")
        fl_layout.addWidget(lbl)

        btn_row = QHBoxLayout()
        btn_ok = QPushButton("✅ ONAYLA (Space)", self.floating_panel)
        btn_ok.setStyleSheet("background-color: #7c3aed; color: #ffffff; font-weight: 700; padding: 8px; border-radius: 6px;")
        btn_ok.clicked.connect(self.rpa.approve_and_continue)
        btn_row.addWidget(btn_ok)

        btn_stop = QPushButton("🚨 DURDUR (ESC)", self.floating_panel)
        btn_stop.setStyleSheet("background-color: #dc2626; color: #ffffff; font-weight: 700; padding: 8px; border-radius: 6px;")
        btn_stop.clicked.connect(self.trigger_emergency_stop)
        btn_row.addWidget(btn_stop)

        fl_layout.addLayout(btn_row)
        self.floating_panel.show()

    def _hide_floating_panic_button(self):
        if self.floating_panel:
            try:
                self.floating_panel.close()
            except Exception:
                pass
            self.floating_panel = None

    def trigger_emergency_stop(self):
        self.rpa.trigger_emergency_stop()
        self._hide_floating_panic_button()
        self.btn_beyan_start.setEnabled(True)
        self.btn_beyan_pause.setEnabled(False)
        self.btn_beyan_resume.setEnabled(False)

    # ------------------------------------------------------------------
    # RPA ENGINE CALLBACK HANDLERS
    # ------------------------------------------------------------------
    def _on_rpa_status_change(self, status: str, msg: str):
        def _u():
            if status == "RUNNING":
                self.c1.set_value("⚙️ ÇALIŞIYOR")
            elif status == "WAITING_APPROVAL":
                self.c1.set_value("👁️ ONAY BEKLİYOR")
            elif status == "PAUSED":
                self.c1.set_value("⏸️ DURAKLATILDI")
            elif status == "STOPPED":
                self.c1.set_value("🛑 DURDURULDU")
            else:
                self.c1.set_value("HAZIR")
            self._log(msg)
        QTimer.singleShot(0, _u)

    def _on_rpa_queue_update(self, idx: int, item: Dict[str, Any]):
        QTimer.singleShot(0, self._refresh_tree_view)

    def _on_rpa_step_progress(self, alan_adi: str, cur: int, tot: int):
        pass

    def _on_rpa_approval_needed(self, item: Dict[str, Any]):
        pass

    def _on_rpa_finish(self, success: bool, msg: str):
        def _u():
            self._hide_floating_panic_button()
            self.btn_beyan_start.setEnabled(True)
            self.btn_beyan_pause.setEnabled(False)
            self.btn_beyan_resume.setEnabled(False)
            if success:
                self.show_info("Tamamlandı", msg)
            else:
                self.show_warning("Bilgi", msg)
        QTimer.singleShot(0, _u)
