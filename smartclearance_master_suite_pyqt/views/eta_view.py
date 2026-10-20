# -*- coding: utf-8 -*-
"""
eta_view.py — ETA Tarihi Yaklaşan Dosyalar & Varış Radarı (PyQt6)
Excel / Evrim iş takip tarihleri üzerinden yerel varış radarı.
Aciliyet renk rozetleri (kırmızı/kehribar/yeşil), filtreler, kokpit entegrasyonu ve Excel export.
"""

import os
import re
import datetime
from datetime import datetime as dt, date
from typing import Optional, List, Dict, Any

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit,
    QPushButton, QTableWidget, QTableWidgetItem, QHeaderView,
    QFrame, QMessageBox, QFileDialog, QMenu, QAbstractItemView
)
from PyQt6.QtCore import Qt, QTimer
from PyQt6.QtGui import QColor, QFont, QCursor, QShortcut, QKeySequence

try:
    import openpyxl
    import openpyxl.styles
    import openpyxl.utils
except ImportError:
    openpyxl = None

from .base_view import BaseView


class ETAView(BaseView):
    """ETA Tarihi Yaklaşan Dosyalar Sekmesi (PyQt6)."""

    def __init__(self, parent=None, main_window=None):
        super().__init__(parent, main_window)
        self.eta_filter_days = 7
        self._build_ui()

    @property
    def engine(self):
        if self.main_window and hasattr(self.main_window, "engine"):
            return self.main_window.engine
        return None

    @property
    def evrim_records(self):
        if self.main_window and hasattr(self.main_window, "evrim_records"):
            return self.main_window.evrim_records
        return []

    @property
    def kantar_tracker(self):
        if self.main_window and hasattr(self.main_window, "kantar_tracker"):
            return self.main_window.kantar_tracker
        return None

    # ------------------------------------------------------------------
    # UI SETUP
    # ------------------------------------------------------------------
    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(14, 10, 14, 10)
        layout.setSpacing(10)

        # 1. Üst Banner
        banner = QFrame(self)
        banner.setStyleSheet("background-color: #e0f2fe; border: 1px solid #bae6fd; border-radius: 10px; padding: 8px;")
        b_layout = QHBoxLayout(banner)
        b_layout.setContentsMargins(16, 10, 16, 10)

        b_left = QVBoxLayout()
        b_left.setSpacing(4)
        lbl_b_title = QLabel("🚢 ETA Tarihi Yaklaşan Dosyalar", banner)
        lbl_b_title.setStyleSheet("font-size: 16px; font-weight: 800; color: #0369a1; border: none;")
        b_left.addWidget(lbl_b_title)

        self.lbl_b_sub = QLabel(
            "Önümüzdeki 7 gün içerisinde varış yapması beklenen tüm dosyalarınız listelenmektedir. Kaynak: Excel ETA / varış tarihi.",
            banner
        )
        self.lbl_b_sub.setStyleSheet("font-size: 11px; color: #0284c7; border: none;")
        b_left.addWidget(self.lbl_b_sub)
        b_layout.addLayout(b_left, stretch=1)

        b_right = QHBoxLayout()
        b_right.setSpacing(10)

        self.btn_add_kantar = QPushButton("⚖️ Seçilenleri Kantar Takibe Aktar", banner)
        self.btn_add_kantar.setStyleSheet("""
            QPushButton {
                background-color: #2563eb;
                color: #ffffff;
                font-weight: 700;
                font-size: 11.5px;
                padding: 8px 16px;
                border-radius: 6px;
                border: none;
            }
            QPushButton:hover {
                background-color: #1d4ed8;
            }
            QPushButton:pressed {
                background-color: #1e40af;
            }
        """)
        self.btn_add_kantar.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        self.btn_add_kantar.setToolTip(
            "Seçili satırlardaki konteynerleri Kantar Takip havuzuna aktarır.\n"
            "• Tek satır seçebilirsiniz\n"
            "• Ctrl + Tık ile istediğiniz satırları seçebilirsiniz\n"
            "• Ctrl + A ile tüm satırları seçebilirsiniz"
        )
        self.btn_add_kantar.clicked.connect(self.transfer_selected_to_kantar_tracking)
        b_right.addWidget(self.btn_add_kantar)

        btn_excel_export = QPushButton("📊 Excel İndir", banner)
        btn_excel_export.setStyleSheet("background-color: #10b981; color: #ffffff; font-weight: 700; padding: 8px 16px; border-radius: 6px; border: none;")
        btn_excel_export.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn_excel_export.clicked.connect(self.export_eta_yaklasanlar_excel)
        b_right.addWidget(btn_excel_export)

        badge_box = QVBoxLayout()
        badge_box.setSpacing(0)
        badge_box.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_eta_yaklasan_count = QLabel("0", banner)
        self.lbl_eta_yaklasan_count.setStyleSheet("font-size: 22px; font-weight: 900; color: #0369a1; border: none;")
        self.lbl_eta_yaklasan_count.setAlignment(Qt.AlignmentFlag.AlignCenter)
        badge_box.addWidget(self.lbl_eta_yaklasan_count)

        lbl_badge_title = QLabel("YAKLAŞAN YÜK", banner)
        lbl_badge_title.setStyleSheet("font-size: 9px; font-weight: 800; color: #0284c7; border: none;")
        lbl_badge_title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        badge_box.addWidget(lbl_badge_title)
        b_right.addLayout(badge_box)

        b_layout.addLayout(b_right)
        layout.addWidget(banner)

        # 2. Kontrol Çubuğu & Filtreler
        ctrl_bar = QFrame(self)
        ctrl_bar.setStyleSheet("background-color: #111827; border: 1px solid #1e293b; border-radius: 8px; padding: 4px;")
        cb_layout = QHBoxLayout(ctrl_bar)
        cb_layout.setContentsMargins(10, 6, 10, 6)
        cb_layout.setSpacing(8)

        lbl_filter = QLabel("Filtre:", ctrl_bar)
        lbl_filter.setStyleSheet("font-weight: 700; color: #38bdf8; border: none;")
        cb_layout.addWidget(lbl_filter)

        self.btn_f_7d = QPushButton("⚡ Önümüzdeki 7 Gün", ctrl_bar)
        self.btn_f_7d.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        self.btn_f_7d.clicked.connect(lambda: self.set_eta_days_filter(7))
        cb_layout.addWidget(self.btn_f_7d)

        self.btn_f_14d = QPushButton("📅 14 Gün", ctrl_bar)
        self.btn_f_14d.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        self.btn_f_14d.clicked.connect(lambda: self.set_eta_days_filter(14))
        cb_layout.addWidget(self.btn_f_14d)

        self.btn_f_delayed = QPushButton("🔴 Gecikenler", ctrl_bar)
        self.btn_f_delayed.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        self.btn_f_delayed.clicked.connect(lambda: self.set_eta_days_filter(-1))
        cb_layout.addWidget(self.btn_f_delayed)

        self.btn_f_all = QPushButton("🌐 Tümü", ctrl_bar)
        self.btn_f_all.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        self.btn_f_all.clicked.connect(lambda: self.set_eta_days_filter(999))
        cb_layout.addWidget(self.btn_f_all)

        self._update_filter_button_styles()

        cb_layout.addStretch()

        lbl_s = QLabel("🔍 Ara:", ctrl_bar)
        lbl_s.setStyleSheet("font-weight: 700; color: #94a3b8; border: none; background: transparent;")
        cb_layout.addWidget(lbl_s)

        self.txt_eta_search = QLineEdit(ctrl_bar)
        self.txt_eta_search.setPlaceholderText("Gemi, Dosya, Ref, Liman...")
        self.txt_eta_search.setFixedWidth(220)
        self.txt_eta_search.setStyleSheet("background-color: #1e293b; color: #f8fafc; padding: 4px 8px; border: 1px solid #334155; border-radius: 6px;")
        self.txt_eta_search.textChanged.connect(self.refresh_eta_yaklasanlar_table)
        cb_layout.addWidget(self.txt_eta_search)

        layout.addWidget(ctrl_bar)

        # 3. Tablo
        self.tree_eta = QTableWidget(self)
        self.tree_eta.setColumnCount(9)
        self.tree_eta.setHorizontalHeaderLabels([
            "DOSYA NO", "REFERANS NO", "ACENTA", "LİMAN",
            "GEMİ", "KONTEYNER / KONŞ.", "ETA TARİHİ", "KALAN GÜN", "İŞLEM"
        ])
        self.tree_eta.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.tree_eta.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.tree_eta.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.tree_eta.setAlternatingRowColors(True)
        self.tree_eta.verticalHeader().setVisible(False)
        self.tree_eta.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.tree_eta.customContextMenuRequested.connect(self.show_eta_context_menu)
        self.tree_eta.cellDoubleClicked.connect(lambda r, c: self.on_eta_row_double_click(r))
        self.shortcut_select_all = QShortcut(QKeySequence(QKeySequence.StandardKey.SelectAll), self.tree_eta)
        self.shortcut_select_all.activated.connect(self.tree_eta.selectAll)

        header = self.tree_eta.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Interactive)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.Interactive)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.Interactive)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.Interactive)
        header.setSectionResizeMode(4, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(5, QHeaderView.ResizeMode.Interactive)
        header.setSectionResizeMode(6, QHeaderView.ResizeMode.Interactive)
        header.setSectionResizeMode(7, QHeaderView.ResizeMode.Interactive)
        header.setSectionResizeMode(8, QHeaderView.ResizeMode.Interactive)

        self.tree_eta.setColumnWidth(0, 100)
        self.tree_eta.setColumnWidth(1, 130)
        self.tree_eta.setColumnWidth(2, 140)
        self.tree_eta.setColumnWidth(3, 110)
        self.tree_eta.setColumnWidth(5, 160)
        self.tree_eta.setColumnWidth(6, 110)
        self.tree_eta.setColumnWidth(7, 110)
        self.tree_eta.setColumnWidth(8, 90)

        layout.addWidget(self.tree_eta, stretch=1)

    # ------------------------------------------------------------------
    # FILTERS & STYLES
    # ------------------------------------------------------------------
    def set_eta_days_filter(self, days: int):
        self.eta_filter_days = days
        self._update_filter_button_styles()

        if days == 7:
            self.lbl_b_sub.setText("Önümüzdeki 7 gün içerisinde varış yapması beklenen tüm açık dosyalarınız listelenmektedir.")
        elif days == 14:
            self.lbl_b_sub.setText("Önümüzdeki 14 gün içerisinde varış yapması beklenen tüm açık dosyalarınız listelenmektedir.")
        elif days == -1:
            self.lbl_b_sub.setText("ETA tarihi geçmiş veya geciken tüm açık dosyalarınız listelenmektedir.")
        else:
            self.lbl_b_sub.setText("Sistemdeki tüm ETA tarihli açık dosyalarınız listelenmektedir.")

        self.refresh_eta_yaklasanlar_table()

    def _update_filter_button_styles(self):
        active_style = "background-color: #0284c7; color: #ffffff; font-weight: 700; padding: 4px 10px; border-radius: 6px;"
        inactive_style = "background-color: #f1f5f9; color: #334155; font-weight: 500; padding: 4px 10px; border-radius: 6px;"
        delayed_active = "background-color: #dc2626; color: #ffffff; font-weight: 700; padding: 4px 10px; border-radius: 6px;"

        self.btn_f_7d.setStyleSheet(active_style if self.eta_filter_days == 7 else inactive_style)
        self.btn_f_14d.setStyleSheet(active_style if self.eta_filter_days == 14 else inactive_style)
        self.btn_f_delayed.setStyleSheet(delayed_active if self.eta_filter_days == -1 else inactive_style)
        self.btn_f_all.setStyleSheet(active_style if self.eta_filter_days == 999 else inactive_style)

    # ------------------------------------------------------------------
    # DATA REFRESH
    # ------------------------------------------------------------------
    def on_activate(self):
        self.refresh_eta_yaklasanlar_table()

    def refresh_table(self):
        self.refresh_eta_yaklasanlar_table()

    def on_eta_status_changed(self):
        QTimer.singleShot(0, self.refresh_eta_yaklasanlar_table)

    def refresh_eta_yaklasanlar_table(self):
        today = dt.now().date()
        search_q = self.txt_eta_search.text().strip().upper()

        if self.engine and not getattr(self.engine, "records", None):
            if hasattr(self.engine, "fetch_excel_records"):
                self.engine.fetch_excel_records()

        engine_records = getattr(self.engine, "records", []) if self.engine else []
        pool_records = [r for r in engine_records if not r.get("kapanis")] if engine_records else []

        # Merge with Evrim records
        for ev in self.evrim_records:
            d_no = ev.get("dosya_no")
            if d_no and not any(r.get("dosya") == d_no for r in pool_records):
                statu = str(ev.get("statu") or "").lower()
                if "kapan" not in statu and "teslim" not in statu:
                    pool_records.append({
                        "dosya": d_no,
                        "ref": ev.get("hankook_ref") or ev.get("musteri_referans") or "—",
                        "acenta": ev.get("acenta") or ev.get("nakliyeci") or "—",
                        "liman": ev.get("giris_gumruk") or ev.get("liman") or "EVYAP",
                        "gemi": ev.get("gemi_adi") or ev.get("gemi") or "—",
                        "konteyner": ev.get("konteyner_no") or ev.get("konteyner") or "—",
                        "bl": ev.get("bl_no") or ev.get("bl") or "—",
                        "eta": ev.get("eta") or ev.get("varis_tarihi") or ""
                    })

        approaching_items = []
        for rec in pool_records:
            eta_raw = str(rec.get("eta") or "").strip()
            if not eta_raw:
                continue

            clean_str = eta_raw.replace("T", " ").split(" ")[0].strip()
            eta_dt = None
            for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y.%m.%d"):
                try:
                    eta_dt = dt.strptime(clean_str[:10], fmt).date()
                    break
                except Exception:
                    pass

            if not eta_dt:
                continue

            fark_gun = (eta_dt - today).days
            dosya = str(rec.get("dosya") or "—")
            ref = str(rec.get("ref") or "—")
            acenta = str(rec.get("acenta") or "—")
            liman = str(rec.get("liman") or "EVYAP")
            gemi = str(rec.get("gemi") or "—")
            cntr_bl = str(rec.get("konteyner") or rec.get("bl") or "—")
            eta_formatted = eta_dt.strftime("%d.%m.%Y")

            if fark_gun == 0:
                kalan_txt = "🟢 Bugün"
                tag = "green_today"
            elif fark_gun > 0:
                kalan_txt = f"🔴 {fark_gun} Gün" if fark_gun <= 2 else f"🟡 {fark_gun} Gün"
                tag = "red_urgent" if fark_gun <= 2 else ("yellow_soon" if fark_gun <= 5 else "normal")
            else:
                kalan_txt = f"🔴 Gecikti ({abs(fark_gun)}g)"
                tag = "red_urgent"

            if self.eta_filter_days == 7 and not (0 <= fark_gun <= 7):
                continue
            elif self.eta_filter_days == 14 and not (0 <= fark_gun <= 14):
                continue
            elif self.eta_filter_days == -1 and fark_gun >= 0:
                continue

            if search_q:
                haystack = f"{dosya} {ref} {acenta} {liman} {gemi} {cntr_bl} {eta_formatted}".upper()
                if search_q not in haystack:
                    continue

            approaching_items.append({
                "dosya": dosya,
                "ref": ref,
                "acenta": acenta,
                "liman": liman,
                "gemi": gemi,
                "konteyner_bl": cntr_bl,
                "eta": eta_formatted,
                "kalan_gun": kalan_txt,
                "fark_gun": fark_gun,
                "tag": tag
            })

        approaching_items.sort(key=lambda x: x["fark_gun"])
        self.lbl_eta_yaklasan_count.setText(str(len(approaching_items)))

        # Populate Table
        self.tree_eta.setRowCount(len(approaching_items))
        for row_idx, item in enumerate(approaching_items):
            tag = item["tag"]
            if tag == "green_today":
                bg_color = QColor("#f0fdf4")
                fg_color = QColor("#15803d")
            elif tag == "red_urgent":
                bg_color = QColor("#fff1f2")
                fg_color = QColor("#be123c")
            elif tag == "yellow_soon":
                bg_color = QColor("#fefce8")
                fg_color = QColor("#a16207")
            else:
                bg_color = QColor("#ffffff")
                fg_color = QColor("#0f172a")

            values = [
                item["dosya"], item["ref"], item["acenta"], item["liman"],
                item["gemi"], item["konteyner_bl"], item["eta"], item["kalan_gun"], "Detay ➡️"
            ]

            for col_idx, val in enumerate(values):
                cell = QTableWidgetItem(str(val))
                cell.setBackground(bg_color)
                cell.setForeground(fg_color)
                if col_idx in [0, 1, 3, 6, 7, 8]:
                    cell.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                else:
                    cell.setTextAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)

                if col_idx in [0, 1, 7]:
                    f = cell.font()
                    f.setBold(True)
                    cell.setFont(f)

                self.tree_eta.setItem(row_idx, col_idx, cell)

    # ------------------------------------------------------------------
    # DOUBLE CLICK & ACTIONS
    # ------------------------------------------------------------------
    def on_eta_row_double_click(self, row: int):
        d_item = self.tree_eta.item(row, 0)
        r_item = self.tree_eta.item(row, 1)
        dosya = d_item.text().strip() if d_item else ""
        ref = r_item.text().strip() if r_item else ""

        # Switch to Cockpit
        self.switch_tab("COCKPIT")
        cockpit_view = self.main_window.views.get("COCKPIT") if self.main_window else None
        if cockpit_view:
            search_term = ref if ref and ref != "—" else dosya
            if hasattr(cockpit_view, "txt_search"):
                cockpit_view.txt_search.setText(search_term)
            if hasattr(cockpit_view, "_filter_table"):
                cockpit_view._filter_table()

            # Select matching row
            table = getattr(cockpit_view, "table", None)
            if table:
                for r in range(table.rowCount()):
                    row_dosya = table.item(r, 0).text() if table.item(r, 0) else ""
                    row_ref = table.item(r, 1).text() if table.item(r, 1) else ""
                    if (dosya and dosya in row_dosya) or (ref and ref in row_ref):
                        table.selectRow(r)
                        if hasattr(cockpit_view, "open_detail_drawer"):
                            cockpit_view.open_detail_drawer()
                        break

    # ------------------------------------------------------------------
    # KANTAR TRACKING BATCH TRANSFER
    # ------------------------------------------------------------------
    def transfer_selected_to_kantar_tracking(self):
        """Seçili tekli veya çoklu (Ctrl+Tık, Ctrl+A) konteynerleri Kantar Takip havuzuna ekler."""
        selected_rows = sorted(set(idx.row() for idx in self.tree_eta.selectedIndexes()))
        if not selected_rows:
            self.show_warning(
                "Seçim Yapılmadı",
                "Lütfen Kantar Takip havuzuna aktarmak istediğiniz en az bir satır seçin.\n\n"
                "💡 İpuçları:\n"
                "• Tek satır seçmek için: Satıra tıklayın.\n"
                "• Birden fazla satır seçmek için: Ctrl tuşuna basılı tutarak istediğiniz satırlara tıklayın.\n"
                "• Tüm satırları seçmek için: Ctrl + A tuşlarına basın."
            )
            return

        if not self.kantar_tracker:
            self.show_warning("Kantar Servisi", "Kantar takip servisi aktif değil.")
            return

        added_cntrs = []
        updated_cntrs = []
        skipped_rows = []

        for row in selected_rows:
            d_item = self.tree_eta.item(row, 0)
            r_item = self.tree_eta.item(row, 1)
            c_item = self.tree_eta.item(row, 5)

            dosya = d_item.text().strip() if d_item else "—"
            ref = r_item.text().strip() if r_item else "—"
            raw_cntr = c_item.text().strip() if c_item else ""

            if not raw_cntr or raw_cntr in ("—", "-", "None", "null"):
                skipped_rows.append(dosya)
                continue

            # Ayrıştırma: Birden fazla konteyner virgül, boşluk, artı veya slaş ile ayrılmış olabilir
            tokens = re.split(r"[\s,;+/]+", raw_cntr)
            valid_found = []
            for t in tokens:
                clean_t = re.sub(r"[^A-Z0-9]", "", t.upper())
                if len(clean_t) >= 6:
                    valid_found.append(clean_t)

            if not valid_found:
                clean_raw = re.sub(r"[^A-Z0-9]", "", raw_cntr.upper())
                if len(clean_raw) >= 6:
                    valid_found = [clean_raw]

            if not valid_found:
                skipped_rows.append(f"{dosya} ({raw_cntr})")
                continue

            for cntr_no in valid_found:
                try:
                    res = self.kantar_tracker.add_to_tracking(
                        konteyner_no=cntr_no,
                        dosya_no=dosya,
                        ref=ref,
                        kaynak="otomatik",
                        firma="WH-HANKOOK"
                    )
                    if res.get("success"):
                        if res.get("created"):
                            added_cntrs.append(cntr_no)
                        else:
                            updated_cntrs.append(cntr_no)
                except Exception as e:
                    print(f"[ETA-KANTAR-ERR] {cntr_no} aktarılamadı: {e}")

        # Eğer Kantar Takip sekmesi açıksa tablosunu anında yenile
        if self.main_window and hasattr(self.main_window, "views"):
            k_view = self.main_window.views.get("KANTAR_TAKIP")
            if k_view and hasattr(k_view, "refresh_kantar_takip_table"):
                k_view.refresh_kantar_takip_table()

        total_count = len(added_cntrs) + len(updated_cntrs)
        if total_count == 0:
            self.show_warning("Uyarı", "Seçilen satırlarda geçerli bir konteyner numarası bulunamadı.")
            return

        msg = f"⚖️ Toplam {total_count} adet konteyner Kantar Takip havuzuna aktarıldı!\n\n"
        if added_cntrs:
            preview = ", ".join(added_cntrs[:6])
            if len(added_cntrs) > 6:
                preview += f" ... (+{len(added_cntrs) - 6} daha)"
            msg += f"• Yeni Eklenen: {len(added_cntrs)} konteyner ({preview})\n"
        if updated_cntrs:
            msg += f"• Zaten Takipte Olan / Güncellenen: {len(updated_cntrs)} konteyner\n"
        if skipped_rows:
            msg += f"• Atlanan (Konteyner nosu bulunamayan): {len(skipped_rows)} dosya\n"

        msg += "\nŞimdi 'Kantar Takip' sekmesine geçip durumları görmek ister misiniz?"

        reply = QMessageBox.question(
            self,
            "Kantar Takip Aktarımı Tamamlandı",
            msg,
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.Yes
        )
        if reply == QMessageBox.StandardButton.Yes:
            self.switch_tab("KANTAR_TAKIP")

    # ------------------------------------------------------------------
    # CONTEXT MENU
    # ------------------------------------------------------------------
    def show_eta_context_menu(self, pos):
        selected_rows = sorted(set(idx.row() for idx in self.tree_eta.selectedIndexes()))
        if not selected_rows:
            return

        menu = QMenu(self)

        # 1. Kantar Takibe Aktarma Aksiyonu (Çoklu veya Tekli)
        if len(selected_rows) > 1:
            menu.addAction(
                f"⚖️ Seçili {len(selected_rows)} Satırı Kantar Takibe Aktar",
                self.transfer_selected_to_kantar_tracking
            )
        else:
            row = selected_rows[0]
            cntr = self.tree_eta.item(row, 5).text().strip() if self.tree_eta.item(row, 5) else ""
            label = f"⚖️ Kantar Takibe Aktar: {cntr}" if cntr and cntr != "—" else "⚖️ Kantar Takibe Aktar"
            menu.addAction(label, self.transfer_selected_to_kantar_tracking)

        menu.addAction("🔘 Tüm Satırları Seç (Ctrl+A)", self.tree_eta.selectAll)
        menu.addSeparator()

        row = selected_rows[0]
        dosya = self.tree_eta.item(row, 0).text() if self.tree_eta.item(row, 0) else ""
        ref = self.tree_eta.item(row, 1).text() if self.tree_eta.item(row, 1) else ""
        cntr = self.tree_eta.item(row, 5).text() if self.tree_eta.item(row, 5) else ""

        menu.addAction(f"📋 Operasyon Kokpitinde Aç: {dosya}", lambda: self.on_eta_row_double_click(row))
        menu.addSeparator()

        if cntr and cntr != "—":
            menu.addAction(f"🔍 Kantar Tikesini Anlık Sorgula: {cntr}", lambda: self._query_kantar_cntr(cntr))
            menu.addAction(f"📄 Kantar PDF'ini Aç / İndir: {cntr}", lambda: self._download_kantar_pdf(cntr))
            menu.addSeparator()

        menu.addAction("🌐 Kantar Takip Sekmesine Git", lambda: self.switch_tab("KANTAR_TAKIP"))
        menu.addAction("📋 Referans / Dosya Kopyala", lambda: self.copy_to_clipboard(f"{dosya} - {ref}"))
        menu.addSeparator()
        menu.addAction("📊 ETA Listesini Excel'e Aktar", self.export_eta_yaklasanlar_excel)
        menu.exec(self.tree_eta.viewport().mapToGlobal(pos))

    def _query_kantar_cntr(self, cntr_no: str):
        clean_c = re.sub(r"[^A-Z0-9]", "", str(cntr_no or "").upper())
        if not clean_c or len(clean_c) < 4:
            self.show_warning("Uyarı", "Geçerli bir konteyner numarası bulunamadı.")
            return

        if not self.kantar_tracker:
            self.show_warning("Kantar Servisi", "Kantar takip servisi aktif değil.")
            return

        def _worker():
            res = self.kantar_tracker.query_single_now(clean_c)
            def _done():
                if res.get("success"):
                    hazir = res.get("hazir")
                    msg = f"📦 Konteyner: {clean_c}\n"
                    msg += f"📊 Durum: {'🟢 Tartım Hazır' if hazir else '⏳ Tartım Bekliyor'}\n"
                    msg += f"🧾 Tartım No: {res.get('tartim_no') or '—'}\n"
                    msg += f"📅 Tarih: {res.get('tarih') or '—'}\n"
                    msg += f"🚛 Plaka: {res.get('plaka') or '—'}\n"
                    msg += f"⚖️ DBA (VGM): {res.get('dba') or '—'}\n"
                    self.show_info("DP World & Körfez Kantar Sonucu", msg)
                else:
                    self.show_error("Hata", f"Kantar sorgulanamadı:\n{res.get('error')}")
            QTimer.singleShot(0, _done)

        import threading
        threading.Thread(target=_worker, daemon=True).start()

    def _download_kantar_pdf(self, cntr_no: str):
        clean_c = re.sub(r"[^A-Z0-9]", "", str(cntr_no or "").upper())
        if not clean_c or len(clean_c) < 4:
            self.show_warning("Uyarı", "Geçerli bir konteyner numarası bulunamadı.")
            return

        if not self.kantar_tracker:
            self.show_warning("Kantar Servisi", "Kantar takip servisi aktif değil.")
            return

        fn = getattr(self.kantar_tracker, "download_and_get_pdf", None) or getattr(self.kantar_tracker, "download_pdf_tike", None)
        if not fn:
            return

        def _worker():
            pdf_path = fn(clean_c)
            def _done():
                if pdf_path and os.path.exists(pdf_path):
                    try:
                        os.startfile(pdf_path)
                    except Exception:
                        self.show_info("PDF İndirildi", f"Kantar tikesi başarıyla indirildi:\n{pdf_path}")
                else:
                    self.show_warning("Tike Bulunamadı", f"'{clean_c}' konteyneri için henüz bir kantar tikesi veya PDF oluşturulmamış.\nLiman tartımının tamamlanması bekleniyor.")
            QTimer.singleShot(0, _done)

        import threading
        threading.Thread(target=_worker, daemon=True).start()

    # ------------------------------------------------------------------
    # EXCEL EXPORT
    # ------------------------------------------------------------------
    def export_eta_yaklasanlar_excel(self):
        if not openpyxl:
            self.show_error("Hata", "openpyxl modülü yüklü değil.")
            return

        desktop = os.path.join(os.path.expanduser("~"), "Desktop")
        fname = f"ETA_Yaklasan_Dosyalar_{dt.now().date().strftime('%d_%m_%Y')}.xlsx"
        initial_path = os.path.join(desktop, fname)

        dest_path, _ = QFileDialog.getSaveFileName(
            self, "ETA Raporunu Excel Olarak Kaydet", initial_path, "Excel Dosyası (*.xlsx)"
        )
        if not dest_path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ETA Yaklaşan Dosyalar"

        headers = ["Dosya No", "Referans No", "Acenta", "Liman", "Gemi", "Konteyner / Konş.", "ETA Tarihi", "Kalan Gün"]
        ws.append(headers)

        header_fill = openpyxl.styles.PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        header_font = openpyxl.styles.Font(color="FFFFFF", bold=True, size=10)
        thin_border = openpyxl.styles.Border(
            left=openpyxl.styles.Side(style="thin", color="E2E8F0"),
            right=openpyxl.styles.Side(style="thin", color="E2E8F0"),
            top=openpyxl.styles.Side(style="thin", color="E2E8F0"),
            bottom=openpyxl.styles.Side(style="thin", color="E2E8F0")
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")

        row_num = 2
        for r in range(self.tree_eta.rowCount()):
            row_vals = [
                self.tree_eta.item(r, 0).text() if self.tree_eta.item(r, 0) else "",
                self.tree_eta.item(r, 1).text() if self.tree_eta.item(r, 1) else "",
                self.tree_eta.item(r, 2).text() if self.tree_eta.item(r, 2) else "",
                self.tree_eta.item(r, 3).text() if self.tree_eta.item(r, 3) else "",
                self.tree_eta.item(r, 4).text() if self.tree_eta.item(r, 4) else "",
                self.tree_eta.item(r, 5).text() if self.tree_eta.item(r, 5) else "",
                self.tree_eta.item(r, 6).text() if self.tree_eta.item(r, 6) else "",
                self.tree_eta.item(r, 7).text() if self.tree_eta.item(r, 7) else "",
            ]
            ws.append(row_vals)
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_num, column=col_idx)
                c.border = thin_border
                c.alignment = openpyxl.styles.Alignment(
                    horizontal="center" if col_idx in [1, 2, 4, 7, 8] else "left",
                    vertical="center"
                )
            row_num += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        try:
            wb.save(dest_path)
            reply = QMessageBox.question(
                self, "Excel İndirildi",
                f"ETA Yaklaşan Dosyalar başarıyla kaydedildi:\n{dest_path}\n\nDosya açılsın mı?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.Yes:
                try:
                    os.startfile(dest_path)
                except Exception:
                    pass
        except Exception as e:
            self.show_error("Hata", f"Excel kaydedilemedi:\n{e}")
