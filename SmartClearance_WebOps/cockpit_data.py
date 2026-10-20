# -*- coding: utf-8 -*-
"""Operasyon Kokpiti veri okuma (Excel → Evrim JSON → paket içi CSV)."""
from __future__ import annotations

import csv
import json
import os
import re
from pathlib import Path
from typing import Any, Dict, List

PACKAGE_ROOT = Path(__file__).resolve().parent
CSV_FALLBACK = PACKAGE_ROOT / "data" / "Operasyon_Kokpiti.csv"
# Üst repodan da dene (monorepo içinde çalışıyorsa)
ROOT_CSV = PACKAGE_ROOT.parent / "Operasyon_Kokpiti.csv"
CONFIG_DIR = Path(os.environ.get("LOCALAPPDATA", Path.home())) / "SmartClearance"
EVRIM_PATH = CONFIG_DIR / "evrim_live_data.json"


def _load_json(path: Path, default=None):
    if default is None:
        default = []
    if not path.exists():
        return default
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def _from_excel() -> List[Dict[str, Any]]:
    try:
        import openpyxl
    except ImportError:
        return []

    cfg = CONFIG_DIR / "config.json"
    excel_path = ""
    if cfg.exists():
        try:
            excel_path = str((_load_json(cfg, {}) or {}).get("excel_path") or "")
        except Exception:
            excel_path = ""
    if not excel_path or not os.path.exists(excel_path):
        return []

    out: List[Dict[str, Any]] = []
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet = "GÜNCEL RAPOR" if "GÜNCEL RAPOR" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet]

        def clean_date(d_val):
            if not d_val or str(d_val).lower() == "none":
                return ""
            d_str = str(d_val).strip()
            if " " in d_str:
                d_str = d_str.split()[0]
            if "-" in d_str:
                parts = d_str.split("-")
                if len(parts) == 3 and len(parts[0]) == 4:
                    return f"{parts[2]}.{parts[1]}.{parts[0]}"
            return d_str

        for r in range(2, ws.max_row + 1):
            dosya = str(ws.cell(r, 1).value or "").strip()
            ref = str(ws.cell(r, 3).value or "").strip()
            bayi = str(ws.cell(r, 4).value or "").strip()
            bl = str(ws.cell(r, 6).value or "").strip()
            ordino = clean_date(ws.cell(r, 7).value)
            tareks = clean_date(ws.cell(r, 9).value)
            tescil = clean_date(ws.cell(r, 12).value)
            hat = str(ws.cell(r, 13).value or "").strip().upper()
            kapanis = clean_date(ws.cell(r, 14).value)
            vergi = clean_date(ws.cell(r, 15).value)
            liman = str(ws.cell(r, 22).value or "").strip()
            if not (dosya or ref or bl):
                continue
            is_tescil = bool(tescil)
            is_tax = is_tescil and not vergi
            if kapanis:
                stage, title = "5_KAPANDI", "Kapandı / Arşiv"
            elif is_tescil and hat == "KIRMIZI":
                stage, title = "3_MUAYENE", "Kırmızı Hat (Muayenede)"
            elif is_tescil:
                stage, title = "3_TESCILDE", "Tescilde"
            elif ordino or tareks:
                stage, title = "2_ORDINO_TAREKS", "Ordino / Tareks"
            else:
                stage, title = "1_EVRAK", "Evrak"
            out.append(
                {
                    "dosya": dosya,
                    "ref": ref,
                    "bl": bl,
                    "stage": stage,
                    "stage_title": title,
                    "hat": hat,
                    "vergi_mail": "Vergi bekliyor" if is_tax else "",
                    "is_tax_pending": is_tax,
                    "bayi": bayi,
                    "liman": liman,
                    "beyanname_no": "",
                    "kapandi": stage == "5_KAPANDI",
                    "kaynak": "excel",
                }
            )
    except Exception:
        return []
    return out


def _from_evrim() -> List[Dict[str, Any]]:
    rows = _load_json(EVRIM_PATH, [])
    out = []
    for r in rows or []:
        if not isinstance(r, dict):
            continue
        kapandi = bool(r.get("kapanis") or r.get("kapandi"))
        hat = str(r.get("hat") or "").strip()
        tescil = str(r.get("tescil_no") or r.get("tescil_tarihi") or "").strip()
        is_tax = bool(r.get("is_tax_pending"))
        if kapandi:
            stage, title = "5_KAPANDI", "Kapandı / Arşiv"
        elif "KIRMIZI" in hat.upper():
            stage, title = "3_MUAYENE", "Kırmızı Hat"
        elif tescil or r.get("beyanname_no"):
            stage, title = "3_TESCILDE", "Tescilde"
        else:
            stage, title = "1_EVRAK", "Evrak"
        out.append(
            {
                "dosya": str(r.get("dosya_no") or "").strip(),
                "ref": str(r.get("referans_no") or r.get("hankook_ref") or "").strip(),
                "bl": str(r.get("bl_no") or "").strip(),
                "stage": stage,
                "stage_title": title,
                "hat": hat,
                "vergi_mail": "Vergi bekliyor" if is_tax else "",
                "is_tax_pending": is_tax,
                "bayi": str(r.get("musteri_unvan") or r.get("firma") or "").strip(),
                "liman": str(r.get("liman") or r.get("gumruk") or "").strip(),
                "beyanname_no": str(r.get("beyanname_no") or "").strip(),
                "kapandi": kapandi,
                "kaynak": "evrim",
            }
        )
    return out


def _from_csv() -> List[Dict[str, Any]]:
    path = CSV_FALLBACK if CSV_FALLBACK.exists() else ROOT_CSV
    if not path.exists():
        return []
    out = []
    with open(path, "r", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f, delimiter=";"):
            stage_title = (row.get("Aşama") or "").strip()
            st_u = stage_title.upper()
            if "KAPAND" in st_u:
                stage = "5_KAPANDI"
            elif "KIRMIZI" in st_u or "MUAYENE" in st_u:
                stage = "3_MUAYENE"
            elif "TESCİL" in st_u or "TESCIL" in st_u:
                stage = "3_TESCILDE"
            elif "ORDİNO" in st_u or "ORDINO" in st_u or "TAREKS" in st_u:
                stage = "2_ORDINO_TAREKS"
            else:
                stage = "1_EVRAK"
            vergi = (row.get("Vergi Maili") or "").strip()
            out.append(
                {
                    "dosya": (row.get("Dosya No") or "").strip(),
                    "ref": (row.get("Referans No") or "").strip(),
                    "bl": (row.get("Konşimento") or "").strip(),
                    "stage": stage,
                    "stage_title": stage_title,
                    "hat": (row.get("Hat") or "").strip(),
                    "vergi_mail": vergi,
                    "is_tax_pending": "BEKLE" in vergi.upper(),
                    "bayi": (row.get("Bayi / Alıcı") or "").strip(),
                    "liman": (row.get("Liman") or "").strip(),
                    "beyanname_no": "",
                    "kapandi": stage == "5_KAPANDI",
                    "kaynak": "csv",
                }
            )
    return out


def load_cockpit_records(q: str = "", filter_type: str = "all", limit: int = 200) -> Dict[str, Any]:
    source = "empty"
    records = _from_excel()
    if records:
        source = "excel"
    else:
        records = _from_evrim()
        if records:
            source = "evrim"
        else:
            records = _from_csv()
            if records:
                source = "csv"

    q_clean = (q or "").strip().upper()
    ft = (filter_type or "all").strip().lower()
    filtered = []
    for r in records:
        kapandi = bool(r.get("kapandi"))
        hat = str(r.get("hat") or "").upper()
        is_tax = bool(r.get("is_tax_pending"))
        stage = str(r.get("stage") or "")
        if ft == "active" and kapandi:
            continue
        if ft == "closed" and not kapandi:
            continue
        if ft == "tax" and not is_tax:
            continue
        if ft == "tescil" and not stage.startswith("3"):
            continue
        if ft == "red" and "KIRMIZI" not in hat and "MUAYENE" not in str(r.get("stage_title") or "").upper():
            continue
        if ft == "ordino" and not stage.startswith("2"):
            continue
        if q_clean:
            blob = " ".join(str(r.get(k) or "") for k in ("dosya", "ref", "bl", "bayi", "liman", "stage_title", "hat")).upper()
            if q_clean not in blob:
                continue
        filtered.append(r)
        if len(filtered) >= limit:
            break

    return {
        "source": source,
        "count": len(filtered),
        "stats": {
            "total": len(records),
            "active": sum(1 for r in records if not r.get("kapandi")),
            "tax": sum(1 for r in records if r.get("is_tax_pending")),
            "closed": sum(1 for r in records if r.get("kapandi")),
        },
        "items": filtered,
    }


def extract_beyanname_candidates(text: str) -> List[str]:
    found, seen = [], set()
    for part in re.split(r"[\s,;/\n\r]+", text or ""):
        clean = re.sub(r"[^A-Z0-9]", "", part.upper())
        if len(clean) >= 15 and clean not in seen:
            seen.add(clean)
            found.append(clean)
    return found


def load_evrim_raw() -> List[Dict[str, Any]]:
    """GTB/Kantar auto_sync için ham Evrim listesi."""
    data = _load_json(EVRIM_PATH, [])
    return data if isinstance(data, list) else []


def load_excel_raw_for_sync() -> List[Dict[str, Any]]:
    """Kantar sync için Excel satırlarını konteyner alanıyla döner."""
    try:
        import openpyxl
    except ImportError:
        return []

    cfg = CONFIG_DIR / "config.json"
    excel_path = ""
    if cfg.exists():
        try:
            excel_path = str((_load_json(cfg, {}) or {}).get("excel_path") or "")
        except Exception:
            excel_path = ""
    if not excel_path or not os.path.exists(excel_path):
        return []

    out: List[Dict[str, Any]] = []
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet = "GÜNCEL RAPOR" if "GÜNCEL RAPOR" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet]
        for r in range(2, ws.max_row + 1):
            dosya = str(ws.cell(r, 1).value or "").strip()
            ref = str(ws.cell(r, 3).value or "").strip()
            bayi = str(ws.cell(r, 4).value or "").strip()
            liman = str(ws.cell(r, 22).value or "").strip()
            konteyner = str(ws.cell(r, 21).value or "").strip()
            beyan = ""
            # Evrim eşlemesi için beyanname yoksa boş
            if not (dosya or ref or konteyner):
                continue
            out.append(
                {
                    "dosya": dosya,
                    "dosya_no": dosya,
                    "ref": ref,
                    "hankook_ref": ref,
                    "bayi": bayi,
                    "liman": liman,
                    "konteyner": konteyner,
                    "beyanname_no": beyan,
                }
            )
            # GTB sync için de beyanname sütunu varsa (bazı şablonlarda farklı) — atlanır
    except Exception:
        return []
    return out
